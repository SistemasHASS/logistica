# -*- coding: utf-8 -*-
"""
Genera DOCUMENTACION_DIAGRAMA_BD_DETALLE.docx:
  - DER detallado por modulo (todas las columnas con su tipo de dato)
  - Ficha por tabla: cada columna con tipo de dato, restricciones y detalle
Entrada: doc_output/schema.json (producido por extraer_esquema.py)
"""
import os, json, math
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(HERE, "..", "docs")
FIG = os.path.join(HERE, "doc_output", "figs_bd")
os.makedirs(FIG, exist_ok=True)

def _find(name):
    for base in (os.path.join(HERE, "doc_output"), os.path.join(HERE, "..", "doc_output")):
        p = os.path.join(base, name)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(name)

SCHEMA = json.load(open(_find("schema.json"), encoding="utf-8"))

# ---------- clasificacion por modulo ----------
def modulo(t):
    n = t.lower()
    if n.startswith(("wh_", "spring_", "logistica_colamigracion",
                     "logistica_sincronizacion", "logistica_maestroitem")):
        return "Integracion ERP Spring"
    if n.startswith("logistica_consol"):
        return "Consolidacion OC/OS"
    if "servicio" in n or n.endswith(("aprobacionos", "flujoaprobacionos",
        "usuarioaprobadoros", "osconformidad", "aprobacionss")) \
        or "confirmacionservicio" in n or "conformidadservicio" in n \
        or "historialprecioservicio" in n or "plantillaconformidadservicio" in n:
        return "Servicios"
    if "consolidacion" in n or "itempendienteconsolidacion" in n \
        or "consolidado" in n or n == "reqconsolidado":
        return "Consolidacion de Requerimientos"
    if "aprobacion" in n or "flujo" in n or "aprobadores" in n \
        or "saldopendiente" in n or "adjuntosorden" in n \
        or n in ("logistica_areas", "logistica_usuariosporarea"):
        return "Aprobaciones"
    if n in ("logistica_req", "logistica_dreq") or "reqcommodity" in n \
        or "solicitudcreacionitem" in n or "historialanulaciones" in n \
        or "logcambioestadorequerimiento" in n or n == "logistica_requerimientos" \
        or n == "logistica_requerimientosporarea":
        return "Requerimientos"
    if "requerimientotransferencia" in n:
        return "Almacen e Inventario"
    if "recepcion" in n or "despacho" in n or "kardex" in n or "stock" in n \
        or "inventario" in n or "transaccion" in n or "devolucion" in n \
        or "reingreso" in n or "almacen" in n or n == "logistica_item" \
        or "conformidadalmacen" in n or "plantillaalmacen" in n:
        return "Almacen e Inventario"
    if "cotizacion" in n or "solicitudcompra" in n or "ordencompra" in n \
        or "ocadjunto" in n or "seguimientoordencompra" in n \
        or "incrementooc" in n or "proveedor" in n or n == "logistica_entrega":
        return "Compras"
    if "notificacion" in n or "correolog" in n or "alerta" in n:
        return "Notificaciones"
    return "Administracion y Configuracion"

MODULOS_ORDEN = [
    "Requerimientos", "Aprobaciones", "Consolidacion de Requerimientos",
    "Compras", "Consolidacion OC/OS", "Servicios", "Almacen e Inventario",
    "Notificaciones", "Administracion y Configuracion", "Integracion ERP Spring",
]
MODULO_TABLAS = {m: [] for m in MODULOS_ORDEN}
for t in SCHEMA:
    MODULO_TABLAS[modulo(t)].append(t)
for m in MODULO_TABLAS:
    MODULO_TABLAS[m].sort(key=lambda x: x.lower())

DESCR_MODULO = {
    "Requerimientos": "Requerimientos internos de bienes/servicios, su detalle por item, commodity y adjuntos.",
    "Aprobaciones": "Flujos de aprobacion por area, aprobadores, niveles y log de acciones.",
    "Consolidacion de Requerimientos": "Agrupacion de requerimientos aprobados para compra conjunta.",
    "Compras": "Solicitudes de cotizacion, cotizaciones de proveedores, solicitudes de compra, ordenes de compra y seguimiento.",
    "Consolidacion OC/OS": "Consolidacion de ordenes de compra y de servicio generadas en el nuevo flujo.",
    "Servicios": "Solicitudes de servicio, cotizaciones de servicio, ordenes de servicio y conformidad.",
    "Almacen e Inventario": "Recepciones, despachos, kardex, stock, transacciones de inventario, devoluciones y reingresos.",
    "Notificaciones": "Notificaciones y log de correos del sistema.",
    "Administracion y Configuracion": "Usuarios, roles, perfiles, areas, menus, permisos, parametros y auditoria.",
    "Integracion ERP Spring": "Tablas espejo y colas de sincronizacion con el ERP Spring.",
}

# ---------- helpers docx ----------
def estilos(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10)
    for i, (sz, col) in enumerate([(20, "1F4E79"), (14, "1F4E79"), (12, "2E74B5"), (11, "404040")], 1):
        h = doc.styles[f"Heading {i}"]
        h.font.size = Pt(sz); h.font.color.rgb = RGBColor.from_string(col); h.font.bold = True

def _shd(color):
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), color)
    return el

def tabla(doc, headers, rows, size=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        c._tc.get_or_add_tcPr().append(_shd("1F4E79"))
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size); r.font.name = "Consolas" if i == 1 else "Calibri"
    return t

def portada(doc, titulo, subtitulo):
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0x1F,0x4E,0x79)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitulo); r.font.size = Pt(16); r.font.color.rgb = RGBColor(0x40,0x40,0x40)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Base de datos LOGISTICA - SQL Server | Documento generado desde los scripts SQL del proyecto")
    r.font.size = Pt(10); r.italic = True
    doc.add_page_break()

# ---------- figura ER detallada por modulo ----------
def figura_er_modulo(tablas, fname):
    n = len(tablas)
    ncols = min(4, max(2, math.ceil(math.sqrt(n))))
    BW, ROW_H, HDR, GAPX, GAPY = 3.6, 0.24, 0.42, 0.9, 0.9
    filas = [tablas[i:i+ncols] for i in range(0, n, ncols)]
    pos = {}; y = 0
    for fila in filas:
        hmax = max(HDR + ROW_H * len(SCHEMA[t]["cols"]) for t in fila)
        x = 0
        for t in fila:
            pos[t] = (x, y); x += BW + GAPX
        y += hmax + GAPY
    W = ncols * (BW + GAPX); H = y + 0.5
    fig, ax = plt.subplots(figsize=(W * 0.55, H * 0.55), dpi=170)
    ax.set_xlim(-0.3, W); ax.set_ylim(-0.3, H); ax.axis("off"); ax.invert_yaxis()
    for t, (x, y0) in pos.items():
        cols = SCHEMA[t]["cols"]
        h = HDR + ROW_H * len(cols)
        ax.add_patch(FancyBboxPatch((x, y0), BW, h, boxstyle="round,pad=0.03",
                                    fc="#FDFEFE", ec="#2E74B5", lw=1.2))
        ax.add_patch(plt.Rectangle((x, y0), BW, HDR, fc="#1F4E79", ec="#1F4E79"))
        ax.text(x + BW/2, y0 + HDR/2, t, ha="center", va="center",
                fontsize=6.2, fontweight="bold", color="white")
        fks_cols = {c.lower() for f in SCHEMA[t].get("fks", []) for c in (f["cols"] or "").split(",")}
        for i, c in enumerate(cols):
            yy = y0 + HDR + i * ROW_H + ROW_H/2
            pk = "PK" in (c.get("flags") or "")
            fk = c["name"].lower() in fks_cols or c["name"].lower() in {s.strip() for fc in fks_cols for s in fc.split("+")}
            marca = ("PK " if pk else "") + ("FK " if fk else "")
            ax.text(x + 0.08, yy, f"{marca}{c['name']}", fontsize=4.6, va="center",
                    fontweight="bold" if pk else "normal",
                    color="#C00000" if pk else ("#B45F06" if fk else "#222222"))
            ax.text(x + BW - 0.08, yy, c["type"], fontsize=4.2, va="center",
                    ha="right", color="#777777")
    for t, (x, y0) in pos.items():
        for fk in SCHEMA[t].get("fks", []):
            ref = fk["ref"]
            if ref in pos:
                x2, y2 = pos[ref]
                ax.annotate("",
                    xy=(x2 + BW/2, y2 - 0.05),
                    xytext=(x + BW/2, y0 + HDR + ROW_H * len(SCHEMA[t]["cols"]) + 0.05),
                    arrowprops=dict(arrowstyle="->", color="#B45F06", lw=0.7,
                                    alpha=0.75, connectionstyle="arc3,rad=0.2"))
    fig.savefig(fname, bbox_inches="tight", facecolor="white")
    plt.close(fig)

# =====================================================================
doc = Document(); estilos(doc)
portada(doc, "Diagrama de Base de Datos - Detalle por Tabla",
        "Esquema completo con tipos de datos de cada columna")

doc.add_heading("1. Generalidades", 1)
doc.add_paragraph(
    f"Motor: Microsoft SQL Server. Base de datos: LOGISTICA. "
    f"El esquema contiene {len(SCHEMA)} tablas agrupadas en {len(MODULOS_ORDEN)} modulos "
    "funcionales. El detalle se obtuvo de los scripts CREATE TABLE y ALTER TABLE del "
    "proyecto api_logistica (carpeta SQL).")
doc.add_paragraph(
    "Convenciones: las columnas PK aparecen en rojo en los diagramas y con marca PK en las "
    "fichas; las columnas FK en naranja. La columna Restricciones consolida NOT NULL, "
    "DEFAULT, IDENTITY y demas atributos declarados en el script. La columna Detalle "
    "muestra valores permitidos o comentarios presentes en el script cuando existen.")

doc.add_heading("2. Indice de modulos y tablas", 1)
rows = []
for m in MODULOS_ORDEN:
    for t in MODULO_TABLAS[m]:
        rows.append([m, t, len(SCHEMA[t]["cols"]),
                     SCHEMA[t].get("src", "") or "-"])
tabla(doc, ["Modulo", "Tabla", "N° columnas", "Script origen"], rows, size=7)
doc.add_page_break()

nfig = 0
for i, mod in enumerate(MODULOS_ORDEN, 3):
    tablas = MODULO_TABLAS[mod]
    if not tablas:
        continue
    doc.add_heading(f"{i}. Modulo: {mod}", 1)
    doc.add_paragraph(DESCR_MODULO.get(mod, ""))
    doc.add_paragraph(f"{len(tablas)} tablas.")

    fname = os.path.join(FIG, f"bd_{i}_{mod.split()[0].lower()}.png")
    figura_er_modulo(tablas, fname)
    doc.add_picture(fname, width=Inches(7.0))
    p = doc.add_paragraph(f"Figura {i-2}. DER del modulo {mod} (PK en rojo, FK en naranja).")
    p.runs[0].font.size = Pt(8); p.runs[0].italic = True
    nfig += 1

    doc.add_heading(f"{i}.1 Detalle de tablas", 2)
    for t in tablas:
        info = SCHEMA[t]
        doc.add_heading(t, 3)
        met = doc.add_paragraph()
        met.add_run("Script: ").bold = True
        met.add_run(info.get("src") or "-").font.size = Pt(8)
        met.add_run("   |   Columnas: ").bold = True
        met.add_run(str(len(info["cols"]))).font.size = Pt(8)
        rows = []
        for c in info["cols"]:
            flags = (c.get("flags") or "").strip()
            rows.append([c["name"], c["type"], flags or "-", c.get("note") or "-"])
        tabla(doc, ["Columna", "Tipo de dato", "Restricciones", "Detalle"], rows, size=7.5)
        if info.get("fks"):
            p = doc.add_paragraph()
            p.add_run("Llaves foraneas: ").bold = True
            p.add_run("; ".join(
                f"{f['cols']} -> {f['ref']}.{f.get('refcols','')}" for f in info["fks"])
            ).font.size = Pt(8)

doc.save(os.path.join(DOCS, "DOCUMENTACION_DIAGRAMA_BD_DETALLE.docx"))
print("OK", nfig, "figuras,", len(SCHEMA), "tablas")
