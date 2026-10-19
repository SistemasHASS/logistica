# -*- coding: utf-8 -*-
"""
Genera MANUAL_CARGA_MASIVA_USUARIOS.docx y .pdf
Explica como funciona la carga masiva de usuarios con roles:
- DONDE se ejecuta (SQL Server, no desde la app web)
- COMO funciona el SP LOGISTICA_cargaMasivaConfiguracion
- Que roles van en cada tabla
- Paso a paso completo Excel -> JSON -> SQL -> ejecucion
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
from reportlab.lib import colors

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(BASE, "docs")
os.makedirs(OUT, exist_ok=True)

AZUL = "1F4E78"
GRIS = "F2F2F2"

# Bloques: ("h1"|"h2"|"p"|"bullet"|"table"|"code", ...)
C = [
    ("h1", "MANUAL DE CARGA MASIVA - USUARIOS Y ROLES"),
    ("p", "Sistema: App Logistica | Base de datos: LOGISTICA (SQL Server)"),
    ("p", "Este documento explica DONDE se ejecuta la carga masiva, COMO "
          "funciona el procedimiento y COMO se asignan los roles a los usuarios "
          "en cada tabla del sistema."),

    # ================= SECCION 1 =================
    ("h1", "1. DESDE DONDE SE HACE LA CARGA MASIVA"),
    ("p", "Punto clave: la carga masiva NO se hace desde un modulo de la "
          "aplicacion web. Se ejecuta directamente en la base de datos usando "
          "SQL Server Management Studio (SSMS)."),
    ("h2", "1.1 Dos formas de cargar usuarios"),
    ("table",
     ["Forma", "Donde se ejecuta", "Para que sirve"],
     [
         ["Alta manual (uno por uno)",
          "App web -> modulo 'Admin Logistica' (administradores) y pantallas "
          "Aprobaciones OC / Aprobaciones OS",
          "Registrar o corregir usuarios individuales"],
         ["Carga masiva (todo de una vez)",
          "SSMS (SQL Server) -> base de datos LOGISTICA -> "
          "EXEC LOGISTICA_cargaMasivaConfiguracion",
          "Configurar docenas/cientos de usuarios, areas y aprobaciones "
          "en una sola ejecucion"],
     ]),
    ("p", "El modulo 'Admin Logistica' de la app sigue existiendo y sirve para "
          "el alta manual, pero cuando hay que configurar todo el sistema "
          "(muchos usuarios, varias empresas, flujos de aprobacion) la carga "
          "masiva desde BD ahorra dias de digitacion."),
    ("h2", "1.2 Que herramienta se usa"),
    ("bullet", [
        "SQL Server Management Studio (SSMS) conectado a la BD LOGISTICA.",
        "Permiso para ejecutar INSERT/UPDATE en las tablas LOGISTICA_*.",
        "El SP LOGISTICA_cargaMasivaConfiguracion creado previamente "
        "(archivo api_logistica/SQL/750_SP_CARGA_MASIVA_CONFIGURACION.sql).",
        "Opcional pero recomendado: el Excel tecnico lleno + el script "
        "generar_json_carga_masiva.py que arma el EXEC automaticamente.",
    ]),

    # ================= SECCION 2 =================
    ("h1", "2. COMO FUNCIONA EL SP DE CARGA MASIVA"),
    ("p", "El SP LOGISTICA_cargaMasivaConfiguracion recibe UN solo parametro: "
          "un texto JSON con toda la configuracion. En una sola transaccion "
          "procesa 8 secciones. Si algo falla, se revierte TODO (ningun dato "
          "queda a medias)."),
    ("h2", "2.1 Secciones del JSON y tabla destino"),
    ("table",
     ["Seccion JSON", "Tabla destino", "Clave para NO duplicar"],
     [
         ["areas", "LOGISTICA_Areas", "ruc + nombre del area"],
         ["usuariosAcceso", "LOGISTICA_AdminUsuarios", "usuario o nrodocumento"],
         ["usuariosPorArea", "LOGISTICA_UsuariosPorArea",
          "ruc + documentoidentidad + idarea"],
         ["flujoReq", "LOGISTICA_FlujoAprobacionPorArea",
          "ruc + idarea + tipoRequerimiento + rolAprobador"],
         ["flujoOC", "LOGISTICA_FlujoAprobacionOC",
          "montoMinimo + montoMaximo + nivel"],
         ["aprobadoresOC", "LOGISTICA_UsuariosAprobadoresOC",
          "documentoIdentidad + rol"],
         ["flujoOS", "LOGISTICA_FlujoAprobacionOS",
          "montoDesde + montoHasta"],
         ["aprobadoresOS", "LOGISTICA_UsuarioAprobadorOS",
          "documentoIdentidad"],
     ]),
    ("h2", "2.2 Comportamiento"),
    ("bullet", [
        "Es RE-EJECUTABLE: si el registro ya existe lo actualiza; si no, lo "
        "inserta. Se puede correr varias veces sin duplicar.",
        "En usuariosPorArea y flujoReq se puede mandar el idarea numerico O "
        "el nombre del area ('area': 'ALMACEN'); el SP busca el id por "
        "ruc + nombre automaticamente.",
        "Si un usuario referencia un area que no existe, NO se inserta y "
        "aparece en el resultado como advertencia (nada se pierde en silencio).",
        "Las claves pueden ir en texto plano ('clave') y el SP las hashea con "
        "SHA2_256, o ya hasheadas ('claveHash').",
        "Al terminar devuelve un JSON con: resumen por seccion "
        "(insertados / actualizados / omitidos) y lista de advertencias.",
    ]),
    ("h2", "2.3 Ejemplo de la respuesta del SP"),
    ("code",
"""{"status":"success","mensaje":"Carga masiva ejecutada",
 "resumen":[
   {"seccion":"AREAS","insertados":4,"actualizados":16,"omitidos":0},
   {"seccion":"USUARIOS_ACCESO","insertados":12,"actualizados":0,"omitidos":3},
   {"seccion":"USUARIOS_POR_AREA","insertados":30,"actualizados":5,"omitidos":0},
   {"seccion":"FLUJO_REQ_AREA","insertados":15,"actualizados":0,"omitidos":0},
   {"seccion":"FLUJO_OC","insertados":0,"actualizados":3,"omitidos":0},
   {"seccion":"APROBADORES_OC","insertados":2,"actualizados":0,"omitidos":0}],
 "advertencias":[
   {"seccion":"USUARIOS_POR_AREA",
    "detalle":"Area no encontrada [ACOPIO] para 12345678 PEREZ JUAN"}]}"""),

    # ================= SECCION 3 =================
    ("h1", "3. COMO SE ASIGNAN LOS ROLES (3 capas)"),
    ("p", "En el sistema un usuario puede tener hasta 3 definiciones de rol, "
          "en tablas distintas. Entender esto es clave para llenar bien el JSON:"),
    ("h2", "3.1 Capa 1 - Cuenta de acceso (LOGISTICA_AdminUsuarios)"),
    ("p", "Es el usuario y clave para entrar a la app. Campos: usuario "
          "(normalmente el DNI), clave, nombre, email, idrol + rol, ruc, "
          "idempresa. El rol de esta tabla define los permisos generales "
          "de la cuenta."),
    ("h2", "3.2 Capa 2 - Rol dentro del area (LOGISTICA_UsuariosPorArea)"),
    ("p", "Vincula a la persona con su area y define que hace dentro del "
          "flujo de requerimientos. Campos clave: rol, esJefeArea (1=es jefe "
          "del area), esAprobador (1=puede aprobar requerimientos)."),
    ("h2", "3.3 Capa 3 - Aprobador de OC / OS"),
    ("p", "LOGISTICA_UsuariosAprobadoresOC y LOGISTICA_UsuarioAprobadorOS "
          "definen quienes firman las ordenes de compra y servicio, con su "
          "rol aprobador, rango de monto y nivel (orden de firma). "
          "Esta configuracion es GLOBAL (no lleva RUC)."),
    ("h2", "3.4 Catalogo de roles"),
    ("table",
     ["Codigo (idrol)", "Nombre", "Uso tipico"],
     [
         ["SOLICITANTE", "Solicitante", "Crea requerimientos en su area"],
         ["APLOGIST", "Aprobador Consumo", "Jefe/aprueba req de consumo y compra"],
         ["JEFE_AREA", "Jefe de Area", "Rol aprobador en el flujo por area"],
         ["ALLOGIST", "Almacen", "Recepcion, despachos, transferencias"],
         ["OPLOGIST", "Operativo Campo", "Personal de campo que solicita"],
         ["LOLOGIST", "Operador Logistica", "Compras, cotizaciones, OC/OS"],
         ["EMLOGIST", "Operador Licitaciones", "Licitaciones/empaque"],
         ["JLOLOGIST", "Jefe Logistica", "Jefatura de logistica"],
         ["JEMLOGIST", "Jefe Licitaciones", "Jefatura licitaciones/empaque"],
         ["FINANZAS", "Finanzas", "Aprueba OC/OS por monto"],
         ["GERENTE", "Gerencia", "Aprobacion gerencial"],
         ["ADLOGIST", "Admin Logistica", "Administrador funcional"],
         ["TILOGIST", "Admin Sistema", "Administrador total del sistema"],
     ]),
    ("h2", "3.5 Combinaciones tipicas por tipo de persona"),
    ("table",
     ["Tipo de persona", "usuariosAcceso.idrol", "usuariosPorArea.rol", "esJefe / esAprob"],
     [
         ["Solicitante (crea req consumo/compra)", "SOLICITANTE", "SOLICITANTE", "0 / 0"],
         ["Jefe de area que aprueba consumo y compra", "APLOGIST", "APLOGIST", "1 / 1"],
         ["Almacenero que hace transferencias", "ALLOGIST", "ALLOGIST", "0 / 0"],
         ["Comprador de logistica", "LOLOGIST", "LOLOGIST", "0 / 0"],
         ["Aprobador de OC/OS (finanzas)", "FINANZAS", "(no requiere area)", "-"],
         ["Administrador del sistema", "TILOGIST o ADLOGIST", "-", "-"],
     ]),

    # ================= SECCION 4 =================
    ("h1", "4. PROCESO PASO A PASO"),
    ("h2", "Paso 0 - Crear el SP (una sola vez por base de datos)"),
    ("code",
"""-- En SSMS, ejecutar el archivo:
--   api_logistica/SQL/750_SP_CARGA_MASIVA_CONFIGURACION.sql
-- Crea el SP LOGISTICA_cargaMasivaConfiguracion."""),
    ("h2", "Paso 1 - Levantar la informacion"),
    ("p", "Enviar a cada jefatura el Excel FORMATO_USUARIO_LEVANTAMIENTO_"
          "LOGISTICA.xlsx. Cada fila es una persona con sus casillas SI/NO "
          "de lo que crea y lo que aprueba."),
    ("h2", "Paso 2 - Trasladar al Excel tecnico"),
    ("p", "Llenar FORMATO_TECNICO_CONFIGURACION_LOGISTICA.xlsx con los datos "
          "validados: hojas AREAS, CUENTAS_ACCESO, USUARIOS_POR_AREA, "
          "FLUJO_REQ_AREA, FLUJO_OC, APROBADORES_OC, FLUJO_OS, APROBADORES_OS."),
    ("h2", "Paso 3 - Generar el SQL de carga"),
    ("code",
"""cd logistica/scripts
python generar_json_carga_masiva.py
-- genera:
--   doc_output/carga_masiva_config.json
--   doc_output/CARGA_MASIVA.sql   (el EXEC listo para SSMS)"""),
    ("h2", "Paso 4 - Ejecutar en SSMS"),
    ("p", "Abrir doc_output/CARGA_MASIVA.sql en SSMS conectado a la BD "
          "LOGISTICA y ejecutar. Equivale a correr:"),
    ("code",
"""EXEC [dbo].[LOGISTICA_cargaMasivaConfiguracion] @json = N'{...}';"""),
    ("h2", "Paso 5 - Revisar el resultado"),
    ("p", "El SP devuelve un JSON. Revisar: status=success, los insertados por "
          "seccion, y la lista de advertencias (usuarios con area no "
          "encontrada o cuentas sin clave)."),
    ("h2", "Paso 6 - Verificacion en BD"),
    ("code",
"""SELECT a.nombre AS area, u.documentoidentidad, u.nombreCompleto,
       u.rol, u.esJefeArea, u.esAprobador
FROM LOGISTICA_UsuariosPorArea u
JOIN LOGISTICA_Areas a ON a.idarea = u.idarea
WHERE u.activo = 1 ORDER BY a.nombre;

SELECT ruc, idarea, tipoRequerimiento, rolAprobador
FROM LOGISTICA_FlujoAprobacionPorArea WHERE activo = 1;

SELECT documentoIdentidad, nombreCompleto, rol, montoMinimo,
       montoMaximo, nivel
FROM LOGISTICA_UsuariosAprobadoresOC WHERE activo = 1;

SELECT documentoIdentidad, nombre, rol, nivelMaximo
FROM LOGISTICA_UsuarioAprobadorOS WHERE activo = 1;"""),
    ("h2", "Paso 7 - Prueba funcional"),
    ("p", "Ingresar a la app con un usuario de cada tipo y verificar: el "
          "solicitante crea requerimiento, el jefe de area lo ve para aprobar, "
          "y las OC/OS llegan a los aprobadores configurados."),

    # ================= SECCION 5 =================
    ("h1", "5. EJEMPLO COMPLETO DEL JSON"),
    ("code",
"""EXEC [dbo].[LOGISTICA_cargaMasivaConfiguracion] @json = N'
{
  "usuarioCreacion": "CARGA_INICIAL",
  "areas": [
    {"ruc":"20481121966","nombre":"ACOPIOS","activo":1}
  ],
  "usuariosAcceso": [
    {"ruc":"20481121966","usuario":"48073280","clave":"MiClave2026",
     "nombre":"JULCA BARRANTES JULIO CESAR","email":"jjulca@hassperu.com",
     "idrol":"APLOGIST","rol":"Aprobador Consumo",
     "nrodocumento":"48073280","idempresa":"000008"}
  ],
  "usuariosPorArea": [
    {"ruc":"20481121966","documentoidentidad":"48073280",
     "nombreCompleto":"JULCA BARRANTES JULIO CESAR",
     "area":"SERVICIO DE TRANSPORTE","rol":"APLOGIST",
     "esJefeArea":1,"esAprobador":1,"activo":1},
    {"ruc":"20481121966","documentoidentidad":"44123456",
     "nombreCompleto":"PEREZ LOPEZ MARIA",
     "area":"RIEGO","rol":"SOLICITANTE",
     "esJefeArea":0,"esAprobador":0,"activo":1}
  ],
  "flujoReq": [
    {"ruc":"20481121966","area":"RIEGO","tipoRequerimiento":"CONSUMO",
     "rolAprobador":"JEFE_AREA"},
    {"ruc":"20481121966","area":"RIEGO","tipoRequerimiento":"COMPRA",
     "rolAprobador":"JEFE_AREA"},
    {"ruc":"20481121966","area":"ALMACEN","tipoRequerimiento":"TRANSFERENCIA",
     "rolAprobador":"JEFE_AREA"}
  ],
  "flujoOC": [
    {"montoMinimo":0,"montoMaximo":10000,"nivel":1,
     "rolAprobador":"FINANZAS","nombreRol":"Finanzas",
     "descripcion":"OC hasta S/ 10,000"},
    {"montoMinimo":10001,"montoMaximo":999999999,"nivel":1,
     "rolAprobador":"GERENTE","nombreRol":"Gerencia",
     "descripcion":"OC mayor a S/ 10,000"}
  ],
  "aprobadoresOC": [
    {"documentoIdentidad":"41780114","nombreCompleto":"MACEN MOYA MARCO",
     "email":"mmacen@hassperu.com","rol":"FINANZAS",
     "montoMinimo":0,"montoMaximo":999999999,"nivel":1,"activo":1}
  ],
  "flujoOS": [
    {"nombre":"OS hasta S/ 50,000","montoDesde":0,"montoHasta":50000,
     "niveles":1,"activo":1}
  ],
  "aprobadoresOS": [
    {"documentoIdentidad":"41780114","nombre":"MACEN MOYA MARCO",
     "rol":"FINANZAS","nivelMaximo":1,"activo":1}
  ]
}';"""),

    # ================= SECCION 6 =================
    ("h1", "6. ALTERNATIVA: USUARIOS DESDE EL ERP"),
    ("p", "Si los usuarios ya existen en el ERP (BDERP_Agro_Hass), las cuentas "
          "de acceso pueden crearse con la sincronizacion en vez de mandarlas "
          "en el JSON: existe el SP SP_Logistica_Usuarios_SincronizarERP y el "
          "endpoint POST /api/logistica/admin-logistica/usuarios/sincronizar-erp. "
          "Asi las claves vienen del ERP y no hace falta generar claves "
          "temporales. Despues de sincronizar, el resto de la configuracion "
          "(areas, roles por area, flujos) se hace con la carga masiva."),

    # ================= SECCION 7 =================
    ("h1", "7. BUENAS PRACTICAS"),
    ("bullet", [
        "Hacer backup de la BD LOGISTICA antes de la primera carga.",
        "Probar primero en ambiente de desarrollo/pruebas.",
        "El documentoidentidad es texto: conservar ceros a la izquierda.",
        "No dejar claves reales escritas en el Excel o el JSON: usar claves "
        "temporales y pedir cambio en el primer ingreso (o sincronizar ERP).",
        "Revisar SIEMPRE las advertencias del resultado: ahi aparecen los "
        "registros que no se cargaron y por que.",
        "La carga es re-ejecutable: si falto un dato, se corrige el Excel, "
        "se regenera y se vuelve a correr completo sin miedo a duplicar.",
        "La configuracion de OC/OS es global para toda la BD: si hay dos "
        "empresas, los mismos aprobadores aplican a ambas.",
    ]),
]

# ================================================================
# DOCX
# ================================================================
def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def generar_docx(ruta):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)

    for bloque in C:
        tipo = bloque[0]
        if tipo == "h1":
            par = doc.add_heading(bloque[1], level=1)
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
                run.font.size = Pt(16)
        elif tipo == "h2":
            par = doc.add_heading(bloque[1], level=2)
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
                run.font.size = Pt(12)
        elif tipo == "p":
            doc.add_paragraph(bloque[1])
        elif tipo == "bullet":
            for item in bloque[1]:
                doc.add_paragraph(item, style='List Bullet')
        elif tipo == "table":
            headers, rows = bloque[1], bloque[2]
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = 'Table Grid'
            for i, h in enumerate(headers):
                cell = t.rows[0].cells[i]
                cell.text = h
                set_cell_bg(cell, AZUL)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9)
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
                    for r in cells[i].paragraphs[0].runs:
                        r.font.size = Pt(9)
        elif tipo == "code":
            par = doc.add_paragraph()
            run = par.add_run(bloque[1])
            run.font.name = 'Consolas'
            run.font.size = Pt(8)
            pPr = par._p.get_or_add_pPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'F2F2F2')
            pPr.append(shd)

    doc.save(ruta)


# ================================================================
# PDF
# ================================================================
def generar_pdf(ruta):
    doc = SimpleDocTemplate(ruta, pagesize=letter,
                            leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    st_h1 = ParagraphStyle('h1x', parent=styles['Heading1'],
                           textColor=HexColor('#1F4E78'), fontSize=15, spaceAfter=8)
    st_h2 = ParagraphStyle('h2x', parent=styles['Heading2'],
                           textColor=HexColor('#1F4E78'), fontSize=12, spaceBefore=8, spaceAfter=4)
    st_p = ParagraphStyle('px', parent=styles['Normal'], fontSize=9, leading=12)
    st_code = ParagraphStyle('cx', parent=styles['Code'], fontName='Courier',
                             fontSize=6.5, leading=8, backColor=HexColor('#F2F2F2'),
                             borderPadding=4, leftIndent=6)

    story = []
    for bloque in C:
        tipo = bloque[0]
        if tipo == "h1":
            story.append(Paragraph(bloque[1], st_h1))
        elif tipo == "h2":
            story.append(Paragraph(bloque[1], st_h2))
        elif tipo == "p":
            story.append(Paragraph(bloque[1], st_p))
            story.append(Spacer(1, 4))
        elif tipo == "bullet":
            for item in bloque[1]:
                story.append(Paragraph(f"&bull; {item}", st_p))
            story.append(Spacer(1, 4))
        elif tipo == "table":
            headers, rows = bloque[1], bloque[2]
            data = [[Paragraph(f"<b>{h}</b>", st_p) for h in headers]]
            for row in rows:
                data.append([Paragraph(str(v), st_p) for v in row])
            ncols = len(headers)
            col_w = (7.0 * inch) / ncols
            t = Table(data, colWidths=[col_w] * ncols)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), HexColor('#1F4E78')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]))
            story.append(t)
            story.append(Spacer(1, 6))
        elif tipo == "code":
            story.append(Preformatted(bloque[1], st_code))
            story.append(Spacer(1, 6))

    doc.build(story)


docx_path = os.path.join(OUT, "MANUAL_CARGA_MASIVA_USUARIOS.docx")
pdf_path = os.path.join(OUT, "MANUAL_CARGA_MASIVA_USUARIOS.pdf")
generar_docx(docx_path)
generar_pdf(pdf_path)
print(f"DOCX: {docx_path}")
print(f"PDF : {pdf_path}")
