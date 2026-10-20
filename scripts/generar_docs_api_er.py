# -*- coding: utf-8 -*-
"""
Genera 3 documentos Word adicionales:
  1. DOCUMENTACION_ENDPOINTS_API.docx   - endpoints por flujo, con SP y JSON de entrada
  2. DOCUMENTACION_SPS_JSON.docx        - catalogo de SPs y estructura del JSON @json
  3. DOCUMENTACION_MODELO_ER_COMPLETO.docx - DER detallado por modulo (todas las columnas)
Entradas: doc_output/api_map.json, doc_output/schema.json
"""
import os, re, json, math
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "doc_output")
DOCS = os.path.join(HERE, "..", "docs")
FIG = os.path.join(OUT, "figs_er")
os.makedirs(FIG, exist_ok=True)

def _find(name):
    for base in (os.path.join(HERE, "doc_output"), os.path.join(HERE, "..", "doc_output")):
        p = os.path.join(base, name)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(name)

API = json.load(open(_find("api_map.json"), encoding="utf-8"))
SCHEMA = json.load(open(_find("schema.json"), encoding="utf-8"))

UC_MAP, SP_MAP, SP_INFO = API["uc_map"], API["sp_map"], API["sp_info"]

# ---------- resolucion endpoint -> SPs ----------
def sps_de_endpoint(e):
    sps = list(e.get("sp_directos", []))
    for call in e["usecase_calls"]:
        repo = UC_MAP.get(call, call)
        s = SP_MAP.get(repo) or SP_MAP.get(call)
        if s and s not in sps:
            sps.append(s)
    return sps

# ---------- agrupacion de controladores por flujo ----------
FLUJOS = [
    ("Requerimientos", ["RequerimientosController", "LogisticaController"]),
    ("Aprobaciones", ["AprobacionesPorAreaController", "AprobacionesController",
                      "AprobacionSCController", "AprobacionOrdenController"]),
    ("Consolidacion de requerimientos", ["ConsolidacionController"]),
    ("Compras: cotizacion, solicitud y orden de compra",
     ["CotizacionController", "SolicitudCompraController", "OrdenCompraController",
      "AprobacionOCController", "GestionOCOSController", "ProveedorEmailController"]),
    ("Servicios: solicitud, cotizacion y orden de servicio",
     ["SolicitudServicioController", "OrdenServicioController",
      "AprobacionOSController", "SeguimientoOSController"]),
    ("Recepcion y conformidad", ["RecepcionOCController", "ConformidadAlmacenController"]),
    ("Almacen, inventario, despachos y devoluciones",
     ["KardexController", "GestionInventarioController", "ListasStockController",
      "ReingresoController", "DevolucionConsumoController", "DevolucionProveedorController"]),
    ("Dashboards y reportes",
     ["DashboardLogisticaController", "DashboardComprasController",
      "DashboardJlologistController", "ReportesComprasController"]),
    ("Administracion y configuracion",
     ["AdminLogisticaController", "ConfiguracionPermisoController", "ConfigMenuController",
      "ConfigSmtpController", "BrandingController", "PlantillaCorreoController",
      "PlantillaPdfController", "UnidadMedidaController", "TipoCambioController",
      "AppDeploymentsController"]),
]
CTRL2FLUJO = {}
for flujo, ctrls in FLUJOS:
    for c in ctrls:
        CTRL2FLUJO[c] = flujo

# ---------- helpers docx ----------
def estilos(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10)
    for i, (sz, col) in enumerate([(20, "1F4E79"), (14, "1F4E79"), (12, "2E74B5"), (11, "404040")], 1):
        h = doc.styles[f"Heading {i}"]
        h.font.size = Pt(sz); h.font.color.rgb = RGBColor.from_string(col); h.font.bold = True

def tabla(doc, headers, rows, widths=None, size=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255,255,255)
        c._tc.get_or_add_tcPr().append(_shd("1F4E79"))
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
    return t

def _shd(color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), color)
    return el

def portada(doc, titulo, subtitulo):
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0x1F,0x4E,0x79)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitulo); r.font.size = Pt(16); r.font.color.rgb = RGBColor(0x40,0x40,0x40)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sistema de Logistica - Aplicacion logistica (Angular) / api_logistica (.NET) / BD SQL Server LOGISTICA")
    r.font.size = Pt(10); r.italic = True
    doc.add_page_break()

def campos_json_str(sp):
    info = SP_INFO.get(sp)
    if not info:
        return ""
    campos = sorted(set(list(info["json_fields"].keys()) + info["openjson"]))
    txt = ", ".join(campos)
    return txt if len(txt) <= 350 else txt[:347] + "..."

# =====================================================================
# DOC 1: ENDPOINTS POR FLUJO
# =====================================================================
doc = Document(); estilos(doc)
portada(doc, "Documentacion de Endpoints API", "Catalogo de endpoints REST por flujo funcional")

doc.add_heading("1. Arquitectura de la capa API", 1)
doc.add_paragraph(
    "La API api_logistica (ASP.NET) expone endpoints REST bajo la ruta base /api. "
    "El patron general de implementacion es: el endpoint recibe el cuerpo de la peticion "
    "como JsonElement ([FromBody]), lo serializa a texto JSON y lo envia a un Stored "
    "Procedure de SQL Server como parametro @json NVARCHAR(MAX). El SP lee los campos con "
    "JSON_VALUE(@json,'$.campo') u OPENJSON y devuelve el resultado serializado con FOR JSON PATH.")
doc.add_paragraph(
    "Cadena de invocacion: Controller -> UseCase -> Repository -> Stored Procedure (@json). "
    "En los controladores de administracion el SP se invoca directamente con SqlCommand.")

doc.add_heading("2. Resumen de controladores", 1)
rows = []
for c in API["controllers"]:
    flujo = CTRL2FLUJO.get(c["controlador"], "Otros")
    rows.append([c["controlador"], "/" + c["base"], len(c["endpoints"]), flujo])
tabla(doc, ["Controlador", "Ruta base", "N° endpoints", "Flujo"], rows)

doc.add_heading("3. Endpoints por flujo", 1)
for flujo, ctrls in FLUJOS:
    doc.add_heading(flujo, 2)
    for c in API["controllers"]:
        if c["controlador"] not in ctrls or not c["endpoints"]:
            continue
        doc.add_heading(c["controlador"], 3)
        rows = []
        for e in c["endpoints"]:
            sps = sps_de_endpoint(e)
            rows.append([e["verbo"], e["ruta"], e["metodo"],
                         "\n".join(sps) if sps else "(consulta directa / EF)"])
        tabla(doc, ["Verbo", "Ruta", "Metodo C#", "Stored Procedure(s)"], rows)
        # detalle JSON por endpoint
        for e in c["endpoints"]:
            sps = sps_de_endpoint(e)
            if not sps:
                continue
            campos = campos_json_str(sps[0])
            p = doc.add_paragraph()
            r = p.add_run(f"{e['verbo']} {e['ruta']}")
            r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x2E,0x74,0xB5)
            p2 = doc.add_paragraph()
            r2 = p2.add_run("JSON de entrada: " + (campos if campos else "{} (sin campos requeridos)"))
            r2.font.size = Pt(8); r2.font.name = "Consolas"
            p2.paragraph_format.left_indent = Inches(0.25)

doc.add_heading("4. Servicios Angular que consumen la API", 1)
doc.add_paragraph(
    "Relacion de URLs invocadas desde los servicios del frontend (src/app/services). "
    "Todas usan environment.baseUrl como prefijo.")
por_url = {}
for s in API["servicios"]:
    por_url.setdefault(s["url"], []).append(s["servicio"])
rows = [[u, ", ".join(sorted(set(sv)))] for u, sv in sorted(por_url.items())]
tabla(doc, ["URL invocada", "Servicio Angular"], rows, size=7)

doc.save(os.path.join(DOCS, "DOCUMENTACION_ENDPOINTS_API.docx"))
print("OK endpoints")

# =====================================================================
# DOC 2: SPs Y ESTRUCTURA JSON
# =====================================================================
doc = Document(); estilos(doc)
portada(doc, "Documentacion de Stored Procedures", "Estructura del parametro @json por procedimiento")

doc.add_heading("1. Patron de parametros JSON", 1)
doc.add_paragraph(
    "Todos los SPs del sistema reciben un unico parametro @json NVARCHAR(MAX) con el cuerpo "
    "de la peticion. Dentro del SP los campos se extraen con JSON_VALUE(@json,'$.campo') y los "
    "arreglos con OPENJSON(@json,'$.items'). La salida se serializa con FOR JSON PATH. "
    "A continuacion se documenta, para cada SP, los campos JSON que consume, con el tipo SQL "
    "en el que se convierten, y los endpoints que lo invocan.")

# endpoints -> sp (invertir)
sp2ep = {}
for c in API["controllers"]:
    for e in c["endpoints"]:
        for sp in sps_de_endpoint(e):
            sp2ep.setdefault(sp, []).append(f"{e['verbo']} {e['ruta']}")

sps_referenciados = set(sp2ep)

doc.add_heading("2. Stored Procedures invocados por la API", 1)
for flujo, ctrls in FLUJOS:
    sps_flujo = []
    for c in API["controllers"]:
        if c["controlador"] in ctrls:
            for e in c["endpoints"]:
                for sp in sps_de_endpoint(e):
                    if sp not in sps_flujo:
                        sps_flujo.append(sp)
    if not sps_flujo:
        continue
    doc.add_heading(flujo, 2)
    for sp in sps_flujo:
        info = SP_INFO.get(sp)
        h = doc.add_heading(sp, 3)
        meta = doc.add_paragraph()
        eps = sp2ep.get(sp, [])
        meta.add_run("Endpoints: ").bold = True
        meta.add_run("; ".join(eps[:4]) + ("..." if len(eps) > 4 else "")).font.size = Pt(8)
        if info:
            srcs = sorted(set(info["archivos"]))
            p = doc.add_paragraph()
            p.add_run("Definido en: ").bold = True
            p.add_run(", ".join(srcs)).font.size = Pt(8)
            if info["params"]:
                p = doc.add_paragraph()
                p.add_run("Parametros: ").bold = True
                p.add_run(", ".join(info["params"])).font.size = Pt(8)
        # tabla de campos JSON
        if info and (info["json_fields"] or info["openjson"]):
            rows = [[f"$.{k}", v or "(inferido)"] for k, v in info["json_fields"].items()]
            rows += [[f"$.{k}", "array/objeto (OPENJSON)"] for k in info["openjson"]]
            tabla(doc, ["Campo JSON", "Tipo/lectura"], rows, size=8)
        else:
            doc.add_paragraph("Sin campos JSON extraidos (sin parametros de entrada).").runs[0].font.size = Pt(8)

doc.add_heading("3. Otros SPs definidos en los scripts SQL", 1)
doc.add_paragraph(
    "SPs presentes en los scripts del proyecto que no fueron referenciados directamente "
    "por los repositorios analizados (utilidades, migraciones, SPs llamados desde otros SPs "
    "o desde el ERP).")
otros = sorted(s for s in SP_INFO if s not in sps_referenciados)
rows = []
for sp in otros:
    info = SP_INFO[sp]
    nf = len(info["json_fields"]) + len(info["openjson"])
    rows.append([sp, ", ".join(sorted(set(info["archivos"]))[:2]), nf])
tabla(doc, ["Stored Procedure", "Script(s)", "N° campos JSON"], rows, size=7)

doc.save(os.path.join(DOCS, "DOCUMENTACION_SPS_JSON.docx"))
print("OK sps")

# =====================================================================
# DOC 3: MODELO ER COMPLETO DETALLADO
# =====================================================================
def modulo(t):
    n = t.lower()
    if n.startswith(("wh_", "spring_", "logistica_colamigracion",
                     "logistica_sincronizacion", "logistica_maestroitem")):
        return "Integracion ERP Spring"
    if n.startswith("logistica_consol"):
        return "Consolidacion OC/OS"
    if "servicio" in n or n.endswith(("aprobacionos", "flujoaprobacionos",
        "usuarioaprobadoros", "osconformidad", "aprobacionss")) \
        or "confirmacionservicio" in n or "conformidadservicio" in n \
        or "historialprecioservicio" in n or "plantillaconformidadservicio" in n:
        return "Servicios"
    if "consolidacion" in n or "itempendienteconsolidacion" in n \
        or "consolidado" in n or n == "reqconsolidado":
        return "Consolidacion de Requerimientos"
    if "aprobacion" in n or "flujo" in n or "aprobadores" in n \
        or "saldopendiente" in n or "adjuntosorden" in n \
        or n in ("logistica_areas", "logistica_usuariosporarea"):
        return "Aprobaciones"
    if n in ("logistica_req", "logistica_dreq") or "reqcommodity" in n \
        or "solicitudcreacionitem" in n or "historialanulaciones" in n \
        or "logcambioestadorequerimiento" in n or n == "logistica_requerimientos" \
        or n == "logistica_requerimientosporarea":
        return "Requerimientos"
    if "requerimientotransferencia" in n:
        return "Almacen e Inventario"
    if "recepcion" in n or "despacho" in n or "kardex" in n or "stock" in n \
        or "inventario" in n or "transaccion" in n or "devolucion" in n \
        or "reingreso" in n or "almacen" in n or n == "logistica_item" \
        or "conformidadalmacen" in n or "plantillaalmacen" in n:
        return "Almacen e Inventario"
    if "cotizacion" in n or "solicitudcompra" in n or "ordencompra" in n \
        or "ocadjunto" in n or "seguimientoordencompra" in n \
        or "incrementooc" in n or "proveedor" in n or n == "logistica_entrega":
        return "Compras"
    if "notificacion" in n or "correolog" in n or "alerta" in n:
        return "Notificaciones"
    return "Administracion y Configuracion"

MODULOS_ORDEN = [
    "Requerimientos", "Aprobaciones", "Consolidacion de Requerimientos",
    "Compras", "Consolidacion OC/OS", "Servicios", "Almacen e Inventario",
    "Notificaciones", "Administracion y Configuracion", "Integracion ERP Spring",
]
MODULO_TABLAS = {m: [] for m in MODULOS_ORDEN}
for t in SCHEMA:
    MODULO_TABLAS[modulo(t)].append(t)
for m in MODULO_TABLAS:
    MODULO_TABLAS[m].sort(key=lambda x: x.lower())

def figura_er_modulo(mod, tablas, fname):
    """ER detallado: caja por tabla con todas las columnas."""
    n = len(tablas)
    ncols = min(4, max(2, math.ceil(math.sqrt(n))))
    BW, ROW_H, HDR, GAPX, GAPY = 3.6, 0.24, 0.42, 0.9, 0.9
    # ordenar y calcular posiciones por filas de rejilla
    filas = [tablas[i:i+ncols] for i in range(0, n, ncols)]
    pos = {}
    y = 0
    for fila in filas:
        hmax = max(HDR + ROW_H * len(SCHEMA[t]["cols"]) for t in fila)
        x = 0
        for t in fila:
            pos[t] = (x, y); x += BW + GAPX
        y += hmax + GAPY
    W = ncols * (BW + GAPX)
    H = y + 0.5
    fig, ax = plt.subplots(figsize=(W * 0.55, H * 0.55), dpi=170)
    ax.set_xlim(-0.3, W); ax.set_ylim(-0.3, H); ax.axis("off")
    ax.invert_yaxis()

    # cajas
    for t, (x, y0) in pos.items():
        cols = SCHEMA[t]["cols"]
        h = HDR + ROW_H * len(cols)
        ax.add_patch(FancyBboxPatch((x, y0), BW, h, boxstyle="round,pad=0.03",
                                    fc="#FDFEFE", ec="#2E74B5", lw=1.2))
        ax.add_patch(plt.Rectangle((x, y0), BW, HDR, fc="#1F4E79", ec="#1F4E79"))
        ax.text(x + BW/2, y0 + HDR/2, t, ha="center", va="center",
                fontsize=6.2, fontweight="bold", color="white")
        for i, c in enumerate(cols):
            yy = y0 + HDR + i * ROW_H + ROW_H/2
            pk = "PK" in (c.get("flags") or "")
            fk = any(c["name"].lower() in (f["cols"] or "").lower() for f in SCHEMA[t].get("fks", []))
            marca = ("PK " if pk else "") + ("FK " if fk else "")
            ax.text(x + 0.08, yy, f"{marca}{c['name']}", fontsize=4.6, va="center",
                    fontweight="bold" if pk else "normal",
                    color="#C00000" if pk else ("#B45F06" if fk else "#222222"))
            ax.text(x + BW - 0.08, yy, c["type"], fontsize=4.2, va="center",
                    ha="right", color="#777777")

    # flechas FK dentro del modulo
    cj = {"connectionstyle": "arc3,rad=0.15"}
    for t, (x, y0) in pos.items():
        for fk in SCHEMA[t].get("fks", []):
            ref = fk["ref"]
            if ref in pos:
                x2, y2 = pos[ref]
                colsh = len(SCHEMA[ref]["cols"])
                ax.annotate("",
                    xy=(x2 + BW/2, y2 - 0.05),
                    xytext=(x + BW/2, y0 + HDR + ROW_H * len(SCHEMA[t]["cols"]) + 0.05),
                    arrowprops=dict(arrowstyle="->", color="#B45F06", lw=0.7,
                                    alpha=0.75, connectionstyle="arc3,rad=0.2"))

    fig.savefig(fname, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return fname

doc = Document(); estilos(doc)
portada(doc, "Modelo Entidad-Relacion Completo", "Diagramas ER detallados por modulo (todas las columnas)")

doc.add_heading("1. Descripcion general", 1)
doc.add_paragraph(
    f"La base de datos LOGISTICA (SQL Server) contiene {len(SCHEMA)} tablas agrupadas en "
    f"{len(MODULOS_ORDEN)} modulos funcionales. Cada figura muestra las tablas del modulo con "
    "todas sus columnas (nombre y tipo de dato), marcando PK en rojo y FK en naranja. "
    "Las flechas naranjas indican llaves foraneas hacia otra tabla del mismo modulo; las "
    "relaciones con tablas de otros modulos se listan debajo de cada figura.")

nfig = 0
for i, mod in enumerate(MODULOS_ORDEN, 2):
    tablas = MODULO_TABLAS[mod]
    if not tablas:
        continue
    doc.add_heading(f"{i}. Modulo: {mod}", 1)
    doc.add_paragraph(f"{len(tablas)} tablas.")
    fname = os.path.join(FIG, f"er_{i}_{mod.split()[0].lower()}.png")
    figura_er_modulo(mod, tablas, fname)
    doc.add_picture(fname, width=Inches(7.0))
    doc.add_paragraph(f"Figura {i-1}. DER detallado del modulo {mod}.").runs[0].font.size = Pt(8)
    nfig += 1
    # relaciones internas y externas
    internas, externas = [], []
    conj = set(tablas)
    for t in tablas:
        for fk in SCHEMA[t].get("fks", []):
            row = [t, fk["cols"], fk["ref"], fk.get("refcols", "")]
            (internas if fk["ref"] in conj else externas).append(row)
    if internas:
        doc.add_paragraph("Relaciones dentro del modulo:").runs[0].bold = True
        tabla(doc, ["Tabla", "Columna FK", "Referencia", "Columna ref."], internas, size=7)
    if externas:
        doc.add_paragraph("Relaciones hacia otros modulos:").runs[0].bold = True
        tabla(doc, ["Tabla", "Columna FK", "Referencia", "Columna ref."], externas, size=7)
    doc.add_page_break()

doc.save(os.path.join(DOCS, "DOCUMENTACION_MODELO_ER_COMPLETO.docx"))
print("OK er", nfig, "figuras")
print("=== FIN ===")
