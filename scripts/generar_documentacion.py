# -*- coding: utf-8 -*-
"""
Genera documentacion Word del sistema Logistica:
 - Diagrama Entidad-Relacion de la base de datos
 - Diagrama de casos de uso con actores
 - Descripcion de flujos
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Ellipse, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, "doc_output")
os.makedirs(OUT_DIR, exist_ok=True)

# =====================================================================
# HELPERS GRAFICOS
# =====================================================================

def draw_entity(ax, x, y, w, title, cols, color="#dbe9f6", edge="#2e5f8a", title_fs=8, col_fs=6.5):
    """Dibuja una tabla tipo ER: titulo + columnas."""
    row_h = 0.30
    title_h = 0.34
    h = title_h + row_h * len(cols)
    # cuerpo
    ax.add_patch(FancyBboxPatch((x, y - h), w, h,
                 boxstyle="round,pad=0.02,rounding_size=0.05",
                 fc="white", ec=edge, lw=1.0, zorder=2))
    # cabecera
    ax.add_patch(FancyBboxPatch((x, y - title_h), w, title_h,
                 boxstyle="round,pad=0.02,rounding_size=0.05",
                 fc=color, ec=edge, lw=1.0, zorder=3))
    ax.text(x + w / 2, y - title_h / 2, title, ha="center", va="center",
            fontsize=title_fs, fontweight="bold", zorder=4)
    for i, c in enumerate(cols):
        cy = y - title_h - row_h * (i + 0.5)
        ax.text(x + 0.12, cy, c, ha="left", va="center",
                fontsize=col_fs, family="monospace", zorder=4)
    return {"x": x, "y": y, "w": w, "h": h, "cx": x + w / 2, "cy": y - h / 2}


def arrow(ax, p1, p2, label="", rad=0.0, color="#555555"):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>",
                 mutation_scale=9, lw=1.0, color=color,
                 connectionstyle=f"arc3,rad={rad}", zorder=1))
    if label:
        mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
        ax.text(mx, my + 0.12, label, fontsize=6, ha="center", color=color,
                style="italic", zorder=5)


def right(b):  return (b["x"] + b["w"], b["cy"])
def left(b):   return (b["x"], b["cy"])
def top(b):    return (b["cx"], b["y"])
def bottom(b): return (b["cx"], b["y"] - b["h"])


# =====================================================================
# FIGURA 1: DER - FLUJO PRINCIPAL COMPRAS + ALMACEN
# =====================================================================
fig, ax = plt.subplots(figsize=(24, 16))
ax.set_xlim(0, 25.3); ax.set_ylim(0, 16); ax.axis("off")
ax.set_title("Diagrama Entidad-Relacion - Flujo principal: Requerimiento -> Compra -> Almacen",
             fontsize=13, fontweight="bold", pad=12)

e = {}
# --- Columna 1: Requerimientos ---
e["req"] = draw_entity(ax, 0.3, 15.4, 3.4, "logistica_req", [
    "PK id / idrequerimiento", "RequisicionNumero", "ruc (empresa)", "idarea / idfundo",
    "idalmacen / destino", "tipo / itemtipo", "estado / estados", "dniregistra",
    "fecharegistro", "fechaAprobacion", "glosa / prioridad"], "#f4d9d0")
e["dreq"] = draw_entity(ax, 0.3, 11.4, 3.4, "logistica_dreq", [
    "PK id", "FK idrequerimiento", "codigo / idproducto", "iddescripcion",
    "tipoclasificacion", "cantidad", "idcentrocosto", "idproyecto / idlabor"], "#f4d9d0")
e["reqadj"] = draw_entity(ax, 0.3, 8.2, 3.4, "LOGISTICA_RequerimientoAdjuntos", [
    "PK id", "FK idrequerimiento", "nombreArchivo", "contenidoB64"], "#f4d9d0")
e["reqcm"] = draw_entity(ax, 0.3, 6.0, 3.4, "LOGISTICA_ReqCommodity", [
    "PK idRequerimiento", "tipo / idarea", "estado", "usuarioRegistra"], "#f4d9d0")
e["dreqcm"] = draw_entity(ax, 0.3, 3.9, 3.4, "LOGISTICA_DReqCommodity", [
    "PK idDetalle", "FK idRequerimiento", "commodity / cantidad", "estado"], "#f4d9d0")

# --- Columna 2: Aprobaciones ---
e["areas"] = draw_entity(ax, 4.6, 15.4, 3.3, "LOGISTICA_Areas", [
    "PK idarea", "nombreArea", "activo"], "#fde9c8")
e["upxarea"] = draw_entity(ax, 4.6, 13.6, 3.3, "LOGISTICA_UsuariosPorArea", [
    "PK id", "FK idarea", "documentoidentidad", "nombre", "esJefeArea", "activo"], "#fde9c8")
e["flujoarea"] = draw_entity(ax, 4.6, 10.9, 3.3, "LOGISTICA_FlujoAprobacionPorArea", [
    "PK id", "FK idarea", "nivelOrden", "tipoAprobador", "activo"], "#fde9c8")
e["logarea"] = draw_entity(ax, 4.6, 8.6, 3.3, "LOGISTICA_LogAprobacionesPorArea", [
    "PK id", "FK idRequerimiento", "FK idarea", "nivelAprobacion",
    "estado", "jefeAreaAsignado", "fechaAprobacion"], "#fde9c8")
e["saldo"] = draw_entity(ax, 4.6, 5.9, 3.3, "LOGISTICA_SaldoPendienteAprobacion", [
    "PK idSolicitud", "FK idrequerimiento", "usuarioSolicita", "estado",
    "usuarioResuelve"], "#fde9c8")
e["saldodet"] = draw_entity(ax, 4.6, 3.5, 3.3, "LOGISTICA_SaldoPendienteAprobacionDet", [
    "PK id", "FK idSolicitud", "codigo", "cantidadSolicitada",
    "cantidadDespachada", "saldoPendiente"], "#fde9c8")

# --- Columna 3: Consolidacion ---
e["pend"] = draw_entity(ax, 8.9, 15.4, 3.5, "LOGISTICA_ItemPendienteConsolidacion", [
    "PK id", "FK idDetalle / idrequerimiento", "tipoRequerimiento", "codigoItem",
    "cantidad", "estadoConsolidacion", "FK idConsolidacion"], "#d9ead3")
e["concab"] = draw_entity(ax, 8.9, 12.2, 3.5, "LOGISTICA_ConsolidacionCab", [
    "PK idConsolidacion", "codigo (CONS-...)", "estado / estadoProceso",
    "usuarioCreador", "totalItems"], "#d9ead3")
e["condet"] = draw_entity(ax, 8.9, 9.7, 3.5, "LOGISTICA_ConsolidacionDet", [
    "PK idDetConsolidacion", "FK idConsolidacion", "codigoItem",
    "cantidadTotal", "unidadMedida", "estado"], "#d9ead3")
e["conorg"] = draw_entity(ax, 8.9, 7.2, 3.5, "LOGISTICA_ConsolidacionOrigen", [
    "PK idOrigen", "FK idDetConsolidacion", "FK idConsolidacion",
    "tipoOrigen (COMPRA/CONSUMO)", "requerimientoOrigen", "cantidadConsolidada"], "#d9ead3")

# --- Columna 4: Cotizaciones y Compras ---
e["solcot"] = draw_entity(ax, 13.4, 15.4, 3.5, "LOGISTICA_SolicitudCotizacion", [
    "PK idSolicitudCotizacion", "noSolicitud", "FK idConsolidacion",
    "estado", "usuarioGenera"], "#d0e0f0")
e["solcotdet"] = draw_entity(ax, 13.4, 13.3, 3.5, "LOGISTICA_SolicitudCotizacionDet", [
    "PK idDetSolicitud", "FK idSolicitudCotizacion", "codigoItem", "cantidad"], "#d0e0f0")
e["cot"] = draw_entity(ax, 13.4, 11.4, 3.5, "LOGISTICA_Cotizacion", [
    "PK idCotizacion", "FK idSolicitudCotizacion", "proveedor / ruc",
    "moneda / total", "estado", "seleccionada"], "#d0e0f0")
e["cotdet"] = draw_entity(ax, 13.4, 9.4, 3.5, "LOGISTICA_CotizacionDet", [
    "PK id", "FK idCotizacion", "codigoItem", "precioUnitario"], "#d0e0f0")
e["prov"] = draw_entity(ax, 13.4, 7.5, 3.5, "LOGISTICA_Proveedor", [
    "PK id / codigo", "ruc", "razonSocial", "email", "FK idCategoria"], "#d0e0f0")
e["sc"] = draw_entity(ax, 13.4, 5.4, 3.5, "LOGISTICA_SolicitudCompra", [
    "PK idSolicitud", "numeroSolicitud", "tipo", "estado",
    "usuarioSolicita / usuarioAprueba"], "#d0e0f0")
e["scdet"] = draw_entity(ax, 13.4, 3.2, 3.5, "LOGISTICA_SolicitudCompraDet", [
    "PK idDetalle", "FK idSolicitud", "codigo", "cantidad / aprobada",
    "estado"], "#d0e0f0")
e["scadj"] = draw_entity(ax, 13.4, 1.3, 3.5, "LOGISTICA_SolicitudCompraAdjuntos", [
    "PK id", "FK idSolicitud", "nombreArchivo", "contenidoB64"], "#d0e0f0")

# --- Columna 5: Orden de Compra y aprobacion ---
e["oc"] = draw_entity(ax, 17.9, 15.4, 3.6, "LOGISTICA_OrdenCompra", [
    "PK idOrden", "numeroOrden", "FK idSolicitud", "proveedor / rucProveedor",
    "moneda / totalOrden", "estado", "usuarioAprueba"], "#e8d9f0")
e["ocdet"] = draw_entity(ax, 17.9, 12.6, 3.6, "LOGISTICA_OrdenCompraDet", [
    "PK idDetalle", "FK idOrden", "codigo", "cantidad / recibida",
    "precioUnitario / total", "estado"], "#e8d9f0")
e["oclog"] = draw_entity(ax, 17.9, 10.4, 3.6, "LOGISTICA_OrdenCompraLog", [
    "PK idLog", "FK idOrden", "accion", "estadoAnterior/Nuevo", "usuarioAccion"], "#e8d9f0")
e["flujooc"] = draw_entity(ax, 17.9, 8.5, 3.6, "LOGISTICA_FlujoAprobacionOC", [
    "PK id", "nivelOrden", "montoMin / montoMax", "rolAprobador", "activo"], "#e8d9f0")
e["aproc"] = draw_entity(ax, 17.9, 6.2, 3.6, "LOGISTICA_AprobacionOC", [
    "PK id", "FK idOrden", "nivelAprobacion", "aprobador",
    "estado", "fechaAprobacion"], "#e8d9f0")
e["aprocu"] = draw_entity(ax, 17.9, 4.2, 3.6, "LOGISTICA_UsuariosAprobadoresOC", [
    "PK id", "documentoidentidad", "rol / nivel", "montoMaximo", "activo"], "#e8d9f0")
e["segoc"] = draw_entity(ax, 17.9, 2.1, 3.6, "LOGISTICA_SeguimientoOrdenCompra", [
    "PK id", "FK idOrden", "etapa / estado", "fecha", "usuario"], "#e8d9f0")

# --- Columna 6: Recepcion / Kardex / Despacho ---
e["recep"] = draw_entity(ax, 21.6 + 0.0, 15.4, 3.4, "LOGISTICA_RecepcionOC", [
    "PK idRecepcion", "FK idOrden / numeroOrden", "fechaRecepcion",
    "usuarioRecibe", "estado", "numeroNI"], "#cfe2f3")
e["recepdet"] = draw_entity(ax, 21.6, 13.0, 3.4, "LOGISTICA_RecepcionOCDet", [
    "PK id", "FK idRecepcion", "codigo", "cantidadRecibida", "cantidadAceptada"], "#cfe2f3")
e["trans"] = draw_entity(ax, 21.6, 10.8, 3.4, "LOGISTICA_TransaccionInventario", [
    "PK idTransaccion", "tipoTransaccion", "FK idAlmacen", "documentoOrigen",
    "usuario", "fecha"], "#cfe2f3")
e["transdet"] = draw_entity(ax, 21.6, 8.6, 3.4, "LOGISTICA_DetalleTransaccionInv.", [
    "PK id", "FK idTransaccion", "FK idItem", "cantidad", "costoUnitario"], "#cfe2f3")
e["kardex"] = draw_entity(ax, 21.6, 6.5, 3.4, "LOGISTICA_Kardex", [
    "PK id", "FK idItem / idAlmacen", "tipoMovimiento", "cantidad",
    "saldoCantidad / saldoValor", "documentoReferencia"], "#cfe2f3")
e["stock"] = draw_entity(ax, 21.6, 4.3, 3.4, "LOGISTICA_Stock", [
    "PK id", "FK idItem / idAlmacen", "stockActual", "stockDisponible",
    "costoPromedio"], "#cfe2f3")
e["desp"] = draw_entity(ax, 21.6, 2.3, 3.4, "LOGISTICA_Despachos / Detalle", [
    "PK idDespacho", "FK idrequerimiento", "cantidadDespachada",
    "usuarioDespacha", "estado"], "#cfe2f3")
e["item"] = draw_entity(ax, 21.6, 0.6, 3.4, "LOGISTICA_Item / Almacen", [
    "PK idItem / idAlmacen", "codigo / descripcion", "unidadMedida", "activo"], "#cfe2f3")

# Flechas principales
arrow(ax, right(e["req"]), left(e["upxarea"]), "")
arrow(ax, bottom(e["req"]), top(e["dreq"]), "1:N")
arrow(ax, bottom(e["dreq"]), top(e["reqadj"]), "1:N")
arrow(ax, bottom(e["reqcm"]), top(e["dreqcm"]), "1:N")
arrow(ax, right(e["dreq"]), left(e["pend"]), "items aprobados", 0.1)
arrow(ax, bottom(e["pend"]), top(e["concab"]), "N:1")
arrow(ax, bottom(e["concab"]), top(e["condet"]), "1:N")
arrow(ax, bottom(e["condet"]), top(e["conorg"]), "1:N")
arrow(ax, right(e["concab"]), left(e["solcot"]), "", 0.1)
arrow(ax, bottom(e["solcot"]), top(e["solcotdet"]), "1:N")
arrow(ax, bottom(e["solcotdet"]), top(e["cot"]), "genera")
arrow(ax, bottom(e["cot"]), top(e["cotdet"]), "1:N")
arrow(ax, left(e["cot"]), right(e["prov"]), "", 0.0)
arrow(ax, right(e["sc"]), left(e["oc"]), "genera", -0.15)
arrow(ax, bottom(e["sc"]), top(e["scdet"]), "1:N")
arrow(ax, bottom(e["scdet"]), top(e["scadj"]), "1:N")
arrow(ax, bottom(e["oc"]), top(e["ocdet"]), "1:N")
arrow(ax, bottom(e["ocdet"]), top(e["oclog"]), "auditoria")
arrow(ax, bottom(e["flujooc"]), top(e["aproc"]), "niveles")
arrow(ax, bottom(e["aproc"]), top(e["aprocu"]), "asignado a")
arrow(ax, right(e["oc"]), left(e["recep"]), "OC aprobada", 0.15)
arrow(ax, bottom(e["recep"]), top(e["recepdet"]), "1:N")
arrow(ax, bottom(e["recepdet"]), top(e["trans"]), "ingreso")
arrow(ax, bottom(e["trans"]), top(e["transdet"]), "1:N")
arrow(ax, bottom(e["transdet"]), top(e["kardex"]), "movimiento")
arrow(ax, bottom(e["kardex"]), top(e["stock"]), "actualiza")
arrow(ax, bottom(e["stock"]), top(e["desp"]), "salida")
arrow(ax, left(e["desp"]), right(e["saldo"]), "saldo pendiente", 0.15)
arrow(ax, bottom(e["saldo"]), top(e["saldodet"]), "1:N")

plt.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "der_principal.png"), dpi=140, bbox_inches="tight")
plt.close(fig)
print("OK der_principal.png")

# =====================================================================
# FIGURA 2: DER - SERVICIOS, CONSOLIDACION OC/OS Y ADMINISTRACION
# =====================================================================
fig, ax = plt.subplots(figsize=(24, 15))
ax.set_xlim(0, 25.3); ax.set_ylim(0, 15); ax.axis("off")
ax.set_title("Diagrama Entidad-Relacion - Servicios, Consolidacion OC/OS, Administracion e Integracion",
             fontsize=13, fontweight="bold", pad=12)

s = {}
# --- Bloque Servicios (izquierda) ---
s["ss"] = draw_entity(ax, 0.3, 14.4, 3.6, "LOGISTICA_SolicitudServicio", [
    "PK idSolicitud", "numeroSolicitud", "tipoServicio", "FK idarea",
    "estado", "usuarioSolicita"], "#f4d9d0")
s["ssdet"] = draw_entity(ax, 0.3, 12.4, 3.6, "LOGISTICA_SolicitudServicioDet", [
    "PK id", "FK idSolicitud", "descripcionServicio", "cantidad"], "#f4d9d0")
s["ssadj"] = draw_entity(ax, 0.3, 10.6, 3.6, "LOGISTICA_SolicitudServicioAdjuntos", [
    "PK id", "FK idSolicitud", "nombreArchivo", "contenidoB64"], "#f4d9d0")
s["aprss"] = draw_entity(ax, 0.3, 8.6, 3.6, "LOGISTICA_AprobacionSS", [
    "PK id", "FK idSolicitud", "nivelAprobacion", "aprobador", "estado"], "#f4d9d0")
s["cotsrv"] = draw_entity(ax, 4.8, 14.4, 3.6, "LOGISTICA_CotizacionServicio", [
    "PK idCotizacion", "FK idSolicitud", "proveedor", "montoTotal", "estado"], "#fde9c8")
s["cotsrvdet"] = draw_entity(ax, 4.8, 12.2, 3.6, "LOGISTICA_CotizacionServicioDet", [
    "PK id", "FK idCotizacion", "descripcion", "precioUnitario"], "#fde9c8")
s["os"] = draw_entity(ax, 4.8, 10.0, 3.6, "LOGISTICA_OrdenServicio", [
    "PK idOS", "numeroOrden", "FK idSolicitud / idCotizacion",
    "proveedor", "montoTotal", "estado"], "#fde9c8")
s["osdet"] = draw_entity(ax, 4.8, 7.6, 3.6, "LOGISTICA_OrdenServicioDet", [
    "PK id", "FK idOS", "descripcionServicio", "cantidad", "precioUnitario"], "#fde9c8")
s["apros"] = draw_entity(ax, 4.8, 5.5, 3.6, "LOGISTICA_AprobacionOS / Flujo", [
    "PK id", "FK idOS", "nivelAprobacion", "aprobador",
    "estado", "LOGISTICA_FlujoAprobacionOS"], "#fde9c8")
s["conf"] = draw_entity(ax, 4.8, 3.2, 3.6, "LOGISTICA_ConformidadServicio", [
    "PK id", "FK idOS", "estado CONFORME/NO_CONFORME",
    "firmaBase64", "fechaFirma"], "#fde9c8")
s["segos"] = draw_entity(ax, 4.8, 1.1, 3.6, "LOGISTICA_SeguimientoOrdenServicio", [
    "PK id", "FK idOS", "etapa / estado", "fecha"], "#fde9c8")

# --- Bloque Consolidacion OC/OS (centro) ---
s["consoloc"] = draw_entity(ax, 9.4, 14.4, 3.6, "LOGISTICA_ConsolOC", [
    "PK idOrden", "numeroOrden", "FK idConsolidacion", "proveedor",
    "moneda / totalOrden", "estado", "pdfGenerado / correoEnviado"], "#d9ead3")
s["consolocdet"] = draw_entity(ax, 9.4, 12.0, 3.6, "LOGISTICA_ConsolOCDet", [
    "PK id", "FK idOrden", "codigo", "cantidad / recibida",
    "precioUnitario / total"], "#d9ead3")
s["consoldist"] = draw_entity(ax, 9.4, 9.9, 3.6, "LOGISTICA_ConsolOCDist", [
    "PK id", "FK idOrden / idDetalle", "distribucion por requerimiento",
    "cantidadAsignada"], "#d9ead3")
s["consolos"] = draw_entity(ax, 9.4, 7.5, 3.6, "LOGISTICA_ConsolOS", [
    "PK idOS", "numeroOrden", "FK idConsolidacion", "proveedor",
    "fechaInicio / fechaFin", "montoTotal", "estado"], "#d9ead3")
s["consolosdet"] = draw_entity(ax, 9.4, 5.2, 3.6, "LOGISTICA_ConsolOSDet", [
    "PK id", "FK idOS", "descripcionServicio", "cantidad", "precioUnitario"], "#d9ead3")
s["consolconf"] = draw_entity(ax, 9.4, 3.1, 3.6, "LOGISTICA_ConsolOSConformidad", [
    "PK id", "FK idOS", "estado", "firmaBase64", "actaBase64"], "#d9ead3")
s["consoladj"] = draw_entity(ax, 9.4, 1.2, 3.6, "LOGISTICA_ConsolAdjunto", [
    "PK id", "FK idOrden", "tipoOrden OC/OS", "nombreArchivo", "contenidoB64"], "#d9ead3")

# --- Bloque Integracion ERP Spring (centro-derecha) ---
s["whos"] = draw_entity(ax, 14.1, 14.4, 3.7, "WH_OrdenServicio (+Detalle)", [
    "Tabla espejo ERP Spring", "numeroOrden", "proveedor", "estadoSync"], "#e8d9f0")
s["whoc"] = draw_entity(ax, 14.1, 12.4, 3.7, "WH_OrdenCompra (ERP)", [
    "Tabla destino ERP Spring", "OC aprobada sincronizada",
    "via SP sincronizarOC"], "#e8d9f0")
s["spring"] = draw_entity(ax, 14.1, 10.3, 3.7, "SPRING_Stock / Kardex / Transacciones", [
    "Tablas espejo inventario ERP", "sincronizacion bidireccional"], "#e8d9f0")
s["cola"] = draw_entity(ax, 14.1, 8.3, 3.7, "LOGISTICA_ColaMigracionSpring", [
    "PK id", "tipoDocumento", "payload JSON", "estado", "reintentos"], "#e8d9f0")
s["synclog"] = draw_entity(ax, 14.1, 6.1, 3.7, "LOGISTICA_SincronizacionLog", [
    "PK id", "tipo", "documento", "resultado", "fecha"], "#e8d9f0")
s["maestro"] = draw_entity(ax, 14.1, 4.2, 3.7, "LOGISTICA_MaestroItemCache / Sync", [
    "Cache del maestro de items ERP", "sincronizado por job 11pm"], "#e8d9f0")
s["hist"] = draw_entity(ax, 14.1, 2.2, 3.7, "LOGISTICA_HistorialPrecioServicio", [
    "PK id", "servicio / proveedor", "precio", "fecha"], "#e8d9f0")

# --- Bloque Administracion / Configuracion (derecha) ---
s["adminusr"] = draw_entity(ax, 18.9, 14.4, 3.4, "LOGISTICA_AdminUsuarios", [
    "PK id", "usuario / documento", "passwordHash", "rol", "activo"], "#d0e0f0")
s["usrsis"] = draw_entity(ax, 18.9, 12.2, 3.4, "LOGISTICA_Usuarios / Roles / Perfiles", [
    "UsuariosRoles", "PerfilesUsuarios", "asignacion de accesos"], "#d0e0f0")
s["menu"] = draw_entity(ax, 18.9, 10.2, 3.4, "LOGISTICA_ConfigMenu", [
    "PK id", "rol / empresa", "items menu JSON", "activo"], "#d0e0f0")
s["perm"] = draw_entity(ax, 18.9, 8.4, 3.4, "LOGISTICA_ConfigPermisos / ConfiguracionPermiso", [
    "permisos por rol", "valorMax aprobacion", "autoAprobacion"], "#d0e0f0")
s["param"] = draw_entity(ax, 18.9, 6.2, 3.4, "LOGISTICA_Parametros / Auditoria", [
    "parametros generales", "LOGISTICA_Auditoria: trazabilidad acciones"], "#d0e0f0")
s["smtp"] = draw_entity(ax, 18.9, 4.4, 3.4, "ConfigSmtp / PlantillasCorreo / PlantillasPdf", [
    "configuracion correo", "plantillas de notificacion y PDF"], "#d0e0f0")
s["brand"] = draw_entity(ax, 18.9, 2.4, 3.4, "LOGISTICA_ConfigBranding / ConfigEmpresa", [
    "logo, colores, RUC", "datos para PDF de OC/OS"], "#d0e0f0")
s["um"] = draw_entity(ax, 18.9, 0.5, 3.4, "LOGISTICA_UnidadMedida / Correlativo", [
    "maestro unidades", "correlativos de documentos"], "#d0e0f0")

# --- Bloque Notificaciones / Inventario extra (extremo derecho) ---
s["notif"] = draw_entity(ax, 22.9, 14.4, 2.3, "LOGISTICA_Notificaciones", [
    "logistica_notificaciones", "NotificacionesStock", "AlertaStock"], "#fce5cd")
s["inv"] = draw_entity(ax, 22.9, 11.9, 2.3, "Inventario (gestion)", [
    "AjusteInventario(+Det)", "ConteoInventario(+Det)",
    "ListaStock(+Det)"], "#fce5cd")
s["devol"] = draw_entity(ax, 22.9, 9.4, 2.3, "Devoluciones", [
    "DevolucionProveedor(+Det)", "devolucion_consumo(+det)",
    "Reingreso(+Detalle)"], "#fce5cd")
s["confalm"] = draw_entity(ax, 22.9, 7.1, 2.3, "Conformidad / Transferencia", [
    "ConformidadAlmacen", "RequerimientoTransferencia",
    "PlantillaAlmacen"], "#fce5cd")
s["increm"] = draw_entity(ax, 22.9, 4.8, 2.3, "Incremento OC", [
    "ConfigIncrementoOC", "ConfigIncrementoOCAuditoria",
    "OCIncrementoLog"], "#fce5cd")
s["logs"] = draw_entity(ax, 22.9, 2.5, 2.3, "Logs varios", [
    "LogCambioEstadoReq.", "HistorialAnulaciones",
    "CorreoLog", "AppDeployments"], "#fce5cd")

# Flechas servicios
arrow(ax, bottom(s["ss"]), top(s["ssdet"]), "1:N")
arrow(ax, bottom(s["ssdet"]), top(s["ssadj"]), "1:N")
arrow(ax, bottom(s["ssadj"]), top(s["aprss"]), "aprueba")
arrow(ax, right(s["ss"]), left(s["cotsrv"]), "", 0.1)
arrow(ax, bottom(s["cotsrv"]), top(s["cotsrvdet"]), "1:N")
arrow(ax, bottom(s["cotsrvdet"]), top(s["os"]), "genera")
arrow(ax, bottom(s["os"]), top(s["osdet"]), "1:N")
arrow(ax, bottom(s["osdet"]), top(s["apros"]), "aprueba")
arrow(ax, bottom(s["apros"]), top(s["conf"]), "conformidad")
arrow(ax, bottom(s["conf"]), top(s["segos"]), "seguimiento")
# Flechas consol
arrow(ax, bottom(s["consoloc"]), top(s["consolocdet"]), "1:N")
arrow(ax, bottom(s["consolocdet"]), top(s["consoldist"]), "distribuye")
arrow(ax, bottom(s["consolos"]), top(s["consolosdet"]), "1:N")
arrow(ax, bottom(s["consolosdet"]), top(s["consolconf"]), "conformidad")
arrow(ax, bottom(s["consolconf"]), top(s["consoladj"]), "")
# Integracion
arrow(ax, right(s["os"]), left(s["whos"]), "sync", 0.15)
arrow(ax, right(s["consoloc"]), left(s["whoc"]), "sync", 0.1)
arrow(ax, bottom(s["whoc"]), top(s["spring"]), "")
arrow(ax, bottom(s["spring"]), top(s["cola"]), "")
arrow(ax, bottom(s["cola"]), top(s["synclog"]), "")
arrow(ax, right(s["whos"]), left(s["adminusr"]), "", 0.1)
arrow(ax, right(s["consolos"]), left(s["whos"]), "", -0.15)

plt.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "der_servicios_admin.png"), dpi=140, bbox_inches="tight")
plt.close(fig)
print("OK der_servicios_admin.png")

# =====================================================================
# FIGURA 3: DIAGRAMA DE CASOS DE USO
# =====================================================================
fig, ax = plt.subplots(figsize=(20, 15))
ax.set_xlim(0, 20); ax.set_ylim(0, 15); ax.axis("off")
ax.set_title("Diagrama de Casos de Uso - Sistema de Logistica", fontsize=14, fontweight="bold", pad=10)


def draw_actor(ax, x, y, name, side="left"):
    """Stick figure simple."""
    ax.add_patch(plt.Circle((x, y + 0.55), 0.16, fc="white", ec="black", lw=1.4, zorder=3))
    ax.plot([x, x], [y + 0.39, y - 0.15], color="black", lw=1.4, zorder=3)
    ax.plot([x - 0.28, x + 0.28], [y + 0.22, y + 0.22], color="black", lw=1.4, zorder=3)
    ax.plot([x, x - 0.22], [y - 0.15, y - 0.55], color="black", lw=1.4, zorder=3)
    ax.plot([x, x + 0.22], [y - 0.15, y - 0.55], color="black", lw=1.4, zorder=3)
    ax.text(x, y - 0.85, name, ha="center", va="top", fontsize=8.5, fontweight="bold",
            wrap=True, zorder=4)


def draw_uc(ax, x, y, text, w=2.6, h=0.62, fs=7.5):
    ax.add_patch(Ellipse((x, y), w, h, fc="#eef5fc", ec="#2e5f8a", lw=1.0, zorder=2))
    ax.text(x, y, text, ha="center", va="center", fontsize=fs, zorder=3)
    return (x, y)


def link(ax, actor_xy, uc_xy):
    ax.plot([actor_xy[0], uc_xy[0]], [actor_xy[1], uc_xy[1]],
            color="#888888", lw=0.8, zorder=1)


# Limites del sistema
ax.add_patch(FancyBboxPatch((2.6, 0.4), 14.6, 14.0,
             boxstyle="round,pad=0.05,rounding_size=0.15",
             fc="none", ec="#333333", lw=1.5, zorder=0))
ax.text(9.9, 14.05, "Sistema de Logistica", ha="center", fontsize=10,
        fontweight="bold", style="italic")

# ----- Grupos (sub-limites) -----
groups = [
    (3.0, 11.5, 4.5, 2.6, "Requerimientos", "#f4d9d0"),
    (3.0, 8.3, 4.5, 3.0, "Aprobaciones", "#fde9c8"),
    (3.0, 3.4, 4.5, 4.7, "Compras", "#d0e0f0"),
    (8.0, 11.5, 4.4, 2.6, "Servicios", "#d9ead3"),
    (8.0, 6.6, 4.4, 4.7, "Almacen / Inventario", "#cfe2f3"),
    (8.0, 0.7, 4.4, 5.7, "Administracion", "#e8d9f0"),
    (12.9, 8.0, 4.0, 6.1, "Consultas / Notificaciones", "#fce5cd"),
    (12.9, 0.7, 4.0, 7.1, "Integracion", "#ead1dc"),
]
for gx, gy, gw, gh, gname, gcolor in groups:
    ax.add_patch(FancyBboxPatch((gx, gy), gw, gh,
                 boxstyle="round,pad=0.03,rounding_size=0.1",
                 fc=gcolor, ec="#666666", lw=0.8, alpha=0.35, zorder=0))
    ax.text(gx + gw / 2, gy + gh - 0.18, gname, ha="center", va="top",
            fontsize=8, fontweight="bold", zorder=1)

# ----- Casos de uso por grupo -----
ucs = {}

def uc_group(x0, y0, texts, dy=0.62):
    res = []
    for i, t in enumerate(texts):
        res.append(draw_uc(ax, x0, y0 - i * dy, t))
    return res

ucs["req"] = uc_group(5.25, 13.35, [
    "Registrar requerimiento",
    "Adjuntar documentos",
    "Consultar mis requerimientos",
    "Anular requerimiento"])

ucs["apr"] = uc_group(5.25, 10.7, [
    "Aprobar / rechazar requerimiento",
    "Aprobar saldo pendiente",
    "Aprobar solicitud de compra",
    "Ver reporte de aprobaciones"])

ucs["com"] = uc_group(5.25, 7.4, [
    "Consolidar requerimientos",
    "Generar solicitud de cotizacion",
    "Registrar y evaluar cotizaciones",
    "Generar solicitud de compra",
    "Generar orden de compra",
    "Enviar OC a aprobacion",
    "Sincronizar OC con ERP"])

ucs["srv"] = uc_group(10.2, 13.35, [
    "Registrar solicitud de servicio",
    "Aprobar solicitud de servicio",
    "Generar orden de servicio",
    "Registrar conformidad de servicio"])

ucs["alm"] = uc_group(10.2, 10.6, [
    "Recepcionar mercaderia (OC)",
    "Despachar requerimientos",
    "Consultar kardex / stock",
    "Gestionar inventario (ajustes)",
    "Transferir entre almacenes",
    "Registrar devolucion / reingreso",
    "Conformidad de almacen"])

ucs["adm"] = uc_group(10.2, 5.8, [
    "Gestionar usuarios y roles",
    "Asignar usuarios por area",
    "Configurar flujos de aprobacion",
    "Configurar menu y permisos",
    "Parametros / SMTP / plantillas",
    "Maestros (items, proveedores)",
    "Consultar auditoria",
    "Gestionar branding / empresas"])

ucs["con"] = uc_group(14.9, 13.55, [
    "Ver dashboard por rol",
    "Generar reportes",
    "Recibir notificaciones",
    "Consultar seguimiento OC/OS"])

ucs["int"] = uc_group(14.9, 7.2, [
    "Sincronizar requerimientos",
    "Sincronizar OC/OS con Spring",
    "Enviar correo a proveedor",
    "Validar credenciales",
    "Cola de migracion / reintentos"])

# ----- Actores -----
actors = {
    "op":   (0.9, 12.6, "Solicitante\n(Operativo)"),
    "jefe": (0.9, 9.8, "Jefe de Area\n(Aprobador)"),
    "log":  (0.9, 5.4, "Analista de\nLogistica"),
    "alm":  (19.0, 9.6, "Almacenero"),
    "fin":  (19.0, 5.6, "Aprobador OC/OS\n(Finanzas/Gerencia)"),
    "conf": (19.0, 7.6, "Usuario\nConformidad"),
    "adm":  (19.0, 3.4, "Administrador\nLogistica"),
    "sis":  (19.0, 1.5, "Administrador\nSistema"),
    "prov": (19.0, 12.6, "Proveedor\n(externo)"),
    "erp":  (19.0, 14.0, "ERP Spring\n(sistema)"),
    "api":  (0.9, 14.0, "API Maestra\n(sistema)"),
}
apos = {}
for k, (x, y, n) in actors.items():
    draw_actor(ax, x, y, n)
    apos[k] = (x, y)

# Asociaciones actor <-> CU (lado izquierdo)
link(ax, (1.2, 12.4), ucs["req"][0]); link(ax, (1.2, 12.4), ucs["req"][1])
link(ax, (1.2, 12.4), ucs["req"][2]); link(ax, (1.2, 12.4), ucs["req"][3])
link(ax, (1.2, 9.6), ucs["apr"][0]); link(ax, (1.2, 9.6), ucs["apr"][1])
link(ax, (1.2, 9.6), ucs["apr"][3])
link(ax, (1.2, 5.2), ucs["com"][0]); link(ax, (1.2, 5.2), ucs["com"][1])
link(ax, (1.2, 5.2), ucs["com"][2]); link(ax, (1.2, 5.2), ucs["com"][3])
link(ax, (1.2, 5.2), ucs["com"][4]); link(ax, (1.2, 5.2), ucs["com"][5])
link(ax, (1.2, 5.2), ucs["com"][6])
link(ax, (1.2, 5.2), ucs["srv"][2]); link(ax, (1.2, 5.2), ucs["srv"][0])
link(ax, (1.2, 13.8), ucs["int"][3])   # API Maestra -> validar credenciales
link(ax, (1.2, 9.6), ucs["srv"][1])   # Jefe aprueba SS
link(ax, (1.2, 12.4), ucs["int"][0])  # solicitante sincroniza

# Asociaciones actor <-> CU (lado derecho)
link(ax, (18.6, 9.4), ucs["alm"][0]); link(ax, (18.6, 9.4), ucs["alm"][1])
link(ax, (18.6, 9.4), ucs["alm"][2]); link(ax, (18.6, 9.4), ucs["alm"][3])
link(ax, (18.6, 9.4), ucs["alm"][4]); link(ax, (18.6, 9.4), ucs["alm"][5])
link(ax, (18.6, 7.4), ucs["alm"][6])  # conformidad almacen
link(ax, (18.6, 7.4), ucs["srv"][3])  # conformidad servicio
link(ax, (18.6, 5.4), ucs["apr"][2])  # finanzas aprueba SC? -> Aprobar solicitud compra
link(ax, (18.6, 5.4), ucs["com"][5])  # enviar OC a aprobacion (aprueba OC)
link(ax, (18.6, 3.2), ucs["adm"][0]); link(ax, (18.6, 3.2), ucs["adm"][1])
link(ax, (18.6, 3.2), ucs["adm"][2]); link(ax, (18.6, 3.2), ucs["adm"][3])
link(ax, (18.6, 3.2), ucs["adm"][4]); link(ax, (18.6, 3.2), ucs["adm"][5])
link(ax, (18.6, 1.3), ucs["adm"][6]); link(ax, (18.6, 1.3), ucs["adm"][7])
link(ax, (18.6, 12.4), ucs["int"][2])  # proveedor <- correo
link(ax, (18.6, 13.8), ucs["int"][1]); link(ax, (18.6, 13.8), ucs["int"][4])  # ERP
link(ax, (18.6, 9.4), ucs["con"][0]); link(ax, (18.6, 9.4), ucs["con"][1])
link(ax, (1.2, 12.4), ucs["con"][0]); link(ax, (1.2, 12.4), ucs["con"][2])
link(ax, (1.2, 5.2), ucs["con"][3])

plt.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "casos_uso.png"), dpi=140, bbox_inches="tight")
plt.close(fig)
print("OK casos_uso.png")

# =====================================================================
# FIGURA 4: FLUJO GENERAL DEL PROCESO
# =====================================================================
fig, ax = plt.subplots(figsize=(22, 8))
ax.set_xlim(0, 22); ax.set_ylim(0, 8); ax.axis("off")
ax.set_title("Flujo general del proceso logistico", fontsize=13, fontweight="bold")


def step(x, y, w, h, title, sub, color):
    ax.add_patch(FancyBboxPatch((x, y), w, h,
                 boxstyle="round,pad=0.04,rounding_size=0.12",
                 fc=color, ec="#444444", lw=1.2, zorder=2))
    ax.text(x + w / 2, y + h * 0.62, title, ha="center", va="center",
            fontsize=10, fontweight="bold", zorder=3)
    ax.text(x + w / 2, y + h * 0.30, sub, ha="center", va="center",
            fontsize=7.5, zorder=3)
    return (x, y, w, h)


f = {}
f["req"] = step(0.3, 4.6, 2.4, 1.5, "1. Requerimiento", "Operativo registra\ncompra / consumo /\nservicio / commodity", "#f4d9d0")
f["apr"] = step(3.4, 4.6, 2.4, 1.5, "2. Aprobacion", "Jefe de area aprueba\n(flujo por niveles)", "#fde9c8")
f["con"] = step(6.5, 4.6, 2.4, 1.5, "3. Consolidacion", "Logistica agrupa\nitems aprobados", "#d9ead3")
f["cot"] = step(9.6, 4.6, 2.4, 1.5, "4. Cotizacion", "Solicitud a proveedores\nevaluacion y seleccion", "#d0e0f0")
f["oc"] = step(12.7, 4.6, 2.4, 1.5, "5. SC / OC / OS", "Solicitud de compra\nOrden de compra/servicio", "#e8d9f0")
f["aproc"] = step(15.8, 4.6, 2.4, 1.5, "6. Aprobacion OC", "Niveles por monto\nFinanzas / Gerencia", "#e8d9f0")
f["erp"] = step(19.0, 4.6, 2.4, 1.5, "7. Sincronizacion", "OC/OS a ERP Spring\ncola de reintentos", "#ead1dc")
f["rec"] = step(19.0, 1.6, 2.4, 1.5, "8. Recepcion", "Almacen recibe\nmercaderia / NI", "#cfe2f3")
f["kar"] = step(15.8, 1.6, 2.4, 1.5, "9. Kardex / Stock", "Ingreso a inventario\nmovimiento kardex", "#cfe2f3")
f["desp"] = step(12.7, 1.6, 2.4, 1.5, "10. Despacho", "Atencion del\nrequerimiento", "#cfe2f3")
f["saldo"] = step(9.6, 1.6, 2.4, 1.5, "11. Saldo pendiente", "Aprobacion de saldo\no nueva consolidacion", "#fce5cd")


def conn(a, b, vertical=False):
    if vertical:
        p1 = (a[0] + a[2] / 2, a[1]); p2 = (b[0] + b[2] / 2, b[1] + b[3])
    else:
        p1 = (a[0] + a[2], a[1] + a[3] / 2); p2 = (b[0], b[1] + b[3] / 2)
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=16,
                 lw=1.6, color="#333333", zorder=1))


conn(f["req"], f["apr"]); conn(f["apr"], f["con"]); conn(f["con"], f["cot"])
conn(f["cot"], f["oc"]); conn(f["oc"], f["aproc"]); conn(f["aproc"], f["erp"])
conn(f["erp"], f["rec"], vertical=True)
p1 = (f["rec"][0], f["rec"][1] + 0.75); p2 = (f["kar"][0] + f["kar"][2], f["kar"][1] + 0.75)
ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=16, lw=1.6, color="#333333"))
conn(f["kar"], f["desp"]); conn(f["desp"], f["saldo"])
# retorno saldo -> consolidacion
ax.add_patch(FancyArrowPatch((f["saldo"][0], f["saldo"][1] + 1.5), (f["con"][0] + 1.2, f["con"][1]),
             arrowstyle="-|>", mutation_scale=14, lw=1.2, color="#999999",
             connectionstyle="arc3,rad=0.35", linestyle="--"))
ax.text(7.6, 2.6, "re-consolidar", fontsize=7, color="#777777", style="italic")

plt.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "flujo_general.png"), dpi=140, bbox_inches="tight")
plt.close(fig)
print("OK flujo_general.png")


# =====================================================================
# DOCUMENTO WORD
# =====================================================================
doc = Document()

# Estilos base
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)

AZUL = RGBColor(0x1F, 0x4E, 0x79)


def h1(text):
    p = doc.add_heading(text, level=1)
    return p


def h2(text):
    return doc.add_heading(text, level=2)


def h3(text):
    return doc.add_heading(text, level=3)


def par(text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return p


def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(htxt)
        r.bold = True
        r.font.size = Pt(9)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # sombreado encabezado
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


# ---------- PORTADA ----------
for _ in range(6):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SISTEMA DE LOGISTICA")
r.font.size = Pt(28); r.bold = True; r.font.color.rgb = AZUL
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Documentacion Tecnica Funcional")
r.font.size = Pt(18); r.font.color.rgb = AZUL
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Modelo de Base de Datos  |  Modelo de Casos de Uso  |  Descripcion de Flujos")
r.font.size = Pt(12); r.italic = True
for _ in range(8):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Componentes: api_logistica (.NET)  +  logistica (Angular)  +  SQL Server (BD LOGISTICA)")
r.font.size = Pt(10)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Fecha de elaboracion: 09/10/2026")
r.font.size = Pt(10)
doc.add_page_break()

# ---------- 1. INTRODUCCION ----------
h1("1. Introduccion")
par("El sistema de Logistica es una aplicacion web que gestiona el ciclo completo de "
    "abastecimiento interno: requerimientos, aprobaciones por area, consolidacion, "
    "cotizaciones, solicitudes y ordenes de compra, ordenes de servicio, recepcion de "
    "mercaderia, almacen, inventario, kardex, despachos, devoluciones y administracion "
    "del sistema.")
par("Arquitectura de la solucion:", bold=True)
bullet("Frontend: aplicacion Angular (proyecto 'logistica') con modulos por rol y guards de acceso.")
bullet("Backend: API REST en .NET (proyecto 'api_logistica') con arquitectura Domain / Application / Infrastructure / Controllers.")
bullet("Base de datos: SQL Server, base de datos LOGISTICA; la logica de negocio reside principalmente en Stored Procedures.")
bullet("Integraciones: ERP Spring (tablas WH_* / SPRING_* y cola de migracion), API Maestra (autenticacion de usuarios), servidor SMTP para correos a proveedores.")
doc.add_paragraph()
doc.add_picture(os.path.join(OUT_DIR, "flujo_general.png"), width=Inches(6.8))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
par("Figura 1. Flujo general del proceso logistico.", italic=True, size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# ---------- 2. MODELO DE BASE DE DATOS ----------
h1("2. Modelo de Base de Datos")
par("La base de datos LOGISTICA (SQL Server) contiene mas de 90 tablas. Se organizan en "
    "los siguientes modulos funcionales. Las tablas principales llevan el prefijo "
    "LOGISTICA_; las tablas logistica_req y logistica_dreq son el nucleo de requerimientos; "
    "las tablas WH_* y SPRING_* son de integracion con el ERP.")

h2("2.1 Diagrama Entidad-Relacion - Flujo principal (Compras y Almacen)")
doc.add_picture(os.path.join(OUT_DIR, "der_principal.png"), width=Inches(6.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
par("Figura 2. DER del flujo principal: requerimiento, aprobacion, consolidacion, "
    "cotizacion, solicitud/orden de compra, recepcion e inventario.", italic=True, size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

h2("2.2 Diagrama Entidad-Relacion - Servicios, Consolidacion OC/OS, Administracion e Integracion")
doc.add_picture(os.path.join(OUT_DIR, "der_servicios_admin.png"), width=Inches(6.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
par("Figura 3. DER de servicios, ordenes generadas por consolidacion, administracion "
    "y tablas de integracion con el ERP Spring.", italic=True, size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

h2("2.3 Catalogo de tablas por modulo")

modulos = [
    ("Requerimientos", [
        ("logistica_req", "Cabecera de requerimientos (compra, consumo, servicio).", "idrequerimiento, RequisicionNumero, ruc, idarea, tipo, estado, dniregistra"),
        ("logistica_dreq", "Detalle de items del requerimiento.", "idrequerimiento, codigo, cantidad, idcentrocosto, idproyecto"),
        ("LOGISTICA_RequerimientoAdjuntos", "Archivos adjuntos del requerimiento (B64).", "idrequerimiento, nombreArchivo, contenidoB64"),
        ("LOGISTICA_ReqCommodity / LOGISTICA_DReqCommodity", "Requerimientos de commodities y su detalle.", "idRequerimiento, tipo, estado"),
        ("LOGISTICA_RequerimientosPorArea", "Asociacion requerimiento-area para flujo de aprobacion.", "idRequerimiento, idarea"),
        ("LOGISTICA_RequerimientoTransferencia", "Requerimientos de transferencia entre almacenes.", "almacenOrigen, almacenDestino, estado"),
        ("LOGISTICA_SolicitudCreacionItem", "Solicitudes de creacion de nuevos items.", "descripcion, estado"),
        ("LOGISTICA_HistorialAnulaciones", "Trazabilidad de anulaciones de requerimientos.", "idRequerimiento, motivo, usuario"),
        ("LOGISTICA_LogCambioEstadoRequerimiento", "Log de cambios de estado del requerimiento.", "idRequerimiento, estadoAnterior, estadoNuevo"),
    ]),
    ("Aprobaciones", [
        ("LOGISTICA_Areas", "Areas / subareas de la empresa.", "idarea, nombreArea"),
        ("LOGISTICA_UsuariosPorArea", "Usuarios asignados a areas; marca jefes de area.", "idarea, documentoidentidad, esJefeArea"),
        ("LOGISTICA_FlujoAprobacion(PorArea)", "Niveles de aprobacion configurados por area.", "idarea, nivelOrden, tipoAprobador"),
        ("LOGISTICA_LogAprobaciones(PorArea)", "Registro de aprobaciones/rechazos realizados.", "idRequerimiento, nivelAprobacion, estado"),
        ("LOGISTICA_SaldoPendienteAprobacion(+Det)", "Saldos pendientes generados en despachos parciales.", "idSolicitud, idrequerimiento, saldoPendiente, estado"),
        ("LOGISTICA_FlujoAprobacionSC / AprobacionSC / UsuariosAprobadoresSC / LogAprobacionesSC", "Flujo multinivel de aprobacion de solicitudes de compra.", "nivelOrden, montoMin/Max, aprobador, estado"),
        ("LOGISTICA_FlujoAprobacionOC / AprobacionOC / UsuariosAprobadoresOC / LogAprobacionesOC", "Flujo multinivel de aprobacion de ordenes de compra.", "nivelOrden, montoMaximo, aprobador, estado"),
        ("LOGISTICA_FlujoAprobacionOS / AprobacionOS / UsuarioAprobadorOS / LogAprobacionesOS", "Flujo de aprobacion de ordenes de servicio.", "idOS, nivelAprobacion, estado"),
        ("LOGISTICA_AprobacionSS", "Aprobacion de solicitudes de servicio.", "idSolicitud, nivelAprobacion, estado"),
        ("LOGISTICA_AprobacionOrden(+Detalle) / LOGISTICA_AdjuntosOrden", "Aprobacion generica de ordenes OC/OS y sus adjuntos.", "idOrden, tipoOrden, estado"),
    ]),
    ("Consolidacion", [
        ("LOGISTICA_ConsolidacionCab", "Cabecera de consolidacion de requerimientos.", "idConsolidacion, codigo, estado, estadoProceso"),
        ("LOGISTICA_ConsolidacionDet", "Lineas consolidadas por item.", "idConsolidacion, codigoItem, cantidadTotal"),
        ("LOGISTICA_ConsolidacionOrigen", "Trazabilidad: lineas origen vs. linea consolidada.", "idDetConsolidacion, tipoOrigen, requerimientoOrigen"),
        ("LOGISTICA_ItemPendienteConsolidacion", "Control de detalles pendientes de consolidar.", "idDetalle, tipoRequerimiento, estadoConsolidacion"),
        ("RequerimientoConsolidadoCompra / DetalleRequerimientoConsolidado / ReqConsolidado", "Esquema alterno de consolidacion de compras.", "no_reqconsolcompra, codigo_item, cantidad_total"),
    ]),
    ("Compras (cotizaciones, SC, OC)", [
        ("LOGISTICA_SolicitudCotizacion(+Det)", "Solicitudes de cotizacion generadas desde consolidacion.", "idConsolidacion, noSolicitud, estado"),
        ("LOGISTICA_Cotizacion(+Det)", "Cotizaciones registradas por proveedor.", "idSolicitudCotizacion, proveedor, total, seleccionada"),
        ("LOGISTICA_CotizacionOC(+Detalle)", "Cotizaciones asociadas a OC (multiempresa).", "idCotizacion, proveedor, monto"),
        ("SolicitudCotizacion / DetalleSolicitudCotizacion / CotizacionProveedor", "Esquema alterno de cotizaciones.", "no_solicitud, codigo_proveedor, seleccionado"),
        ("LOGISTICA_SolicitudCompra(+Det)", "Solicitudes de compra (consolidada, directa, urgente).", "numeroSolicitud, tipo, estado, usuarioSolicita"),
        ("LOGISTICA_SolicitudCompraAdjuntos", "Adjuntos de la solicitud de compra.", "idSolicitud, nombreArchivo"),
        ("LOGISTICA_OrdenCompra(+Det)", "Ordenes de compra.", "numeroOrden, idSolicitud, proveedor, moneda, totalOrden, estado"),
        ("LOGISTICA_OrdenCompraLog", "Historial de cambios de estado de la OC.", "idOrden, accion, estadoAnterior/Nuevo"),
        ("LOGISTICA_OCAdjunto", "Adjuntos de la OC.", "idOrden, nombreArchivo"),
        ("LOGISTICA_SeguimientoOrdenCompra", "Seguimiento/etapas de la OC.", "idOrden, etapa, estado"),
        ("LOGISTICA_ConfigIncrementoOC(+Auditoria) / LOGISTICA_OCIncrementoLog", "Configuracion y log de incrementos de cantidad en OC.", "idOrden, porcentaje, usuario"),
        ("LOGISTICA_Proveedor / LOGISTICA_CategoriaProveedor / ProveedoresEmail", "Maestro de proveedores y correos de contacto.", "ruc, razonSocial, email"),
        ("LOGISTICA_Entrega", "Entregas programadas de proveedores.", "idOrden, fechaEntrega"),
    ]),
    ("Servicios (SS, OS)", [
        ("LOGISTICA_SolicitudServicio(+Det)", "Solicitudes de servicio y su detalle.", "numeroSolicitud, tipoServicio, estado"),
        ("LOGISTICA_SolicitudServicioAdjuntos", "Adjuntos de la solicitud de servicio.", "idSolicitud, nombreArchivo"),
        ("LOGISTICA_CotizacionServicio(+Det)", "Cotizaciones de servicios por proveedor.", "idSolicitud, proveedor, montoTotal"),
        ("LOGISTICA_OrdenServicio(+Det)", "Ordenes de servicio.", "numeroOrden, proveedor, fechaInicio/Fin, montoTotal, estado"),
        ("LOGISTICA_ConformidadServicio / LOGISTICA_OSConformidad", "Conformidad del servicio ejecutado (firma digital).", "idOS, estado, firmaBase64"),
        ("LOGISTICA_ConfirmacionServicioRealizado", "Confirmacion de servicio realizado.", "idOS, fecha, usuario"),
        ("LOGISTICA_SeguimientoOrdenServicio", "Seguimiento de etapas de la OS.", "idOS, etapa, estado"),
        ("LOGISTICA_HistorialPrecioServicio", "Historial de precios de servicios.", "servicio, proveedor, precio"),
        ("LOGISTICA_PlantillaConformidadServicio", "Plantilla de acta de conformidad.", "contenido, activo"),
    ]),
    ("Consolidacion OC/OS (multiempresa)", [
        ("LOGISTICA_ConsolOC(+Det)", "OC generadas desde consolidacion.", "numeroOrden, idConsolidacion, proveedor, totalOrden, estado"),
        ("LOGISTICA_ConsolOCDist", "Distribucion de la OC por requerimiento.", "idOrden, cantidadAsignada"),
        ("LOGISTICA_ConsolOS(+Det)", "OS generadas desde consolidacion.", "numeroOrden, idConsolidacion, proveedor, montoTotal, estado"),
        ("LOGISTICA_ConsolOSConformidad", "Conformidad de OS de consolidacion (firma/acta).", "idOS, estado, firmaBase64, actaBase64"),
        ("LOGISTICA_ConsolAdjunto", "Adjuntos de OC y OS de consolidacion.", "idOrden, tipoOrden, nombreArchivo"),
        ("LOGISTICA_ConsolConfigEmpresa", "Datos de empresa para PDF de OC/OS.", "ruc, razonSocial, logoBase64, smtp*"),
        ("LOGISTICA_ConsolCorreoLog", "Log de correos enviados a proveedores.", "tipoOrden, codigoOrden, destinatario, estado"),
    ]),
    ("Almacen e Inventario", [
        ("LOGISTICA_RecepcionOC(+Det+Log)", "Recepciones de mercaderia contra OC.", "idOrden, numeroNI, cantidadRecibida/Aceptada, estado"),
        ("LOGISTICA_RecepcionOrdenCompra(+Detalle)", "Esquema alterno de recepcion de OC.", "idOrden, cantidad"),
        ("LOGISTICA_Despachos / LOGISTICA_DetalleDespacho", "Despachos de atencion de requerimientos.", "idDespacho, idrequerimiento, cantidadDespachada"),
        ("LOGISTICA_Item / LOGISTICA_Almacen", "Maestros de items y almacenes.", "codigo, descripcion, unidadMedida"),
        ("LOGISTICA_TransaccionInventario(+Detalle)", "Transacciones de ingreso/salida de inventario.", "tipoTransaccion, idAlmacen, documentoOrigen"),
        ("LOGISTICA_Kardex", "Movimientos de kardex valorizado.", "idItem, idAlmacen, tipoMovimiento, saldoCantidad/Valor"),
        ("LOGISTICA_Stock", "Stock actual por item y almacen.", "idItem, idAlmacen, stockActual/Disponible, costoPromedio"),
        ("LOGISTICA_ListaStock(+Det) / LOGISTICA_AlertaStock", "Listas de stock y alertas de disponibilidad.", "idItem, stockMinimo"),
        ("LOGISTICA_AjusteInventario(+Det) / LOGISTICA_ConteoInventario(+Det)", "Ajustes y conteos fisicos de inventario.", "tipoAjuste, cantidad, motivo"),
        ("LOGISTICA_DevolucionProveedor(+Det)", "Devoluciones de mercaderia al proveedor.", "idOrden, motivo, estado"),
        ("logistica_devolucion_consumo(+det)", "Devoluciones internas de consumo.", "idrequerimiento, cantidad"),
        ("LOGISTICA_Reingreso(+Detalle)", "Reingresos de stock al almacen.", "documentoOrigen, cantidad"),
        ("LOGISTICA_ConformidadAlmacen / LOGISTICA_PlantillaAlmacen", "Conformidad de almacen y su plantilla.", "idDocumento, estado, firma"),
    ]),
    ("Administracion y Configuracion", [
        ("LOGISTICA_AdminUsuarios", "Usuarios administradores del modulo logistica.", "usuario, passwordHash, rol, activo"),
        ("LOGISTICA_Usuarios / Roles / Perfiles / UsuariosRoles / PerfilesUsuarios", "Esquema de seguridad legacy.", "usuario, rol, perfil"),
        ("LOGISTICA_Usuario", "Usuarios del dashboard logistico.", "documento, nombre, rol"),
        ("LOGISTICA_ConfigMenu", "Menu dinamico por rol/empresa.", "rol, itemsMenu JSON"),
        ("LOGISTICA_ConfigPermisos / LOGISTICA_ConfiguracionPermiso", "Permisos por rol; montos maximos y autoaprobacion.", "rol, valorMax, autoAprobacion"),
        ("LOGISTICA_Parametros / LOGISTICA_ConfigPDF", "Parametros generales y configuracion de PDF.", "clave, valor"),
        ("LOGISTICA_Auditoria", "Auditoria de acciones del sistema.", "usuario, accion, fecha"),
        ("LOGISTICA_ConfigBranding / LOGISTICA_ConfigEmpresa", "Branding y datos de empresa para documentos.", "ruc, razonSocial, logoBase64"),
        ("ConfigSmtp / PlantillasCorreo / PlantillasPdf", "Configuracion SMTP y plantillas.", "host, puerto, plantilla"),
        ("LOGISTICA_UnidadMedida", "Maestro de unidades de medida.", "codigo, descripcion"),
        ("LOGISTICA_Correlativo", "Correlativos de documentos (CONS, SOL-COT, OC, OS).", "tipo, prefijo, anio, correlativo"),
    ]),
    ("Notificaciones y Logs", [
        ("logistica_notificaciones", "Notificaciones generales a usuarios.", "usuario, mensaje, leida"),
        ("LOGISTICA_NotificacionesStock / LOGISTICA_AlertaStock", "Notificaciones de stock disponible / alertas.", "idItem, mensaje, estado"),
        ("LOGISTICA_CorreoLog", "Log de correos enviados.", "destinatario, asunto, estado"),
    ]),
    ("Integracion ERP Spring y Despliegue", [
        ("WH_OrdenServicio(+Detalle)", "Tabla espejo de OS en el ERP Spring.", "numeroOrden, proveedor, estadoSync"),
        ("WH_OrdenCompra (ERP)", "Destino de OC aprobadas hacia Spring.", "numeroOrden, estadoSync"),
        ("SPRING_Stock / SPRING_Kardex / SPRING_Transacciones", "Espejo de inventario del ERP.", "idItem, cantidad, tipoMovimiento"),
        ("LOGISTICA_ColaMigracionSpring", "Cola de sincronizacion con reintentos.", "tipoDocumento, payload, estado, reintentos"),
        ("LOGISTICA_SincronizacionLog / LOGISTICA_SincronizacionSpringLog", "Log de sincronizaciones.", "tipo, documento, resultado"),
        ("LOGISTICA_MaestroItemCache / LOGISTICA_MaestroItemSync", "Cache del maestro de items del ERP (job nocturno).", "codigo, descripcion, fechaSync"),
        ("AppDeployments", "Registro de versiones/despliegues de la app.", "app, version, fechaDespliegue"),
    ]),
]

for nombre, tablas in modulos:
    h3(nombre)
    add_table(["Tabla", "Descripcion", "Campos clave"], tablas, widths=[2.4, 2.6, 2.2])
    doc.add_paragraph()

h2("2.4 Relaciones principales")
add_table(["Tabla origen", "Relacion", "Tabla destino"], [
    ("logistica_dreq", "N:1 (idrequerimiento)", "logistica_req"),
    ("LOGISTICA_RequerimientoAdjuntos", "N:1", "logistica_req"),
    ("LOGISTICA_UsuariosPorArea", "N:1 (idarea)", "LOGISTICA_Areas"),
    ("LOGISTICA_LogAprobacionesPorArea", "N:1", "logistica_req / LOGISTICA_Areas"),
    ("LOGISTICA_ItemPendienteConsolidacion", "N:1 (idDetalle)", "logistica_dreq"),
    ("LOGISTICA_ConsolidacionDet", "N:1 (idConsolidacion)", "LOGISTICA_ConsolidacionCab"),
    ("LOGISTICA_ConsolidacionOrigen", "N:1", "LOGISTICA_ConsolidacionDet / Cab"),
    ("LOGISTICA_SolicitudCotizacion", "N:1 (idConsolidacion)", "LOGISTICA_ConsolidacionCab"),
    ("LOGISTICA_SolicitudCotizacionDet", "N:1", "LOGISTICA_SolicitudCotizacion"),
    ("LOGISTICA_Cotizacion", "N:1", "LOGISTICA_SolicitudCotizacion"),
    ("LOGISTICA_CotizacionDet", "N:1", "LOGISTICA_Cotizacion"),
    ("LOGISTICA_SolicitudCompraDet / Adjuntos", "N:1 (idSolicitud)", "LOGISTICA_SolicitudCompra"),
    ("LOGISTICA_OrdenCompra", "N:1 (idSolicitud)", "LOGISTICA_SolicitudCompra"),
    ("LOGISTICA_OrdenCompraDet / Log / OCAdjunto", "N:1 (idOrden)", "LOGISTICA_OrdenCompra"),
    ("LOGISTICA_AprobacionOC", "N:1 (idOrden)", "LOGISTICA_OrdenCompra"),
    ("LOGISTICA_RecepcionOCDet / Log", "N:1 (idRecepcion)", "LOGISTICA_RecepcionOC"),
    ("LOGISTICA_RecepcionOC", "N:1 (idOrden)", "LOGISTICA_OrdenCompra"),
    ("LOGISTICA_DetalleTransaccionInventario", "N:1", "LOGISTICA_TransaccionInventario"),
    ("LOGISTICA_Kardex / LOGISTICA_Stock", "N:1", "LOGISTICA_Item / LOGISTICA_Almacen"),
    ("LOGISTICA_SolicitudServicioDet / Adjuntos", "N:1", "LOGISTICA_SolicitudServicio"),
    ("LOGISTICA_CotizacionServicio(Det)", "N:1", "LOGISTICA_SolicitudServicio"),
    ("LOGISTICA_OrdenServicioDet", "N:1 (idOS)", "LOGISTICA_OrdenServicio"),
    ("LOGISTICA_AprobacionOS / ConformidadServicio / SeguimientoOrdenServicio", "N:1", "LOGISTICA_OrdenServicio"),
    ("LOGISTICA_ConsolOCDet / ConsolOCDist / ConsolAdjunto", "N:1 (idOrden)", "LOGISTICA_ConsolOC"),
    ("LOGISTICA_ConsolOSDet / ConsolOSConformidad", "N:1 (idOS)", "LOGISTICA_ConsolOS"),
    ("LOGISTICA_DevolucionProveedor(Det)", "N:1", "LOGISTICA_OrdenCompra / RecepcionOC"),
    ("LOGISTICA_Reingreso(Detalle)", "N:1", "LOGISTICA_TransaccionInventario"),
], widths=[3.0, 1.6, 2.6])
doc.add_page_break()

# ---------- 3. MODELO DE CASOS DE USO ----------
h1("3. Modelo de Casos de Uso")

h2("3.1 Actores del sistema")
add_table(["Actor", "Tipo", "Descripcion", "Rol/Guard en el sistema"], [
    ("Solicitante (Operativo)", "Persona", "Registra requerimientos de compra, consumo, commodity, activos o servicios; consulta sus estados y sincroniza pendientes.", "OperativoGuard / OperativoEmpaqueGuard / EmpaqueGuard"),
    ("Jefe de Area (Aprobador)", "Persona", "Aprueba o rechaza requerimientos de su area segun el flujo por niveles; aprueba saldos pendientes.", "AprobadorGuard"),
    ("Analista de Logistica", "Persona", "Consolida requerimientos, gestiona cotizaciones, genera solicitudes de compra, ordenes de compra y de servicio; da seguimiento.", "LogisticoGuard / OperativoLogisticoGuard"),
    ("Almacenero", "Persona", "Recepciona mercaderia, despacha requerimientos, gestiona inventario, kardex, transferencias, devoluciones y reingresos.", "AlmacenGuard"),
    ("Aprobador OC/OS (Finanzas/Gerencia)", "Persona", "Aprueba solicitudes de compra y ordenes de compra/servicio segun niveles por monto.", "AprobadorOCGuard"),
    ("Usuario Conformidad", "Persona", "Registra la conformidad de servicios y de almacen (firma digital).", "ConformidadServiciosGuard / DevolucionConsumoGuard"),
    ("Administrador Logistica", "Persona", "Administra usuarios del modulo, parametros, plantillas, branding y configuracion avanzada.", "AdminLogisticaGuard"),
    ("Administrador Sistema", "Persona", "Configura roles, flujos, menus dinamicos, permisos, maestros y consulta auditoria.", "AdminGuard / AdministradorGuard"),
    ("Proveedor", "Externo", "Recibe solicitudes de cotizacion y ordenes de compra/servicio por correo electronico.", "Actor externo (correo SMTP)"),
    ("ERP Spring", "Sistema externo", "Recibe OC/OS aprobadas; sincroniza stock, kardex y maestro de items.", "Tablas WH_* / SPRING_* / cola de migracion"),
    ("API Maestra", "Sistema externo", "Valida las credenciales de los usuarios en el login.", "Servicio de autenticacion"),
], widths=[1.5, 0.9, 2.9, 1.9])
doc.add_paragraph()

h2("3.2 Diagrama de casos de uso")
doc.add_picture(os.path.join(OUT_DIR, "casos_uso.png"), width=Inches(6.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
par("Figura 4. Diagrama de casos de uso agrupados por modulo.", italic=True, size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

h2("3.3 Especificacion de casos de uso")

casos = [
    ("CU-01 Registrar requerimiento", "Solicitante (Operativo)",
     "El usuario crea un requerimiento de compra, consumo, commodity, activo fijo, activo menor o servicio.",
     "Usuario autenticado con rol operativo; catalogo de items disponible.",
     ["1. Ingresa a Requerimientos y selecciona el tipo.",
      "2. Completa datos generales (area, centro de costo, almacen, prioridad, glosa).",
      "3. Agrega items con cantidades y unidades de medida.",
      "4. Adjunta documentos si aplica.",
      "5. Guarda/envia; si no hay conexion queda local y luego se sincroniza.",
      "6. El sistema asigna el aprobador segun el flujo del area."],
     "Requerimiento en estado PENDIENTE visible en la bandeja del aprobador."),
    ("CU-02 Aprobar / rechazar requerimiento", "Jefe de Area",
     "El aprobador revisa y decide sobre los requerimientos de su area (flujo multinivel).",
     "Requerimiento PENDIENTE asignado al nivel del aprobador.",
     ["1. Abre la bandeja de aprobaciones (o Aprobaciones por Area).",
      "2. Revisa detalle, cantidades, solicitante y adjuntos.",
      "3. Aprueba (avanza al siguiente nivel o pasa a Logistica) o rechaza con observacion.",
      "4. El sistema registra el log de aprobacion y notifica."],
     "Requerimiento APROBADO/ANULADO; si hay mas niveles pasa al siguiente aprobador."),
    ("CU-03 Consolidar requerimientos", "Analista de Logistica",
     "Agrupa lineas aprobadas de distintos requerimientos en una consolidacion.",
     "Existen detalles aprobados pendientes de consolidacion.",
     ["1. Lista los items pendientes (comodities, activos, servicios, consumo).",
      "2. Filtra por empresa, almacen, item o tipo y selecciona lineas compatibles.",
      "3. Crea la consolidacion (LOGISTICA_ConsolidacionCab/Det/Origen).",
      "4. Decide el siguiente paso: cotizacion, solicitud de compra u orden directa."],
     "Consolidacion registrada; items marcados como CONSOLIDADO."),
    ("CU-04 Gestionar cotizaciones", "Analista de Logistica / Proveedor",
     "Genera solicitudes de cotizacion, registra propuestas de proveedores y selecciona ganador.",
     "Consolidacion en estado PENDIENTE_COTIZACION.",
     ["1. Genera la solicitud de cotizacion desde la consolidacion.",
      "2. Envia a proveedores (correo) y registra sus propuestas.",
      "3. Compara precios, plazos y condiciones; selecciona ganador.",
      "4. Cierra la cotizacion."],
     "Proveedor seleccionado; habilita generacion de SC/OC."),
    ("CU-05 Generar solicitud de compra", "Analista de Logistica",
     "Crea la solicitud de compra (consolidada, directa o urgente) con sus adjuntos.",
     "Items consolidados o requerimiento directo aprobado.",
     ["1. Genera la SC desde consolidacion o de forma directa.",
      "2. Adjunta documentos obligatorios.",
      "3. Envia a aprobacion segun flujo configurado."],
     "SC en estado ENVIADA; ingresa al flujo de aprobacion SC."),
    ("CU-06 Aprobar solicitud de compra", "Aprobador SC / Jefe de Area",
     "Aprueba o rechaza la solicitud de compra segun niveles y montos.",
     "SC en estado ENVIADA.",
     ["1. Revisa la SC y sus adjuntos.",
      "2. Aprueba (siguiente nivel o aprobada) o rechaza con motivo.",
      "3. El sistema registra LOGISTICA_AprobacionSC y notifica."],
     "SC APROBADA lista para generar OC, o RECHAZADA con motivo."),
    ("CU-07 Generar y aprobar orden de compra", "Analista de Logistica / Aprobador OC",
     "Genera la OC desde la SC o consolidacion, la envia a aprobacion multinivel y la sincroniza.",
     "SC aprobada o consolidacion con proveedor seleccionado.",
     ["1. Genera la OC con proveedor, moneda, condiciones y detalle.",
      "2. Envia a aprobacion; el flujo asigna niveles segun monto.",
      "3. Cada aprobador aprueba/rechaza (LOGISTICA_AprobacionOC).",
      "4. Aprobada: se genera PDF, se envia correo al proveedor y se sincroniza al ERP."],
     "OC APROBADA y sincronizada; disponible para recepcion."),
    ("CU-08 Recepcionar mercaderia", "Almacenero",
     "Registra la recepcion fisica de la OC, genera nota de ingreso y actualiza kardex/stock.",
     "OC en estado APROBADA/SINCRONIZADA pendiente de recepcion.",
     ["1. Lista OCs pendientes de recepcion y selecciona una.",
      "2. Registra cantidades recibidas/aceptadas por item.",
      "3. El sistema genera la transaccion de ingreso, actualiza kardex y stock.",
      "4. Si hay conformidad de almacen, la registra."],
     "Recepcion registrada (LOGISTICA_RecepcionOC); stock incrementado."),
    ("CU-09 Despachar requerimiento", "Almacenero",
     "Atiende requerimientos de consumo con el stock disponible.",
     "Requerimiento APROBADO de tipo consumo con stock.",
     ["1. Revisa requerimientos pendientes y valida stock.",
      "2. Despacha la cantidad disponible.",
      "3. Si queda saldo, se genera LOGISTICA_SaldoPendienteAprobacion.",
      "4. El kardex registra la salida."],
     "Requerimiento ATENDIDO total o parcial; saldo pendiente si aplica."),
    ("CU-10 Gestionar ordenes de servicio", "Analista / Aprobador OS / Usuario Conformidad",
     "Ciclo de servicios: solicitud, aprobacion SS, cotizacion, OS, aprobacion OS y conformidad.",
     "Solicitud de servicio registrada.",
     ["1. Registra la solicitud de servicio (SS) y adjuntos.",
      "2. Flujo de aprobacion SS por niveles.",
      "3. Cotiza servicios y genera la OS.",
      "4. Aprobacion OS multinivel; sincroniza a WH_OrdenServicio (ERP).",
      "5. Ejecutado el servicio, el usuario registra la conformidad con firma digital."],
     "OS CONFORME; servicio cerrado y trazado."),
    ("CU-11 Gestionar inventario y kardex", "Almacenero",
     "Consulta kardex, gestiona ajustes, conteos, listas de stock y transferencias entre almacenes.",
     "Items y almacenes registrados.",
     ["1. Consulta kardex por item/almacen (entradas, salidas, saldos valorizados).",
      "2. Registra ajustes de inventario o conteos fisicos.",
      "3. Configura listas de stock y alertas.",
      "4. Transfiere stock entre almacenes (requerimiento de transferencia)."],
     "Inventario consistente; movimientos trazados en kardex."),
    ("CU-12 Registrar devoluciones y reingresos", "Almacenero",
     "Devolucion a proveedor, devolucion de consumo interno y reingreso de material.",
     "Recepcion o despacho previo registrado.",
     ["1. Selecciona el documento origen (OC/requerimiento).",
      "2. Registra cantidades y motivo.",
      "3. El sistema genera la transaccion inversa en kardex y ajusta stock."],
     "Stock y kardex actualizados segun el tipo de devolucion/reingreso."),
    ("CU-13 Administrar el sistema", "Administrador Logistica / Sistema",
     "Gestiona usuarios, areas, roles, flujos de aprobacion, menu dinamico, permisos, maestros, SMTP, plantillas y branding.",
     "Usuario con rol administrador.",
     ["1. Gestiona usuarios y su asignacion a areas/roles.",
      "2. Configura flujos de aprobacion por area y por monto (SC/OC/OS).",
      "3. Configura menu dinamico y permisos por rol.",
      "4. Administra maestros (items, proveedores, unidades), SMTP, plantillas, branding y correlativos.",
      "5. Consulta auditoria y registros de despliegue."],
     "Configuracion aplicada a los modulos correspondientes."),
    ("CU-14 Consultar dashboards, reportes y notificaciones", "Todos los roles",
     "Cada rol accede a dashboards con KPIs, reportes por modulo y notificaciones.",
     "Usuario autenticado.",
     ["1. Accede al dashboard de su rol (logistica, compras, almacen, finanzas, aprobaciones).",
      "2. Genera reportes (saldos, despachos, aprobaciones, compras, requerimientos).",
      "3. Recibe y marca notificaciones (stock, aprobaciones, saldos)."],
     "Informacion consolidada disponible para la gestion."),
]

for titulo, actor, desc, prec, flujo, post in casos:
    h3(titulo)
    add_table(["Campo", "Detalle"], [
        ("Actor principal", actor),
        ("Descripcion", desc),
        ("Precondicion", prec),
        ("Flujo principal", "\n".join(flujo)),
        ("Postcondicion", post),
    ], widths=[1.5, 5.7])
    doc.add_paragraph()

doc.add_page_break()

# ---------- 4. DESCRIPCION DE FLUJOS ----------
h1("4. Descripcion del flujo del sistema")

h2("4.1 Flujo de requerimiento y aprobacion por area")
par("1. El Solicitante registra el requerimiento (logistica_req / logistica_dreq) indicando "
    "tipo (COMPRA o CONSUMO), itemtipo, area, centro de costo, almacen y prioridad. Puede "
    "adjuntar documentos (LOGISTICA_RequerimientoAdjuntos) y, si trabaja sin conexion, el "
    "requerimiento queda local hasta sincronizar.")
par("2. Al registrarse, el motor de aprobaciones consulta LOGISTICA_FlujoAprobacionPorArea "
    "y LOGISTICA_UsuariosPorArea para determinar el jefe de area y los niveles; el documento "
    "queda en estado PENDIENTE y aparece en la bandeja del aprobador.")
par("3. El Jefe de Area aprueba o rechaza. Cada accion queda en LOGISTICA_LogAprobacionesPorArea. "
    "Si existen mas niveles, el requerimiento avanza; al completar todos queda APROBADO. El "
    "rechazo registra motivo y devuelve el estado ANULADO/RECHAZADO.")
par("4. Los cambios de estado se trazan en LOGISTICA_LogCambioEstadoRequerimiento.")

h2("4.2 Flujo de compras (consolidacion -> OC -> recepcion)")
par("1. Consolidacion: Logistica revisa LOGISTICA_ItemPendienteConsolidacion (detalles "
    "aprobados pendientes), agrupa por item y crea LOGISTICA_ConsolidacionCab con su detalle "
    "y trazabilidad en ConsolidacionOrigen.")
par("2. Cotizacion: desde la consolidacion se genera LOGISTICA_SolicitudCotizacion(+Det). "
    "Se registran las propuestas en LOGISTICA_Cotizacion(+Det), se evaluan y se selecciona "
    "el proveedor ganador.")
par("3. Solicitud de compra: se genera LOGISTICA_SolicitudCompra(+Det+Adjuntos) y se envia "
    "al flujo de aprobacion SC (LOGISTICA_FlujoAprobacionSC / AprobacionSC), multinivel por monto.")
par("4. Orden de compra: con la SC aprobada o desde el modulo de consolidacion "
    "(LOGISTICA_ConsolOC), se genera LOGISTICA_OrdenCompra(+Det). El flujo de aprobacion "
    "OC asigna niveles segun monto (LOGISTICA_FlujoAprobacionOC / UsuariosAprobadoresOC); "
    "cada decision se guarda en LOGISTICA_AprobacionOC y LOGISTICA_OrdenCompraLog.")
par("5. Sincronizacion ERP: aprobada la OC se genera el PDF (LOGISTICA_ConfigEmpresa / "
    "ConsolConfigEmpresa), se envia correo al proveedor (ConfigSmtp, LOGISTICA_ConsolCorreoLog) "
    "y se sincroniza al ERP Spring (WH_OrdenCompra via SP de sincronizacion; en caso de error "
    "queda en LOGISTICA_ColaMigracionSpring para reintento).")
par("6. Recepcion: Almacen lista las OCs pendientes, registra cantidades en "
    "LOGISTICA_RecepcionOC(+Det) con numero de nota de ingreso (numeroNI). El ingreso genera "
    "LOGISTICA_TransaccionInventario(+Detalle), actualiza LOGISTICA_Kardex y LOGISTICA_Stock, "
    "y replica al ERP (SPRING_*).")
par("7. Despacho: para requerimientos de consumo, Almacen despacha contra stock "
    "(LOGISTICA_Despachos). Si la cantidad despachada es menor a la solicitada, se crea "
    "LOGISTICA_SaldoPendienteAprobacion(+Det) que debe resolverse (aprobar saldo, cerrar o "
    "re-consolidar).")
par("8. Incrementos de OC: si la configuracion LOGISTICA_ConfigIncrementoOC lo permite, "
    "se pueden incrementar cantidades en la OC quedando log en LOGISTICA_OCIncrementoLog.")

h2("4.3 Flujo de servicios (SS -> OS -> conformidad)")
par("1. El Solicitante registra LOGISTICA_SolicitudServicio(+Det+Adjuntos).")
par("2. La SS pasa por su flujo de aprobacion (LOGISTICA_AprobacionSS).")
par("3. Logistica cotiza el servicio (LOGISTICA_CotizacionServicio) y genera "
    "LOGISTICA_OrdenServicio(+Det) o LOGISTICA_ConsolOS desde consolidacion.")
par("4. La OS pasa por aprobacion multinivel (LOGISTICA_FlujoAprobacionOS / AprobacionOS).")
par("5. Aprobada: se genera PDF, correo al proveedor y se sincroniza a WH_OrdenServicio "
    "en el ERP Spring.")
par("6. Ejecutado el servicio, el Usuario Conformidad registra la conformidad "
    "(LOGISTICA_ConformidadServicio / ConsolOSConformidad) con firma digital y acta, "
    "cerrando el ciclo. LOGISTICA_SeguimientoOrdenServicio traza las etapas.")

h2("4.4 Flujo de almacen e inventario")
bullet("Recepcion de OC: ingreso de mercaderia, nota de ingreso, kardex y stock.")
bullet("Despacho: salida de stock por atencion de requerimientos; genera saldos pendientes.")
bullet("Transferencias: LOGISTICA_RequerimientoTransferencia mueve stock entre almacenes "
       "(incluye almacen virtual); al aprobarse genera las transacciones correspondientes.")
bullet("Ajustes y conteos: LOGISTICA_AjusteInventario y LOGISTICA_ConteoInventario corrigen "
       "diferencias de stock.")
bullet("Devoluciones: LOGISTICA_DevolucionProveedor retorna mercaderia al proveedor; "
       "logistica_devolucion_consumo registra devoluciones internas; LOGISTICA_Reingreso "
       "reincorpora stock al almacen.")
bullet("Conformidad de almacen: LOGISTICA_ConformidadAlmacen documenta la recepcion conforme.")
bullet("Alertas: LOGISTICA_AlertaStock y NotificacionesStock avisan niveles criticos o "
       "nueva disponibilidad; el contador de notificaciones se refresca cada 30 segundos.")

h2("4.5 Flujo de administracion")
par("Los administradores gestionan la seguridad y la configuracion: usuarios y password "
    "(LOGISTICA_AdminUsuarios), usuarios por area y jefes (LOGISTICA_UsuariosPorArea), "
    "flujos de aprobacion por area/monto, menu dinamico por rol (LOGISTICA_ConfigMenu), "
    "permisos y montos maximos (LOGISTICA_ConfigPermisos), parametros, SMTP, plantillas de "
    "correo y PDF, branding por empresa, unidades de medida y correlativos. Toda la "
    "actividad queda en LOGISTICA_Auditoria y los despliegues en AppDeployments.")

h2("4.6 Integraciones")
add_table(["Integracion", "Direccion", "Descripcion"], [
    ("API Maestra", "Entrada", "Validacion de credenciales en el login de la aplicacion."),
    ("ERP Spring", "Salida", "OC/OS aprobadas se sincronizan a WH_OrdenCompra / WH_OrdenServicio; fallos van a la cola de migracion con reintentos."),
    ("ERP Spring", "Entrada", "Maestro de items (cache nocturno), stock y kardex espejo (SPRING_*)."),
    ("SMTP", "Salida", "Envio de solicitudes de cotizacion y ordenes a proveedores; log en CorreoLog / ConsolCorreoLog."),
    ("Tipo de cambio", "Entrada", "Servicio de tipo de cambio para documentos en USD."),
], widths=[1.5, 1.0, 4.7])

doc.add_paragraph()
par("Fin del documento.", italic=True, size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER

OUT_DOCX = os.path.join(os.path.dirname(BASE), "docs", "DOCUMENTACION_BD_CASOS_USO_FLUJOS.docx")
doc.save(OUT_DOCX)
print("DOCX generado:", OUT_DOCX)
