# -*- coding: utf-8 -*-
"""
Lee el Excel FORMATO_TECNICO_CONFIGURACION_LOGISTICA.xlsx y genera:
  1) doc_output/carga_masiva_config.json  -> JSON para el SP
  2) doc_output/CARGA_MASIVA.sql          -> EXEC LOGISTICA_cargaMasivaConfiguracion
     listo para pegar/ejecutar en SSMS.

Flujo:  llenar Excel tecnico -> ejecutar este script -> correr CARGA_MASIVA.sql.
"""
import os
import json
from openpyxl import load_workbook

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
EXCEL = os.path.join(BASE, "docs", "FORMATO_TECNICO_CONFIGURACION_LOGISTICA.xlsx")
OUT_DIR = os.path.join(BASE, "doc_output")


def filas(ws, ncols, inicio=4):
    """Devuelve filas no vacias desde la fila 'inicio' (headers en fila 3)."""
    out = []
    for r in range(inicio, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ncols + 1)]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        out.append(vals)
    return out


def txt(v):
    return None if v is None else str(v).strip()


def num(v):
    if v is None or str(v).strip() == "":
        return None
    return float(v) if "." in str(v) else int(float(v))


def bit(v, default=1):
    if v is None or str(v).strip() == "":
        return default
    return 1 if str(v).strip().upper() in ("1", "SI", "TRUE") else 0


def build_json(wb):
    data = {"usuarioCreacion": "CARGA_MASIVA_EXCEL"}

    # ---- AREAS (ruc, idarea, codigo, nombre, descripcion, mostrarAdmision, activo)
    ws = wb["AREAS"]
    data["areas"] = [
        {"ruc": txt(r[0]), "nombre": txt(r[3]), "descripcion": txt(r[4]),
         "mostrarAdmision": bit(r[5]), "activo": bit(r[6])}
        for r in filas(ws, 7) if txt(r[0]) and txt(r[3])
    ]

    # ---- CUENTAS_ACCESO (ruc, usuario, clave, nombre, email, idrol, rol, nrodoc, idempresa, activo)
    if "CUENTAS_ACCESO" in wb.sheetnames:
        ws = wb["CUENTAS_ACCESO"]
        data["usuariosAcceso"] = [
            {"ruc": txt(r[0]), "usuario": txt(r[1]), "clave": txt(r[2]),
             "nombre": txt(r[3]), "email": txt(r[4]), "idrol": txt(r[5]),
             "rol": txt(r[6]), "nrodocumento": txt(r[7]) or txt(r[1]),
             "idempresa": txt(r[8]), "activo": bit(r[9])}
            for r in filas(ws, 10) if txt(r[1])
        ]

    # ---- USUARIOS_POR_AREA (ruc,dni,nombre,idarea,idsubarea,rol,jefe,aprob,email,tel,activo)
    ws = wb["USUARIOS_POR_AREA"]
    data["usuariosPorArea"] = [
        {"ruc": txt(r[0]), "documentoidentidad": txt(r[1]),
         "nombreCompleto": txt(r[2]), "idarea": num(r[3]),
         "idsubarea": num(r[4]), "rol": txt(r[5]),
         "esJefeArea": bit(r[6], 0), "esAprobador": bit(r[7], 0),
         "email": txt(r[8]), "telefono": txt(r[9]), "activo": bit(r[10])}
        for r in filas(ws, 11) if txt(r[1])
    ]

    # ---- FLUJO_REQ_AREA (ruc,idarea,tipo,rolAprobador,activo)
    ws = wb["FLUJO_REQ_AREA"]
    data["flujoReq"] = [
        {"ruc": txt(r[0]), "idarea": num(r[1]), "tipoRequerimiento": txt(r[2]),
         "rolAprobador": txt(r[3]), "activo": bit(r[4])}
        for r in filas(ws, 5) if txt(r[0]) and txt(r[2])
    ]

    # ---- FLUJO_OC (min,max,nivel,rolAprobador,nombreRol,descripcion,activo)
    ws = wb["FLUJO_OC"]
    data["flujoOC"] = [
        {"montoMinimo": num(r[0]), "montoMaximo": num(r[1]), "nivel": num(r[2]) or 1,
         "rolAprobador": txt(r[3]), "nombreRol": txt(r[4]),
         "descripcion": txt(r[5]), "activo": bit(r[6])}
        for r in filas(ws, 7) if r[0] is not None
    ]

    # ---- APROBADORES_OC (dni,nombre,email,rol,min,max,nivel,activo)
    ws = wb["APROBADORES_OC"]
    data["aprobadoresOC"] = [
        {"documentoIdentidad": txt(r[0]), "nombreCompleto": txt(r[1]),
         "email": txt(r[2]), "rol": txt(r[3]), "montoMinimo": num(r[4]),
         "montoMaximo": num(r[5]), "nivel": num(r[6]) or 1, "activo": bit(r[7])}
        for r in filas(ws, 8) if txt(r[0])
    ]

    # ---- FLUJO_OS (nombre,desde,hasta,niveles,activo)
    ws = wb["FLUJO_OS"]
    data["flujoOS"] = [
        {"nombre": txt(r[0]), "montoDesde": num(r[1]), "montoHasta": num(r[2]),
         "niveles": num(r[3]) or 1, "activo": bit(r[4])}
        for r in filas(ws, 5) if txt(r[0])
    ]

    # ---- APROBADORES_OS (dni,nombre,rol,nivelMaximo,activo)
    ws = wb["APROBADORES_OS"]
    data["aprobadoresOS"] = [
        {"documentoIdentidad": txt(r[0]), "nombre": txt(r[1]), "rol": txt(r[2]),
         "nivelMaximo": num(r[3]) or 1, "activo": bit(r[4])}
        for r in filas(ws, 5) if txt(r[0])
    ]

    return data


def main():
    wb = load_workbook(EXCEL)
    data = build_json(wb)

    os.makedirs(OUT_DIR, exist_ok=True)
    json_path = os.path.join(OUT_DIR, "carga_masiva_config.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    # SQL wrapper: el JSON se escapa duplicando comillas simples
    json_sql = json.dumps(data, ensure_ascii=False).replace("'", "''")
    sql_path = os.path.join(OUT_DIR, "CARGA_MASIVA.sql")
    with open(sql_path, "w", encoding="utf-8") as f:
        f.write("-- Generado desde FORMATO_TECNICO_CONFIGURACION_LOGISTICA.xlsx\n")
        f.write("-- Requiere SP: LOGISTICA_cargaMasivaConfiguracion\n")
        f.write("--   (archivo: api_logistica/SQL/750_SP_CARGA_MASIVA_CONFIGURACION.sql)\n")
        f.write("USE [LOGISTICA]\nGO\n\n")
        f.write(f"EXEC [dbo].[LOGISTICA_cargaMasivaConfiguracion] @json = N'{json_sql}';\nGO\n\n")
        f.write("-- Verificacion\n")
        f.write("SELECT COUNT(*) AS usuariosPorArea FROM LOGISTICA_UsuariosPorArea WHERE activo=1;\n")
        f.write("SELECT COUNT(*) AS flujoArea FROM LOGISTICA_FlujoAprobacionPorArea WHERE activo=1;\n")
        f.write("SELECT COUNT(*) AS aprobadoresOC FROM LOGISTICA_UsuariosAprobadoresOC WHERE activo=1;\n")
        f.write("SELECT COUNT(*) AS aprobadoresOS FROM LOGISTICA_UsuarioAprobadorOS WHERE activo=1;\n")

    # resumen
    print(f"JSON: {json_path}")
    print(f"SQL : {sql_path}")
    for k, v in data.items():
        if isinstance(v, list):
            print(f"  {k}: {len(v)} filas")


if __name__ == "__main__":
    main()
