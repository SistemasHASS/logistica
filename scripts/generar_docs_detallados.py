# -*- coding: utf-8 -*-
"""
Genera 3 documentos Word separados y detallados:
 1. DOCUMENTACION_BASE_DE_DATOS.docx   - diccionario de datos completo + DER
 2. DOCUMENTACION_CASOS_DE_USO.docx    - actores + especificacion de CUs
 3. DOCUMENTACION_FLUJOS_PROCESO.docx  - flujos detallados + diagramas de estado
"""
import os, json
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Ellipse, FancyArrowPatch, Circle
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, "doc_output")
DOCS_DIR = os.path.join(os.path.dirname(BASE), "docs")
os.makedirs(OUT_DIR, exist_ok=True)
SCHEMA = json.load(open(os.path.join(OUT_DIR, "schema.json"), encoding="utf-8"))

AZUL = RGBColor(0x1F, 0x4E, 0x79)

# =====================================================================
# Tablas manuales (preexistentes / del ERP, no creadas en scripts)
# =====================================================================
def T(cols, fks=None):
    return {"cols": [{"name": c[0], "type": c[1], "flags": c[2] if len(c) > 2 else "",
                      "note": c[3] if len(c) > 3 else ""} for c in cols],
            "fks": fks or [], "src": "tabla preexistente"}

SCHEMA["logistica_req"] = T([
    ("id", "INT IDENTITY", "PK ID"),
    ("idrequerimiento", "INT", "", "Identificador del requerimiento"),
    ("RequisicionNumero", "VARCHAR(50)", "", "Numero de requisicion"),
    ("ruc", "VARCHAR(20)", "", "Empresa (RUC) - multiempresa"),
    ("idfundo", "VARCHAR(50)", "", "Fundo / sede"),
    ("idarea", "VARCHAR(20)", "", "Area solicitante"),
    ("idclasificacion", "VARCHAR(50)", "", "Clasificacion del requerimiento"),
    ("prioridad", "VARCHAR(20)", "", "NORMAL / URGENTE / CRITICA"),
    ("idalmacen", "VARCHAR(50)", "", "Almacen origen"),
    ("idalmacen_destino", "VARCHAR(50)", "", "Almacen destino (transferencias)"),
    ("servicio", "VARCHAR(100)", "", "Servicio asociado"),
    ("glosa", "NVARCHAR(500)", "", "Observaciones / glosa"),
    ("dniregistra", "VARCHAR(20)", "", "DNI del usuario que registra"),
    ("fecharegistro", "DATETIME", "", "Fecha de registro"),
    ("fechaAprobacion", "DATETIME", "NULL", "Fecha de aprobacion"),
    ("tipo", "VARCHAR(20)", "", "COMPRA / CONSUMO / SERVICIO"),
    ("itemtipo", "VARCHAR(30)", "", "ITEM / COMMODITY / ACTIVO_FIJO / ACTIVO_MENOR"),
    ("estado", "INT", "", "Estado numerico"),
    ("estados", "VARCHAR(20)", "", "PENDIENTE / APROBADO / ANULADO / ATENDIDO / COMPLETADO"),
    ("eliminado", "BIT", "", "Borrado logico"),
], fks=[])
SCHEMA["logistica_dreq"] = T([
    ("id", "INT IDENTITY", "PK ID"),
    ("idrequerimiento", "INT", "FK", "-> logistica_req.idrequerimiento"),
    ("codigo", "VARCHAR(50)", "", "Codigo del item"),
    ("idproducto", "VARCHAR(50)", "", "Id del producto"),
    ("iddescripcion", "NVARCHAR(300)", "", "Descripcion"),
    ("tipoclasificacion", "VARCHAR(50)", "", "Tipo de clasificacion"),
    ("cantidad", "DECIMAL(18,4)", "", "Cantidad solicitada"),
    ("idproyecto", "VARCHAR(100)", "", "Proyecto"),
    ("idcentrocosto", "VARCHAR(100)", "", "Centro de costo"),
    ("idturno", "VARCHAR(100)", "", "Turno"),
    ("idlabor", "VARCHAR(100)", "", "Labor"),
    ("eliminado", "BIT", "", "Borrado logico"),
], fks=[{"cols": "idrequerimiento", "ref": "logistica_req", "refcols": "idrequerimiento"}])
SCHEMA["WH_OrdenCompra"] = T([
    ("idOrden", "INT", "PK", "Identificador OC en ERP"),
    ("numeroOrden", "VARCHAR(50)", "", "Numero de orden"),
    ("proveedor / rucProveedor", "VARCHAR", "", "Proveedor"),
    ("moneda / totalOrden", "DECIMAL", "", "Importes"),
    ("estado", "VARCHAR(30)", "", "Estado en ERP"),
], fks=[])

# =====================================================================
# MODULOS (orden y asignacion de tablas)
# =====================================================================
def modulo(t):
    n = t.lower()
    if n.startswith(("wh_", "spring_", "logistica_colamigracion",
                     "logistica_sincronizacion", "logistica_maestroitem")):
        return "Integracion ERP Spring"
    if n.startswith("logistica_consol"):
        return "Consolidacion OC/OS"
    if "servicio" in n or n.endswith(("aprobacionos", "flujoaprobacionos",
        "usuarioaprobadoros", "osconformidad", "aprobacionss")) \
        or "confirmacionservicio" in n or "conformidadservicio" in n \
        or "historialprecioservicio" in n or "plantillaconformidadservicio" in n:
        return "Servicios"
    if "consolidacion" in n or "itempendienteconsolidacion" in n \
        or "consolidado" in n or n == "reqconsolidado":
        return "Consolidacion de Requerimientos"
    if "aprobacion" in n or "flujo" in n or "aprobadores" in n \
        or "saldopendiente" in n or "adjuntosorden" in n \
        or n in ("logistica_areas", "logistica_usuariosporarea", "logistica_usuariosporarea"):
        return "Aprobaciones"
    if n in ("logistica_areas",) or "porarea" in n:
        return "Aprobaciones"
    if n in ("logistica_req", "logistica_dreq") or "reqcommodity" in n \
        or "solicitudcreacionitem" in n or "historialanulaciones" in n \
        or "logcambioestadorequerimiento" in n or n == "logistica_requerimientos" \
        or n == "logistica_requerimientosporarea":
        return "Requerimientos"
    if "requerimientotransferencia" in n:
        return "Almacen e Inventario"
    if "recepcion" in n or "despacho" in n or "kardex" in n or "stock" in n \
        or "inventario" in n or "transaccion" in n or "devolucion" in n \
        or "reingreso" in n or "almacen" in n or n == "logistica_item" \
        or "conformidadalmacen" in n or "plantillaalmacen" in n:
        return "Almacen e Inventario"
    if "cotizacion" in n or "solicitudcompra" in n or "ordencompra" in n \
        or "ocadjunto" in n or "seguimientoordencompra" in n \
        or "incrementooc" in n or "proveedor" in n or n == "logistica_entrega":
        return "Compras"
    if "notificacion" in n or "correolog" in n or "alerta" in n:
        return "Notificaciones"
    return "Administracion y Configuracion"

MODULOS_ORDEN = [
    "Requerimientos", "Aprobaciones", "Consolidacion de Requerimientos",
    "Compras", "Consolidacion OC/OS", "Servicios", "Almacen e Inventario",
    "Notificaciones", "Administracion y Configuracion", "Integracion ERP Spring",
]

MODULO_TABLAS = {m: [] for m in MODULOS_ORDEN}
for t in SCHEMA:
    MODULO_TABLAS[modulo(t)].append(t)
for m in MODULO_TABLAS:
    MODULO_TABLAS[m].sort(key=lambda x: x.lower())

DESCRIPCIONES_TABLA = {
    "logistica_req": "Cabecera de los requerimientos de compra, consumo o servicio.",
    "logistica_dreq": "Detalle de items solicitados en cada requerimiento.",
    "LOGISTICA_Areas": "Areas/subareas de la empresa para flujos de aprobacion.",
    "LOGISTICA_UsuariosPorArea": "Asignacion de usuarios a areas; identifica jefes de area (esJefeArea).",
    "LOGISTICA_FlujoAprobacionPorArea": "Configuracion de niveles de aprobacion por area.",
    "LOGISTICA_LogAprobacionesPorArea": "Historial de aprobaciones/rechazos por nivel.",
    "LOGISTICA_FlujoAprobacion": "Flujo de aprobacion legacy.",
    "LOGISTICA_LogAprobaciones": "Log de aprobaciones legacy.",
    "LOGISTICA_SaldoPendienteAprobacion": "Cabecera de saldos pendientes generados por despachos parciales.",
    "LOGISTICA_SaldoPendienteAprobacionDet": "Detalle por item del saldo pendiente.",
    "LOGISTICA_SaldoPendienteItems": "Items pendientes de saldo (version alterna).",
    "LOGISTICA_FlujoAprobacionSC": "Niveles de aprobacion de solicitudes de compra por monto.",
    "LOGISTICA_AprobacionSC": "Registro de aprobaciones de solicitudes de compra.",
    "LOGISTICA_UsuariosAprobadoresSC": "Usuarios habilitados para aprobar SC.",
    "LOGISTICA_LogAprobacionesSC": "Historial de aprobaciones de SC.",
    "LOGISTICA_FlujoAprobacionOC": "Niveles de aprobacion de OC segun rango de monto.",
    "LOGISTICA_AprobacionOC": "Registro de aprobacion por nivel de cada OC.",
    "LOGISTICA_UsuariosAprobadoresOC": "Usuarios aprobadores de OC con rango de monto.",
    "LOGISTICA_LogAprobacionesOC": "Historial de acciones de aprobacion OC.",
    "LOGISTICA_AprobacionOrden": "Aprobacion generica de ordenes (OC/OS).",
    "LOGISTICA_AprobacionOrdenDetalle": "Detalle de la aprobacion generica de ordenes.",
    "LOGISTICA_AdjuntosOrden": "Adjuntos asociados a aprobaciones de ordenes.",
    "FlujosAprobacionOC": "Flujo de aprobacion de OC (esquema legacy).",
    "UsuariosAprobadoresOC": "Aprobadores de OC (esquema legacy).",
    "logistica_aprobadores": "Aprobadores por nivel (esquema legacy).",
    "logistica_flujos": "Flujos de aprobacion (esquema legacy).",
    "logistica_areas": "Areas (esquema legacy).",
    "LOGISTICA_Requerimientos": "Tabla auxiliar de requerimientos.",
    "LOGISTICA_RequerimientosPorArea": "Asociacion requerimiento-area.",
    "LOGISTICA_RequerimientoAdjuntos": "Adjuntos de requerimientos (base64).",
    "LOGISTICA_RequerimientoTransferencia": "Requerimientos de transferencia entre almacenes.",
    "LOGISTICA_ReqCommodity": "Requerimientos de commodities (cabecera).",
    "LOGISTICA_DReqCommodity": "Detalle de requerimientos de commodities.",
    "LOGISTICA_SolicitudCreacionItem": "Solicitudes de creacion de nuevos items en el maestro.",
    "LOGISTICA_HistorialAnulaciones": "Historial de anulaciones de requerimientos.",
    "LOGISTICA_LogCambioEstadoRequerimiento": "Trazabilidad de cambios de estado del requerimiento.",
    "LOGISTICA_ConsolidacionCab": "Cabecera de consolidaciones de requerimientos.",
    "LOGISTICA_ConsolidacionDet": "Lineas consolidadas por item.",
    "LOGISTICA_ConsolidacionOrigen": "Trazabilidad linea origen -> linea consolidada.",
    "LOGISTICA_ItemPendienteConsolidacion": "Control de detalles pendientes/consolidados.",
    "RequerimientoConsolidadoCompra": "Cabecera de consolidacion (esquema alterno).",
    "DetalleRequerimientoConsolidado": "Detalle consolidado (esquema alterno).",
    "ReqConsolidado": "Trazabilidad de consolidacion (esquema alterno).",
    "LOGISTICA_SolicitudCotizacion": "Solicitudes de cotizacion generadas desde consolidacion.",
    "LOGISTICA_SolicitudCotizacionDet": "Detalle de items de la solicitud de cotizacion.",
    "LOGISTICA_Cotizacion": "Cotizaciones de proveedores.",
    "LOGISTICA_CotizacionDet": "Detalle de precios por item de cada cotizacion.",
    "LOGISTICA_CotizacionOC": "Cotizacion asociada a OC (multiempresa).",
    "LOGISTICA_CotizacionOCDetalle": "Detalle de cotizacion OC.",
    "LOGISTICA_solicitud_cotizacion": "Solicitud de cotizacion (esquema snake_case).",
    "LOGISTICA_solicitud_cotizacion_detalle": "Detalle de solicitud de cotizacion (snake_case).",
    "LOGISTICA_cotizacion": "Cotizacion (esquema snake_case).",
    "LOGISTICA_cotizacion_detalle": "Detalle de cotizacion (snake_case).",
    "SolicitudCotizacion": "Solicitud de cotizacion (esquema alterno).",
    "DetalleSolicitudCotizacion": "Detalle de solicitud de cotizacion (esquema alterno).",
    "CotizacionProveedor": "Cotizaciones recibidas de proveedores (esquema alterno).",
    "LOGISTICA_SolicitudCompra": "Solicitudes de compra (consolidada/directa/urgente).",
    "LOGISTICA_SolicitudCompraDet": "Detalle de la solicitud de compra.",
    "LOGISTICA_SolicitudCompraAdjuntos": "Adjuntos de la solicitud de compra.",
    "LOGISTICA_OrdenCompra": "Ordenes de compra.",
    "LOGISTICA_OrdenCompraDet": "Detalle de items de la OC.",
    "LOGISTICA_OrdenCompraLog": "Historial de cambios de estado de la OC.",
    "LOGISTICA_OCAdjunto": "Adjuntos de la OC.",
    "LOGISTICA_SeguimientoOrdenCompra": "Seguimiento de etapas de la OC.",
    "LOGISTICA_ConfigIncrementoOC": "Configuracion de incremento de cantidades en OC.",
    "LOGISTICA_ConfigIncrementoOCAuditoria": "Auditoria de cambios de configuracion de incremento.",
    "LOGISTICA_OCIncrementoLog": "Log de incrementos de cantidad aplicados a OC.",
    "LOGISTICA_Proveedor": "Maestro de proveedores.",
    "LOGISTICA_CategoriaProveedor": "Categorias de proveedores.",
    "ProveedoresEmail": "Correos de contacto de proveedores.",
    "LOGISTICA_Entrega": "Entregas programadas de proveedores.",
    "LOGISTICA_ConsolOC": "OC generadas desde el modulo de consolidacion.",
    "LOGISTICA_ConsolOCDet": "Detalle de items de la OC de consolidacion.",
    "LOGISTICA_ConsolOCDist": "Distribucion de la OC entre requerimientos origen.",
    "LOGISTICA_ConsolOS": "OS generadas desde el modulo de consolidacion.",
    "LOGISTICA_ConsolOSDet": "Detalle de la OS de consolidacion.",
    "LOGISTICA_ConsolOSConformidad": "Conformidad de OS de consolidacion (firma y acta).",
    "LOGISTICA_ConsolAdjunto": "Adjuntos de OC y OS de consolidacion.",
    "LOGISTICA_ConsolConfigEmpresa": "Datos de empresa para PDF/correo de OC-OS.",
    "LOGISTICA_ConsolCorreoLog": "Log de correos enviados desde consolidacion.",
    "LOGISTICA_SolicitudServicio": "Solicitudes de servicio (cabecera).",
    "LOGISTICA_SolicitudServicioDet": "Detalle de la solicitud de servicio.",
    "LOGISTICA_SolicitudServicioAdjuntos": "Adjuntos de la solicitud de servicio.",
    "LOGISTICA_AprobacionSS": "Aprobaciones de solicitudes de servicio.",
    "LOGISTICA_CotizacionServicio": "Cotizaciones de servicios por proveedor.",
    "LOGISTICA_CotizacionServicioDet": "Detalle de cotizacion de servicio.",
    "LOGISTICA_OrdenServicio": "Ordenes de servicio.",
    "LOGISTICA_OrdenServicioDet": "Detalle de la orden de servicio.",
    "LOGISTICA_OrdenServicioDetalle": "Detalle de OS (estructura extendida).",
    "LOGISTICA_FlujoAprobacionOS": "Niveles de aprobacion de OS.",
    "LOGISTICA_AprobacionOS": "Aprobaciones de ordenes de servicio.",
    "LOGISTICA_UsuarioAprobadorOS": "Usuarios aprobadores de OS.",
    "LOGISTICA_LogAprobacionesOS": "Historial de aprobaciones OS.",
    "LOGISTICA_ConformidadServicio": "Conformidad de servicio ejecutado (firma digital).",
    "LOGISTICA_ConfirmacionServicioRealizado": "Confirmacion de realizacion del servicio.",
    "LOGISTICA_OSConformidad": "Conformidad de OS (estructura alterna).",
    "LOGISTICA_SeguimientoOrdenServicio": "Seguimiento de etapas de la OS.",
    "LOGISTICA_HistorialPrecioServicio": "Historial de precios de servicios.",
    "LOGISTICA_PlantillaConformidadServicio": "Plantilla del acta de conformidad.",
    "LOGISTICA_RecepcionOC": "Recepciones de mercaderia contra OC.",
    "LOGISTICA_RecepcionOCDet": "Detalle de items recibidos.",
    "LOGISTICA_RecepcionOCLog": "Log de la recepcion.",
    "LOGISTICA_RecepcionOrdenCompra": "Recepcion de OC (esquema alterno).",
    "LOGISTICA_DetalleRecepcionOrdenCompra": "Detalle de recepcion (esquema alterno).",
    "LOGISTICA_Despachos": "Despachos de atencion de requerimientos.",
    "LOGISTICA_DetalleDespacho": "Detalle de items despachados.",
    "LOGISTICA_Item": "Maestro de items.",
    "LOGISTICA_Almacen": "Maestro de almacenes.",
    "LOGISTICA_TransaccionInventario": "Transacciones de ingreso/salida de inventario.",
    "LOGISTICA_DetalleTransaccionInventario": "Detalle de las transacciones de inventario.",
    "LOGISTICA_Kardex": "Kardex valorizado por item y almacen.",
    "LOGISTICA_Stock": "Stock actual por item y almacen.",
    "LOGISTICA_ListaStock": "Listas de stock configuradas.",
    "LOGISTICA_ListaStockDet": "Detalle de listas de stock.",
    "LOGISTICA_AlertaStock": "Alertas de stock minimo / disponibilidad.",
    "LOGISTICA_AjusteInventario": "Ajustes de inventario (cabecera).",
    "LOGISTICA_AjusteInventarioDet": "Detalle de ajustes de inventario.",
    "LOGISTICA_ConteoInventario": "Conteos fisicos (cabecera).",
    "LOGISTICA_ConteoInventarioDet": "Detalle de conteos fisicos.",
    "LOGISTICA_DevolucionProveedor": "Devoluciones de mercaderia al proveedor.",
    "LOGISTICA_DetalleDevolucionProveedor": "Detalle de devoluciones a proveedor.",
    "logistica_devolucion_consumo": "Devoluciones internas de consumo.",
    "logistica_devolucion_consumo_det": "Detalle de devoluciones de consumo.",
    "LOGISTICA_Reingreso": "Reingresos de material al almacen.",
    "LOGISTICA_DetalleReingreso": "Detalle de reingresos.",
    "LOGISTICA_ConformidadAlmacen": "Conformidad de recepcion de almacen.",
    "LOGISTICA_PlantillaAlmacen": "Plantilla de documentos de almacen.",
    "logistica_notificaciones": "Notificaciones generales a usuarios.",
    "LOGISTICA_NotificacionesStock": "Notificaciones de disponibilidad de stock.",
    "LOGISTICA_CorreoLog": "Log de correos enviados.",
    "LOGISTICA_AdminUsuarios": "Usuarios administradores del modulo logistica.",
    "LOGISTICA_Usuario": "Usuarios registrados (dashboard logistica).",
    "LOGISTICA_Usuarios": "Usuarios (esquema de seguridad).",
    "LOGISTICA_Roles": "Roles del sistema.",
    "LOGISTICA_Perfiles": "Perfiles de usuario.",
    "LOGISTICA_UsuariosRoles": "Asignacion usuario-rol.",
    "LOGISTICA_PerfilesUsuarios": "Asignacion usuario-perfil.",
    "logistica_usuarios": "Usuarios (esquema legacy).",
    "logistica_roles": "Roles (esquema legacy).",
    "logistica_perfiles": "Perfiles (esquema legacy).",
    "logistica_usuario_roles": "Usuario-rol (esquema legacy).",
    "logistica_usuario_roles_adicionales": "Roles adicionales por usuario.",
    "logistica_perfil_roles": "Perfil-roles (esquema legacy).",
    "logistica_usuario_config": "Configuracion por usuario.",
    "logistica_rol_config": "Configuracion por rol.",
    "LOGISTICA_ConfigMenu": "Menu dinamico por rol/empresa.",
    "LOGISTICA_ConfigPermisos": "Permisos por rol (montos, autoaprobacion).",
    "LOGISTICA_ConfiguracionPermiso": "Configuracion de permisos detallada.",
    "LOGISTICA_Parametros": "Parametros generales del sistema.",
    "LOGISTICA_ConfigPDF": "Configuracion de plantillas PDF.",
    "LOGISTICA_Auditoria": "Auditoria de acciones del sistema.",
    "LOGISTICA_ConfigBranding": "Branding por empresa (logo, colores).",
    "LOGISTICA_ConfigEmpresa": "Datos de empresa para documentos.",
    "ConfigSmtp": "Configuracion del servidor SMTP.",
    "PlantillasCorreo": "Plantillas de correo electronico.",
    "PlantillasPdf": "Plantillas de documentos PDF.",
    "LOGISTICA_UnidadMedida": "Maestro de unidades de medida.",
    "LOGISTICA_Correlativo": "Correlativos de numeracion de documentos.",
    "AppDeployments": "Registro de versiones desplegadas de la app.",
    "WH_OrdenServicio": "Ordenes de servicio en el ERP Spring (espejo).",
    "WH_OrdenServicioDetalle": "Detalle de OS en el ERP Spring.",
    "WH_OrdenCompra": "Ordenes de compra en el ERP Spring (destino de sincronizacion).",
    "SPRING_Stock": "Stock espejo del ERP Spring.",
    "SPRING_Kardex": "Kardex espejo del ERP Spring.",
    "SPRING_Transacciones": "Transacciones espejo del ERP Spring.",
    "LOGISTICA_ColaMigracionSpring": "Cola de sincronizacion con ERP (reintentos).",
    "LOGISTICA_SincronizacionLog": "Log de sincronizaciones.",
    "LOGISTICA_SincronizacionSpringLog": "Log de sincronizacion Spring.",
    "LOGISTICA_MaestroItemCache": "Cache del maestro de items del ERP.",
    "LOGISTICA_MaestroItemSync": "Sincronizacion del maestro de items (job nocturno).",
}

# =====================================================================
# HELPERS DOCX
# =====================================================================
def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10)
    for hname, sz in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11), ("Heading 4", 10)):
        try:
            s = d.styles[hname]; s.font.color.rgb = AZUL; s.font.size = Pt(sz)
        except Exception:
            pass
    return d


def portada(doc, titulo, subtitulo):
    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SISTEMA DE LOGISTICA"); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = AZUL
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo); r.font.size = Pt(20); r.font.color.rgb = AZUL
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitulo); r.font.size = Pt(11); r.italic = True
    for _ in range(9):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("api_logistica (.NET)  |  logistica (Angular)  |  SQL Server - BD LOGISTICA")
    r.font.size = Pt(10)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fecha: 09/10/2026"); r.font.size = Pt(10)
    doc.add_page_break()


def shade_cell(cell, color="1F4E79"):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, headers, rows, widths=None, font_size=8.5, hdr_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True; r.font.size = Pt(hdr_size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(c)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(font_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


def par(doc, text, bold=False, italic=False, size=10, center=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def bullet(doc, text):
    return doc.add_paragraph(text, style="List Bullet")


def figura(doc, path, caption, w=6.9):
    doc.add_picture(path, width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par(doc, caption, italic=True, size=9, center=True)


# =====================================================================
# HELPERS DE DIAGRAMAS
# =====================================================================
def draw_entity(ax, x, y, w, title, cols, color="#dbe9f6", edge="#2e5f8a",
                title_fs=8, col_fs=6.5):
    row_h = 0.30; title_h = 0.34
    h = title_h + row_h * len(cols)
    ax.add_patch(FancyBboxPatch((x, y - h), w, h,
                 boxstyle="round,pad=0.02,rounding_size=0.05",
                 fc="white", ec=edge, lw=1.0, zorder=2))
    ax.add_patch(FancyBboxPatch((x, y - title_h), w, title_h,
                 boxstyle="round,pad=0.02,rounding_size=0.05",
                 fc=color, ec=edge, lw=1.0, zorder=3))
    ax.text(x + w / 2, y - title_h / 2, title, ha="center", va="center",
            fontsize=title_fs, fontweight="bold", zorder=4)
    for i, c in enumerate(cols):
        ax.text(x + 0.12, y - title_h - row_h * (i + 0.5), c, ha="left",
                va="center", fontsize=col_fs, family="monospace", zorder=4)
    return {"x": x, "y": y, "w": w, "h": h, "cx": x + w / 2, "cy": y - h / 2}


def arrow(ax, p1, p2, label="", rad=0.0, color="#555555"):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=9,
                 lw=1.0, color=color, connectionstyle=f"arc3,rad={rad}", zorder=1))
    if label:
        mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
        ax.text(mx, my + 0.12, label, fontsize=6, ha="center", color=color,
                style="italic", zorder=5)


def right(b):  return (b["x"] + b["w"], b["cy"])
def left(b):   return (b["x"], b["cy"])
def top(b):    return (b["cx"], b["y"])
def bottom(b): return (b["cx"], b["y"] - b["h"])


def state_diagram(filename, titulo, estados, transiciones, layout=None):
    """estados: lista [(id, etiqueta, color)]
       transiciones: [(origen, destino, etiqueta)]"""
    n = len(estados)
    fig, ax = plt.subplots(figsize=(max(10, n * 2.2), 3.4))
    ax.set_xlim(0, n * 2.2 + 1); ax.set_ylim(0, 4); ax.axis("off")
    ax.set_title(titulo, fontsize=11, fontweight="bold")
    pos = {}
    if layout is None:
        layout = {e[0]: (1.2 + i * 2.2, 2.0) for i, e in enumerate(estados)}
    for sid, etq, color in estados:
        x, y = layout[sid]
        ax.add_patch(FancyBboxPatch((x - 0.85, y - 0.32), 1.7, 0.64,
                     boxstyle="round,pad=0.05,rounding_size=0.3",
                     fc=color, ec="#333333", lw=1.2, zorder=3))
        ax.text(x, y, etq, ha="center", va="center", fontsize=8.5,
                fontweight="bold", zorder=4)
        pos[sid] = (x, y)
    for ti, (a, b, lbl) in enumerate(transiciones):
        x1, y1 = pos[a]; x2, y2 = pos[b]
        if a == b:
            continue
        if y1 == y2 and abs(x2 - x1) > 2.6:
            rad = -0.25 if x2 > x1 else 0.25   # salto: curva por arriba
        elif y1 == y2 and x2 < x1:
            rad = 0.25
        else:
            rad = 0.0
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                     mutation_scale=13, lw=1.3, color="#444444",
                     connectionstyle=f"arc3,rad={rad}", zorder=2,
                     shrinkA=28, shrinkB=28))
        if abs(x2 - x1) > 2.6 and y1 == y2:
            dy = 0.95
        elif rad:
            dy = 0.75
        else:
            dy = 0.55 if ti % 2 == 0 else -0.55
        ax.text((x1 + x2) / 2, (y1 + y2) / 2 + dy,
                lbl, ha="center", fontsize=7.5, color="#555555", style="italic")
    fig.savefig(os.path.join(OUT_DIR, filename), dpi=140, bbox_inches="tight")
    plt.close(fig)


# =====================================================================
# FIGURAS DE ESTADO
# =====================================================================
C_OK = "#d9ead3"; C_PEND = "#fff2cc"; C_BAD = "#f4cccc"; C_PROC = "#d0e0f0"; C_FIN = "#cfe2f3"

state_diagram("est_req.png", "Estados del Requerimiento",
    [("bor", "REGISTRADO", C_PROC), ("pen", "PENDIENTE", C_PEND),
     ("eval", "EN EVALUACION", C_PEND), ("apr", "APROBADO", C_OK),
     ("anu", "ANULADO", C_BAD), ("ate", "ATENDIDO", C_FIN), ("comp", "COMPLETADO", C_FIN)],
    [("bor", "pen", "enviar"), ("pen", "eval", "revision jefatura"),
     ("eval", "apr", "aprueba"), ("eval", "anu", "rechaza"),
     ("apr", "ate", "despacho / compra"), ("ate", "comp", "cierre")])

state_diagram("est_sc.png", "Estados de la Solicitud de Compra",
    [("gen", "GENERADA", C_PROC), ("env", "ENVIADA", C_PEND),
     ("apr", "APROBADA", C_OK), ("rec", "RECHAZADA", C_BAD), ("cot", "EN_COTIZACION", C_PROC)],
    [("gen", "env", "enviar a aprobacion"), ("env", "apr", "aprueba niveles"),
     ("env", "rec", "rechaza"), ("apr", "cot", "genera OC")])

state_diagram("est_oc.png", "Estados de la Orden de Compra",
    [("pen", "PENDIENTE", C_PEND), ("pa", "PENDIENTE_APROBACION", C_PEND),
     ("apr", "APROBADA", C_OK), ("env", "ENVIADA/SYNC", C_PROC),
     ("rp", "RECIBIDA_PARCIAL", C_PROC), ("rt", "RECIBIDA_TOTAL", C_FIN),
     ("anu", "ANULADA", C_BAD)],
    [("pen", "pa", "enviar aprobacion"), ("pa", "apr", "aprueba todos los niveles"),
     ("pa", "anu", "rechaza"), ("apr", "env", "sincroniza ERP"),
     ("env", "rp", "recepcion parcial"), ("env", "rt", "recepcion total"),
     ("rp", "rt", "completa recepcion")])

state_diagram("est_os.png", "Estados de la Orden de Servicio",
    [("pen", "PENDIENTE", C_PEND), ("pa", "PENDIENTE_APROBACION", C_PEND),
     ("apr", "APROBADA", C_OK), ("env", "ENVIADA", C_PROC),
     ("eje", "EN_EJECUCION", C_PROC), ("conf", "CONFORME", C_FIN),
     ("noc", "NO_CONFORME", C_BAD)],
    [("pen", "pa", "enviar aprobacion"), ("pa", "apr", "aprueba"),
     ("pa", "noc", "rechaza"), ("apr", "env", "envio a proveedor"),
     ("env", "eje", "ejecucion"), ("eje", "conf", "conformidad"),
     ("eje", "noc", "no conforme")])

state_diagram("est_cons.png", "Estados de la Consolidacion",
    [("con", "CONSOLIDADO", C_PROC), ("pc", "PENDIENTE_COTIZACION", C_PEND),
     ("ec", "EN_COTIZACION", C_PEND), ("cot", "COTIZADO", C_OK),
     ("og", "ORDEN_GENERADA", C_FIN), ("anu", "ANULADO", C_BAD)],
    [("con", "pc", "registra"), ("pc", "ec", "genera solicitud cotizacion"),
     ("ec", "cot", "selecciona proveedor"), ("cot", "og", "genera OC/OS"),
     ("con", "anu", "anula"), ("pc", "anu", "anula")])

state_diagram("est_desp.png", "Estados del Despacho y Saldo Pendiente",
    [("pen", "PENDIENTE DESPACHO", C_PEND), ("par", "ATENDIDO PARCIAL", C_PROC),
     ("tot", "ATENDIDO TOTAL", C_FIN), ("sal", "SALDO PENDIENTE", C_PEND),
     ("sapr", "SALDO APROBADO", C_OK), ("scer", "SALDO CERRADO", C_FIN)],
    [("pen", "tot", "stock completo"), ("pen", "par", "stock parcial"),
     ("par", "sal", "genera saldo"), ("sal", "sapr", "aprueba jefatura"),
     ("sal", "scer", "cierra saldo"), ("sapr", "tot", "atencion posterior")])

print("figuras de estado OK")

# =====================================================================
# FIGURAS DER (reutilizar las generadas si existen; regenerar simplificado)
# =====================================================================
# Reutilizamos der_principal.png y der_servicios_admin.png ya generados.
DER1 = os.path.join(OUT_DIR, "der_principal.png")
DER2 = os.path.join(OUT_DIR, "der_servicios_admin.png")
CASOS = os.path.join(OUT_DIR, "casos_uso.png")
FLUJO = os.path.join(OUT_DIR, "flujo_general.png")


# #####################################################################
# DOCUMENTO 1: BASE DE DATOS
# #####################################################################
doc = new_doc()
portada(doc, "DOCUMENTACION DE BASE DE DATOS",
        "Modelo Entidad-Relacion, diccionario de datos y relaciones")

doc.add_heading("1. Generalidades", level=1)
par(doc, "Motor de base de datos: Microsoft SQL Server.", bold=True)
par(doc, "Base de datos: LOGISTICA.")
par(doc, "Convenciones:")
bullet(doc, "Las tablas del modulo logístico usan el prefijo LOGISTICA_.")
bullet(doc, "Tablas nucleo de requerimientos: logistica_req (cabecera) y logistica_dreq (detalle).")
bullet(doc, "Tablas de integracion con el ERP Spring: WH_* (ordenes) y SPRING_* (inventario).")
bullet(doc, "Existen esquemas alternos/legacy (snake_case o sin prefijo) mantenidos por "
            "compatibilidad; se indican como 'esquema alterno/legacy' en la descripcion.")
bullet(doc, "La logica de negocio reside mayormente en Stored Procedures; los adjuntos se "
            "almacenan en columnas NVARCHAR(MAX) en base64.")
bullet(doc, "Control multiempresa por columna ruc / idEmpresa; borrado logico con 'eliminado'/'activo'.")

total_cols = sum(len(t["cols"]) for t in SCHEMA.values())
total_fks = sum(len(t["fks"]) for t in SCHEMA.values())
par(doc, f"Resumen: {len(SCHEMA)} tablas documentadas, {total_cols} columnas, "
         f"{total_fks} relaciones FK detectadas en los scripts.")

doc.add_heading("2. Diagrama Entidad-Relacion", level=1)
doc.add_heading("2.1 Flujo principal: Requerimiento -> Compra -> Almacen", level=2)
figura(doc, DER1, "Figura 1. DER del flujo principal de compras y almacen.")
doc.add_page_break()
doc.add_heading("2.2 Servicios, Consolidacion OC/OS, Administracion e Integracion", level=2)
figura(doc, DER2, "Figura 2. DER de servicios, consolidacion OC/OS, administracion e integracion.")
doc.add_page_break()

doc.add_heading("3. Diccionario de datos por modulo", level=1)
MODULO_INTRO = {
    "Requerimientos": "Tablas que registran los requerimientos de compra, consumo, "
                      "commodities, activos y transferencias, con sus adjuntos y trazabilidad.",
    "Aprobaciones": "Configuracion de flujos de aprobacion por area y por monto, "
                    "aprobadores y logs de decisiones (requerimientos, SC, OC, OS).",
    "Consolidacion de Requerimientos": "Agrupacion de items aprobados para generar "
                                       "cotizaciones, solicitudes u ordenes.",
    "Compras": "Cotizaciones, solicitudes de compra, ordenes de compra, proveedores, "
               "seguimiento e incrementos.",
    "Consolidacion OC/OS": "Ordenes de compra y servicio generadas desde el modulo de "
                           "consolidacion, con distribucion, adjuntos, conformidad y correos.",
    "Servicios": "Solicitudes de servicio, cotizaciones, ordenes de servicio, aprobaciones "
                 "y conformidad con firma digital.",
    "Almacen e Inventario": "Recepciones, despachos, kardex, stock, ajustes, conteos, "
                            "devoluciones, reingresos y conformidad de almacen.",
    "Notificaciones": "Notificaciones a usuarios y log de correos.",
    "Administracion y Configuracion": "Seguridad (usuarios, roles, perfiles), menu dinamico, "
                                      "permisos, parametros, SMTP, plantillas, branding y auditoria.",
    "Integracion ERP Spring": "Tablas espejo y cola de sincronizacion con el ERP Spring.",
}

for mod in MODULOS_ORDEN:
    tablas_mod = MODULO_TABLAS[mod]
    if not tablas_mod:
        continue
    doc.add_heading(f"3.{MODULOS_ORDEN.index(mod)+1} {mod} ({len(tablas_mod)} tablas)", level=2)
    par(doc, MODULO_INTRO.get(mod, ""), italic=True)
    for tname in tablas_mod:
        t = SCHEMA[tname]
        doc.add_heading(tname, level=3)
        par(doc, DESCRIPCIONES_TABLA.get(tname, ""), italic=True, size=9)
        rows = [(c["name"], c["type"], c["flags"], c["note"]) for c in t["cols"]]
        add_table(doc, ["Columna", "Tipo", "Restricciones", "Descripcion"], rows,
                  widths=[1.7, 1.5, 1.9, 2.1], font_size=8, hdr_size=8)
        if t["fks"]:
            fk_txt = "; ".join(f"{f['cols']} -> {f['ref']}({f['refcols']})" for f in t["fks"])
            par(doc, "FK: " + fk_txt, italic=True, size=8.5)
        par(doc, f"Fuente: {t['src']}", italic=True, size=8)
    doc.add_page_break()

doc.add_heading("4. Matriz de relaciones (FK)", level=1)
rel_rows = []
for tname in sorted(SCHEMA, key=str.lower):
    for f in SCHEMA[tname]["fks"]:
        rel_rows.append((tname, f["cols"], f"{f['ref']}({f['refcols']})"))
add_table(doc, ["Tabla", "Columna FK", "Referencia"], rel_rows,
          widths=[2.6, 1.8, 2.8])
par(doc, "Nota: ademas de las FK declaradas, existen relaciones logicas por convencion "
         "(por ejemplo idConsolidacion, numeroOrden, idrequerimiento) que no siempre tienen "
         "constraint fisico en la base de datos.", italic=True, size=9)

out1 = os.path.join(DOCS_DIR, "DOCUMENTACION_BASE_DE_DATOS.docx")
doc.save(out1)
print("OK", out1)


# #####################################################################
# DOCUMENTO 2: CASOS DE USO
# #####################################################################

# --- Mini diagramas de CU por paquete ---
def draw_actor_fig(ax, x, y, name):
    ax.add_patch(plt.Circle((x, y + 0.4), 0.13, fc="white", ec="black", lw=1.3, zorder=3))
    ax.plot([x, x], [y + 0.27, y - 0.1], color="black", lw=1.3, zorder=3)
    ax.plot([x - 0.22, x + 0.22], [y + 0.15, y + 0.15], color="black", lw=1.3, zorder=3)
    ax.plot([x, x - 0.18], [y - 0.1, y - 0.45], color="black", lw=1.3, zorder=3)
    ax.plot([x, x + 0.18], [y - 0.1, y - 0.45], color="black", lw=1.3, zorder=3)
    ax.text(x, y - 0.72, name, ha="center", va="top", fontsize=8, fontweight="bold")


def uc_package_diagram(filename, titulo, ucs, links):
    """ucs: lista de etiquetas; links: dict actor_name -> [indices uc]"""
    actors = list(links.keys())
    n_uc = len(ucs)
    h = max(4.5, n_uc * 0.75 + 1.2, len(actors) * 1.3 + 1)
    fig, ax = plt.subplots(figsize=(11, h))
    ax.set_xlim(0, 11); ax.set_ylim(0, h); ax.axis("off")
    ax.set_title(titulo, fontsize=12, fontweight="bold")
    # limite del sistema
    ax.add_patch(FancyBboxPatch((3.2, 0.3), 7.5, h - 0.9,
                 boxstyle="round,pad=0.03,rounding_size=0.1",
                 fc="none", ec="#333333", lw=1.4))
    uc_pos = []
    for i, u in enumerate(ucs):
        y = h - 1.3 - i * 0.75
        ax.add_patch(Ellipse((7.0, y), 5.4, 0.56, fc="#eef5fc", ec="#2e5f8a", lw=1.0))
        ax.text(7.0, y, u, ha="center", va="center", fontsize=8.5)
        uc_pos.append((7.0, y))
    apos = {}
    for i, a in enumerate(actors):
        y = h - 1.2 - i * 1.3
        draw_actor_fig(ax, 1.4, y, a)
        apos[a] = (1.75, y)
    for a, idxs in links.items():
        for i in idxs:
            ax.plot([apos[a][0], uc_pos[i][0] - 2.7], [apos[a][1], uc_pos[i][1]],
                    color="#888888", lw=0.8, zorder=0)
    fig.savefig(os.path.join(OUT_DIR, filename), dpi=140, bbox_inches="tight")
    plt.close(fig)


uc_package_diagram("uc_req_apr.png", "Casos de uso - Requerimientos y Aprobaciones",
    ["Registrar requerimiento (items)", "Registrar requerimiento commodity",
     "Registrar requerimiento activo fijo/menor", "Adjuntar documentos",
     "Consultar mis requerimientos", "Sincronizar pendientes", "Anular requerimiento",
     "Aprobar / rechazar requerimiento", "Aprobar saldo pendiente",
     "Ver reporte de aprobaciones", "Dashboard de aprobaciones"],
    {"Solicitante": [0, 1, 2, 3, 4, 5, 6], "Jefe de Area": [7, 8, 9, 10],
     "API Maestra": [0], "ERP Spring": [5]})

uc_package_diagram("uc_compras.png", "Casos de uso - Compras",
    ["Consolidar requerimientos", "Consolidacion corporativa",
     "Generar solicitud de cotizacion", "Registrar cotizacion de proveedor",
     "Evaluar y seleccionar cotizacion", "Generar solicitud de compra",
     "Aprobar solicitud de compra", "Generar orden de compra",
     "Enviar OC a aprobacion", "Aprobar / rechazar OC", "Incrementar cantidad OC",
     "Adjuntar documentos OC", "Sincronizar OC con ERP", "Seguimiento de OC",
     "Gestionar proveedores"],
    {"Analista de Logistica": [0, 1, 2, 3, 4, 5, 7, 8, 11, 13, 14],
     "Aprobador SC": [6], "Aprobador OC": [9], "Proveedor": [4, 12],
     "ERP Spring": [12]})

uc_package_diagram("uc_servicios.png", "Casos de uso - Servicios",
    ["Registrar solicitud de servicio", "Adjuntar documentos SS",
     "Aprobar solicitud de servicio", "Cotizar servicio",
     "Generar orden de servicio", "Enviar OS a aprobacion",
     "Aprobar / rechazar OS", "Sincronizar OS con ERP",
     "Registrar conformidad de servicio", "Seguimiento de OS",
     "Consultar historial de precios"],
    {"Solicitante": [0, 1], "Jefe de Area": [2], "Analista de Logistica": [3, 4, 5, 9, 10],
     "Aprobador OS": [6], "Usuario Conformidad": [8], "ERP Spring": [7],
     "Proveedor": [7]})

uc_package_diagram("uc_alm_admin.png", "Casos de uso - Almacen, Inventario y Administracion",
    ["Recepcionar mercaderia (OC)", "Registrar conformidad de almacen",
     "Despachar requerimientos", "Consultar kardex", "Consultar stock / listas",
     "Ajustar inventario", "Registrar conteo fisico", "Transferir entre almacenes",
     "Devolucion a proveedor", "Devolucion de consumo", "Registrar reingreso",
     "Recibir notificaciones", "Gestionar usuarios y roles", "Asignar usuarios por area",
     "Configurar flujos de aprobacion", "Configurar menu y permisos",
     "Administrar maestros", "Configurar SMTP / plantillas / branding",
     "Consultar auditoria", "Ver dashboards y reportes"],
    {"Almacenero": [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11],
     "Usuario Conformidad": [1], "Administrador": [12, 13, 14, 15, 16, 17, 18],
     "Todos los roles": [11, 19]})

print("figuras UC OK")

# --- Documento ---
doc = new_doc()
portada(doc, "MODELO DE CASOS DE USO",
        "Actores, diagrama y especificacion detallada de casos de uso")

doc.add_heading("1. Introduccion", level=1)
par(doc, "Este documento describe el modelo de casos de uso del Sistema de Logistica: "
         "los actores que interactuan con la aplicacion, el diagrama de casos de uso "
         "por paquete funcional y la especificacion detallada de cada caso de uso "
         "(flujo basico, flujos alternos, precondiciones, postcondiciones y reglas "
         "de negocio).")
par(doc, "Alcance: modulos de autenticacion, requerimientos, aprobaciones, "
         "consolidacion, cotizaciones, solicitudes y ordenes de compra, solicitudes y "
         "ordenes de servicio, recepcion, despacho, inventario, devoluciones, "
         "notificaciones, reportes y administracion.")

doc.add_heading("2. Actores del sistema", level=1)
doc.add_heading("2.1 Actores humanos", level=2)
add_table(doc, ["Actor", "Descripcion", "Modulos / Pantallas", "Guard (frontend)"], [
    ("Solicitante (Operativo)",
     "Usuario de area que registra requerimientos de compra, consumo, commodities, "
     "activos fijos/menores y servicios; consulta estados, adjunta documentos y "
     "sincroniza pendientes locales.",
     "Requerimientos, Notificaciones, Dashboard usuario",
     "OperativoGuard, OperativoEmpaqueGuard, EmpaqueGuard"),
    ("Jefe de Area (Aprobador)",
     "Aprueba o rechaza requerimientos de su area segun el flujo por niveles; "
     "resuelve saldos pendientes de aprobacion.",
     "Aprobaciones, Aprobaciones por Area, Dashboard Aprobaciones, Reporte de aprobaciones",
     "AprobadorGuard"),
    ("Analista de Logistica",
     "Consolida requerimientos aprobados, gestiona cotizaciones, genera solicitudes y "
     "ordenes de compra y de servicio, da seguimiento y sincroniza con el ERP.",
     "Consolidacion (requerimientos, compras, servicios, corporativa), Cotizaciones, "
     "Solicitudes de Compra/Servicio, Ordenes de Compra/Servicio, Dashboards de logistica, "
     "Reportes de compras",
     "LogisticoGuard, OperativoLogisticoGuard"),
    ("Almacenero",
     "Recepciona mercaderia contra OC, atiende (despacha) requerimientos, gestiona "
     "inventario, kardex, transferencias, listas de stock, devoluciones y reingresos.",
     "Recepcion de Mercaderia, Despachos, Kardex, Gestion de Inventario, Listas de Stock, "
     "Transferencias, Devoluciones, Reingresos, Reportes de almacen",
     "AlmacenGuard, DevolucionConsumoGuard, ReingresoGuard"),
    ("Aprobador OC/OS (Finanzas / Gerencia)",
     "Aprueba o rechaza solicitudes de compra y ordenes de compra/servicio segun "
     "niveles configurados por monto.",
     "Aprobaciones OC, Aprobaciones OS, Dashboard Finanzas",
     "AprobadorOCGuard"),
    ("Usuario Conformidad",
     "Registra la conformidad de los servicios ejecutados y la conformidad de almacen "
     "con firma digital.",
     "Conformidad de Servicios, Conformidad de Almacen",
     "ConformidadServiciosGuard, DevolucionConsumoGuard"),
    ("Administrador Logistica",
     "Administra usuarios del modulo logistico, parametros, plantillas y configuracion "
     "avanzada (modulo admin-logistica).",
     "Admin Logistica, parametros, plantillas",
     "AdminLogisticaGuard"),
    ("Administrador Sistema",
     "Configura roles, areas, flujos de aprobacion, menu dinamico, permisos, maestros "
     "y consulta la auditoria.",
     "Administracion, Maestros, Aprobadores, Configuracion",
     "AdminGuard, AdministradorGuard"),
], widths=[1.3, 2.6, 2.0, 1.6], font_size=8)

doc.add_heading("2.2 Actores de sistema (externos)", level=2)
add_table(doc, ["Actor", "Tipo", "Interaccion"], [
    ("API Maestra", "Sistema externo",
     "Valida credenciales de usuario en el login y devuelve datos de sesion (rol, area, empresa)."),
    ("ERP Spring", "Sistema externo",
     "Recibe OC/OS aprobadas (tablas WH_OrdenCompra / WH_OrdenServicio); envia maestro de "
     "items y stock (SPRING_*); la sincronizacion usa una cola con reintentos."),
    ("Proveedor", "Externo (correo)",
     "Recibe por correo las solicitudes de cotizacion y las ordenes de compra/servicio "
     "aprobadas en PDF."),
    ("Servidor SMTP", "Sistema externo",
     "Envia los correos a proveedores y notificaciones; registra resultado en logs."),
    ("Servicio de Tipo de Cambio", "Sistema externo",
     "Provee el tipo de cambio USD/PEN para documentos en moneda extranjera."),
], widths=[1.6, 1.2, 4.7], font_size=8.5)

doc.add_page_break()
doc.add_heading("3. Diagrama de casos de uso", level=1)
doc.add_heading("3.1 Diagrama general", level=2)
figura(doc, CASOS, "Figura 1. Diagrama general de casos de uso por modulo.")
doc.add_page_break()
doc.add_heading("3.2 Requerimientos y Aprobaciones", level=2)
figura(doc, os.path.join(OUT_DIR, "uc_req_apr.png"),
       "Figura 2. Casos de uso de requerimientos y aprobaciones.", 6.4)
doc.add_heading("3.3 Compras", level=2)
figura(doc, os.path.join(OUT_DIR, "uc_compras.png"),
       "Figura 3. Casos de uso del modulo de compras.", 6.4)
doc.add_heading("3.4 Servicios", level=2)
figura(doc, os.path.join(OUT_DIR, "uc_servicios.png"),
       "Figura 4. Casos de uso del modulo de servicios.", 6.4)
doc.add_heading("3.5 Almacen, Inventario y Administracion", level=2)
figura(doc, os.path.join(OUT_DIR, "uc_alm_admin.png"),
       "Figura 5. Casos de uso de almacen, inventario y administracion.", 6.4)
doc.add_page_break()

doc.add_heading("4. Especificacion de casos de uso", level=1)
par(doc, "Formato de especificacion: identificador, actores, precondiciones, flujo "
         "basico, flujos alternos, postcondiciones, reglas de negocio y tablas "
         "involucradas.", italic=True)

CU = [
 # --- Autenticacion ---
 dict(id="CU-001", nombre="Iniciar sesion", actor="Todos los usuarios",
      secundarios="API Maestra",
      modulo="auth/login",
      desc="El usuario ingresa sus credenciales; el sistema las valida contra la API "
           "Maestra, carga su rol/area/empresa y muestra el menu permitido.",
      prec="Usuario registrado y activo en la API Maestra.",
      flujo=["1. El usuario ingresa usuario y clave.",
             "2. El frontend envia las credenciales a la API.",
             "3. La API consulta la API Maestra para validar.",
             "4. Se obtiene rol, area, empresa y permisos del usuario.",
             "5. Se carga el menu dinamico segun el rol (LOGISTICA_ConfigMenu).",
             "6. Se redirige al dashboard correspondiente al rol."],
      alt=["3a. Credenciales invalidas: se muestra mensaje de error y se reintenta.",
           "4a. Usuario sin area/rol: se muestra menu reducido o mensaje de contacto al administrador.",
           "5a. Sin menu dinamico configurado: se usa el menu legacy por rol."],
      post="Sesion iniciada; token y datos de usuario disponibles.",
      reglas=["El menu se resuelve por rol + empresa + permisos configurados.",
              "Si hay requerimientos locales sin sincronizar, se alerta al usuario."],
      tablas="LOGISTICA_AdminUsuarios, LOGISTICA_UsuariosPorArea, LOGISTICA_ConfigMenu"),
 dict(id="CU-002", nombre="Login de administrador de logistica", actor="Administrador Logistica",
      secundarios="", modulo="admin-login / admin-logistica",
      desc="Acceso independiente al modulo de administracion logistica con usuario "
           "local (LOGISTICA_AdminUsuarios con password hasheado).",
      prec="Usuario administrador registrado en LOGISTICA_AdminUsuarios.",
      flujo=["1. El usuario ingresa credenciales en admin-login.",
             "2. La API valida usuario y hash de password.",
             "3. Se habilita el panel admin-logistica."],
      alt=["2a. Credenciales invalidas: error y reintento."],
      post="Sesion administrativa iniciada.",
      reglas=["El password se almacena hasheado (no en texto plano)."],
      tablas="LOGISTICA_AdminUsuarios, LOGISTICA_Auditoria"),
 # --- Requerimientos ---
 dict(id="CU-010", nombre="Registrar requerimiento", actor="Solicitante",
      secundarios="", modulo="main/requerimientos",
      desc="Crea un requerimiento de tipo COMPRA o CONSUMO con items del catalogo, "
           "commodities, activos fijos o activos menores.",
      prec="Usuario autenticado; catalogo de items disponible (MaestroItem).",
      flujo=["1. Selecciona tipo de requerimiento (compra/consumo) y tipo de item.",
             "2. Completa cabecera: empresa (ruc), fundo, area, almacen, centro de "
             "costo, proyecto, prioridad y glosa.",
             "3. Agrega items: codigo, descripcion, cantidad, unidad de medida.",
             "4. Adjunta archivos si el tipo lo requiere.",
             "5. Guarda: el registro queda en logistica_req / logistica_dreq en "
             "estado PENDIENTE.",
             "6. El sistema determina el aprobador segun LOGISTICA_FlujoAprobacionPorArea "
             "y LOGISTICA_UsuariosPorArea.",
             "7. Se notifica al aprobador asignado."],
      alt=["3a. Item inexistente: puede solicitar creacion de item "
           "(LOGISTICA_SolicitudCreacionItem).",
           "5a. Sin conexion: el requerimiento queda local y se sincroniza luego (CU-012).",
           "6a. Sin flujo configurado: queda pendiente sin asignar; se alerta a administracion."],
      post="Requerimiento PENDIENTE en bandeja del aprobador.",
      reglas=["El borrado es logico (eliminado=1).",
              "Cada detalle puede llevar centro de costo, proyecto, turno y labor."],
      tablas="logistica_req, logistica_dreq, LOGISTICA_RequerimientoAdjuntos, "
             "LOGISTICA_ReqCommodity, LOGISTICA_DReqCommodity, LOGISTICA_SolicitudCreacionItem"),
 dict(id="CU-011", nombre="Adjuntar documentos a un requerimiento", actor="Solicitante",
      secundarios="", modulo="main/requerimientos (modal adjuntos)",
      desc="Adjunta archivos (PDF, imagenes, Excel) al requerimiento; se guardan en "
           "base64 en la base de datos.",
      prec="Requerimiento registrado.",
      flujo=["1. Abre el detalle del requerimiento.",
             "2. Selecciona el archivo a adjuntar.",
             "3. El sistema lo codifica en base64 y lo guarda.",
             "4. El adjunto queda visible para el aprobador."],
      alt=["3a. Archivo excede el tamano maximo: se rechaza con mensaje.",
           "3b. Formato no permitido: se rechaza."],
      post="Adjunto disponible en el requerimiento.",
      reglas=["Los adjuntos obligatorios se validan antes del envio a aprobacion."],
      tablas="LOGISTICA_RequerimientoAdjuntos"),
 dict(id="CU-012", nombre="Sincronizar requerimientos pendientes", actor="Solicitante",
      secundarios="ERP Spring", modulo="main/requerimientos",
      desc="Los requerimientos guardados localmente (sin conexion) se envian al "
           "servidor. El cierre de sesion exige no tener pendientes.",
      prec="Existen requerimientos locales (items, commodities, activos).",
      flujo=["1. El sistema muestra la cantidad de pendientes por tipo.",
             "2. El usuario ejecuta la sincronizacion.",
             "3. Cada requerimiento se registra en el servidor.",
             "4. Al terminar, el contador queda en cero y se permite cerrar sesion."],
      alt=["3a. Falla parcial: quedan solo los no sincronizados; se informa el detalle."],
      post="Sin pendientes locales.",
      reglas=["No se permite cerrar sesion con requerimientos sin sincronizar."],
      tablas="logistica_req, logistica_dreq"),
 dict(id="CU-013", nombre="Anular requerimiento", actor="Solicitante / Administrador",
      secundarios="", modulo="main/requerimientos",
      desc="Anula un requerimiento registrado, registrando motivo y usuario.",
      prec="Requerimiento en estado que permita anulacion (no completado).",
      flujo=["1. Selecciona el requerimiento.",
             "2. Ingresa el motivo de anulacion.",
             "3. El sistema marca el estado ANULADO y registra el historial."],
      alt=["2a. Ya atendido/completado: no se permite anular."],
      post="Requerimiento ANULADO; se notifica al aprobador si estaba en bandeja.",
      reglas=["Toda anulacion queda en LOGISTICA_HistorialAnulaciones y "
              "LOGISTICA_LogCambioEstadoRequerimiento."],
      tablas="logistica_req, LOGISTICA_HistorialAnulaciones, LOGISTICA_LogCambioEstadoRequerimiento"),
 # --- Aprobaciones ---
 dict(id="CU-020", nombre="Aprobar o rechazar requerimiento por area", actor="Jefe de Area",
      secundarios="", modulo="main/aprobaciones, main/aprobaciones-area",
      desc="El jefe de area revisa los requerimientos pendientes de su area y decide "
           "por nivel de aprobacion.",
      prec="Requerimiento PENDIENTE; usuario con esJefeArea en el area.",
      flujo=["1. Ingresa a la bandeja de aprobaciones (se listan pendientes de su area "
             "segun LOGISTICA_UsuariosPorArea y el flujo configurado).",
             "2. Abre el detalle: items, cantidades, solicitante, adjuntos, glosa.",
             "3. Aprueba o rechaza con observacion.",
             "4. Se registra en LOGISTICA_LogAprobacionesPorArea.",
             "5. Si hay siguiente nivel, pasa al siguiente aprobador; si no, queda APROBADO.",
             "6. Se notifica al solicitante y/o al siguiente nivel."],
      alt=["3a. Rechazo: exige motivo; el requerimiento queda ANULADO/RECHAZADO.",
           "5a. Flujo hibrido: puede aprobar jefatura de area o aprobador corporativo "
           "segun configuracion."],
      post="Requerimiento aprobado (total o siguiente nivel) o rechazado.",
      reglas=["Un mismo usuario no puede aprobar dos niveles del mismo documento "
              "salvo autoaprobacion configurada (LOGISTICA_ConfigPermisos).",
              "El orden de niveles se respeta (nivelOrden)."],
      tablas="logistica_req, LOGISTICA_LogAprobacionesPorArea, LOGISTICA_UsuariosPorArea, "
             "LOGISTICA_FlujoAprobacionPorArea, logistica_notificaciones"),
 dict(id="CU-021", nombre="Aprobar saldo pendiente", actor="Jefe de Area / Logistica",
      secundarios="", modulo="main/aprobaciones (saldos)",
      desc="Cuando un despacho no cubre lo solicitado, el saldo queda pendiente de "
           "resolucion: aprobar para atencion posterior o cerrar.",
      prec="Existe LOGISTICA_SaldoPendienteAprobacion en estado PENDIENTE.",
      flujo=["1. Se listan los saldos pendientes con detalle por item "
             "(cantidad solicitada, despachada y saldo).",
             "2. El aprobador revisa y decide: aprobar, rechazar o cerrar el saldo.",
             "3. Se registra la resolucion con usuario y fecha."],
      alt=["2a. Si se aprueba: el saldo puede atenderse luego o pasar a consolidacion.",
           "2b. Si se cierra: el saldo se da por terminado."],
      post="Saldo APROBADO / RECHAZADO / CERRADO.",
      reglas=["El saldo mantiene trazabilidad al requerimiento origen."],
      tablas="LOGISTICA_SaldoPendienteAprobacion, LOGISTICA_SaldoPendienteAprobacionDet"),
 dict(id="CU-022", nombre="Consultar reporte de aprobaciones por area", actor="Jefe de Area",
      secundarios="", modulo="main/reporte-aprobaciones-area, main/dashboard-aprobaciones-area",
      desc="Reporte y dashboard con los requerimientos aprobados/rechazados por area, "
           "nivel y aprobador.",
      prec="Usuario con rol aprobador.",
      flujo=["1. Abre el reporte o dashboard.",
             "2. Filtra por fechas, area o estado.",
             "3. Visualiza el detalle y exporta si es necesario."],
      alt=[],
      post="Informacion de aprobaciones visualizada.",
      reglas=[],
      tablas="LOGISTICA_LogAprobacionesPorArea, LOGISTICA_Areas, LOGISTICA_UsuariosPorArea"),
 # --- Consolidacion ---
 dict(id="CU-030", nombre="Consolidar requerimientos", actor="Analista de Logistica",
      secundarios="", modulo="main/consolidacion-requerimientos, consolidacion-compras, "
      "consolidacion-corporativa",
      desc="Agrupa items aprobados de distintos requerimientos en una consolidacion "
           "para cotizar/comprar en conjunto.",
      prec="Existen detalles aprobados en LOGISTICA_ItemPendienteConsolidacion.",
      flujo=["1. Lista pendientes de consolidacion (items, commodities, activos, "
             "servicios) por empresa/tab.",
             "2. Filtra por empresa, almacen, item o tipo.",
             "3. Selecciona lineas compatibles y crea la consolidacion "
             "(LOGISTICA_ConsolidacionCab/Det).",
             "4. El sistema registra la trazabilidad en LOGISTICA_ConsolidacionOrigen "
             "y marca los pendientes como CONSOLIDADO.",
             "5. La consolidacion queda en estadoProceso PENDIENTE_COTIZACION."],
      alt=["3a. Puede anular lineas o items de la consolidacion (con motivo).",
           "4a. Consolidacion corporativa: agrupa entre empresas del grupo."],
      post="Consolidacion creada con codigo correlativo (CONS-anio- correlativo).",
      reglas=["Solo se consolidan detalles APROBADOS y no consolidados previamente.",
              "El correlativo se controla en LOGISTICA_Correlativo."],
      tablas="LOGISTICA_ItemPendienteConsolidacion, LOGISTICA_ConsolidacionCab, "
             "LOGISTICA_ConsolidacionDet, LOGISTICA_ConsolidacionOrigen, LOGISTICA_Correlativo"),
 # --- Cotizaciones ---
 dict(id="CU-040", nombre="Generar solicitud de cotizacion", actor="Analista de Logistica",
      secundarios="Proveedor (correo)", modulo="main/cotizaciones",
      desc="Genera una solicitud de cotizacion desde una consolidacion y la envia a "
           "proveedores.",
      prec="Consolidacion en estado PENDIENTE_COTIZACION.",
      flujo=["1. Selecciona la consolidacion.",
             "2. Genera LOGISTICA_SolicitudCotizacion(+Det) con correlativo SOL-COT.",
             "3. Selecciona proveedores y envia la solicitud por correo (SMTP).",
             "4. La consolidacion pasa a EN_COTIZACION."],
      alt=["3a. Falla el envio: queda registrado en log de correos para reintento."],
      post="Solicitud de cotizacion GENERADA/ENVIADA.",
      reglas=["La solicitud referencia la consolidacion de origen."],
      tablas="LOGISTICA_SolicitudCotizacion, LOGISTICA_SolicitudCotizacionDet, "
             "LOGISTICA_Correlativo, ConfigSmtp, LOGISTICA_CorreoLog"),
 dict(id="CU-041", nombre="Registrar y evaluar cotizaciones de proveedores",
      actor="Analista de Logistica", secundarios="Proveedor",
      modulo="main/cotizaciones, main/evaluacion-proveedores",
      desc="Registra las propuestas de los proveedores, las compara y selecciona la "
           "ganadora.",
      prec="Solicitud de cotizacion ENVIADA.",
      flujo=["1. Registra la cotizacion del proveedor (precios, plazos, condiciones "
             "de pago, forma de pago).",
             "2. El sistema permite comparar cotizaciones por item.",
             "3. Se selecciona el proveedor ganador (marca 'seleccionada').",
             "4. La solicitud pasa a COTIZADA / CERRADA."],
      alt=["3a. Puede evaluar por precio, plazo y condiciones.",
           "4a. Si ninguna cotizacion es viable, se puede cerrar sin seleccion."],
      post="Proveedor seleccionado; se habilita generacion de OC/SC.",
      reglas=["La cotizacion seleccionada alimenta los precios de la OC."],
      tablas="LOGISTICA_Cotizacion, LOGISTICA_CotizacionDet, LOGISTICA_Proveedor, "
             "LOGISTICA_CotizacionOC"),
 # --- Compras ---
 dict(id="CU-050", nombre="Generar solicitud de compra", actor="Analista de Logistica",
      secundarios="", modulo="main/solicitudes-compra",
      desc="Crea una solicitud de compra (CONSOLIDADA, DIRECTA o URGENTE) con su "
           "detalle y adjuntos.",
      prec="Consolidacion cotizada o necesidad directa autorizada.",
      flujo=["1. Genera la SC desde la consolidacion o de forma directa.",
             "2. Completa almacen, prioridad y observaciones.",
             "3. Adjunta documentos (LOGISTICA_SolicitudCompraAdjuntos).",
             "4. Envia a aprobacion: estado ENVIADA."],
      alt=["4a. Si no requiere aprobacion (monto/permiso), puede pasar directo."],
      post="SC ENVIADA en el flujo de aprobacion.",
      reglas=["numeroSolicitud es unico; la SC registra usuario y fecha de envio."],
      tablas="LOGISTICA_SolicitudCompra, LOGISTICA_SolicitudCompraDet, "
             "LOGISTICA_SolicitudCompraAdjuntos"),
 dict(id="CU-051", nombre="Aprobar solicitud de compra", actor="Aprobador SC",
      secundarios="", modulo="main/aprobaciones (SC)",
      desc="Aprueba o rechaza la SC segun el flujo multinivel configurado por monto.",
      prec="SC en estado ENVIADA.",
      flujo=["1. El sistema asigna el nivel de aprobacion segun "
             "LOGISTICA_FlujoAprobacionSC (montoMin/montoMax).",
             "2. El aprobador revisa detalle y adjuntos.",
             "3. Aprueba o rechaza; se registra en LOGISTICA_AprobacionSC y el log.",
             "4. Si completan todos los niveles: SC APROBADA."],
      alt=["3a. Rechazo: estado RECHAZADA con motivo.",
           "4a. Si el usuario tiene rango suficiente puede aprobar el nivel final."],
      post="SC APROBADA o RECHAZADA.",
      reglas=["Los niveles se asignan por monto total de la SC.",
              "LOGISTICA_UsuariosAprobadoresSC define quien puede aprobar cada nivel."],
      tablas="LOGISTICA_FlujoAprobacionSC, LOGISTICA_AprobacionSC, "
             "LOGISTICA_UsuariosAprobadoresSC, LOGISTICA_LogAprobacionesSC"),
 dict(id="CU-052", nombre="Generar orden de compra", actor="Analista de Logistica",
      secundarios="", modulo="main/ordenes-compra",
      desc="Genera la OC desde la SC aprobada o desde la consolidacion "
           "(LOGISTICA_ConsolOC), con proveedor, moneda, condiciones y detalle.",
      prec="SC aprobada o cotizacion seleccionada.",
      flujo=["1. Genera la OC con numeroOrden correlativo.",
             "2. Completa datos del proveedor (maestro LOGISTICA_Proveedor), moneda, "
             "tipo de cambio, condiciones y forma de pago, lugar y fecha de entrega.",
             "3. El detalle se calcula con precio unitario, descuento, subtotal, IGV "
             "y total por item.",
             "4. La OC queda en estado PENDIENTE/GENERADA."],
      alt=["2a. Moneda USD: se toma el tipo de cambio del servicio.",
           "1a. Desde consolidacion se crea LOGISTICA_ConsolOC con distribucion por "
           "requerimiento (ConsolOCDist)."],
      post="OC registrada y lista para aprobacion.",
      reglas=["El total se calcula como subtotal + IGV (18%).",
              "Toda OC referencia su solicitud o consolidacion de origen."],
      tablas="LOGISTICA_OrdenCompra, LOGISTICA_OrdenCompraDet, LOGISTICA_OrdenCompraLog, "
             "LOGISTICA_ConsolOC, LOGISTICA_ConsolOCDet, LOGISTICA_ConsolOCDist"),
 dict(id="CU-053", nombre="Aprobar orden de compra", actor="Aprobador OC",
      secundarios="", modulo="main/aprobaciones-oc, main/dashboard-finanzas",
      desc="La OC pasa por aprobacion multinivel por monto; cada aprobador registra "
           "su decision.",
      prec="OC enviada a aprobacion (PENDIENTE_APROBACION).",
      flujo=["1. El flujo asigna niveles segun monto (LOGISTICA_FlujoAprobacionOC).",
             "2. Cada aprobador revisa la OC y sus adjuntos (LOGISTICA_OCAdjunto).",
             "3. Aprueba o rechaza; se registra LOGISTICA_AprobacionOC.",
             "4. Al completar los niveles: OC APROBADA.",
             "5. Se genera el PDF de la OC y se envia al proveedor por correo."],
      alt=["3a. Rechazo en cualquier nivel: OC ANULADA/rechazada; se notifica.",
           "5a. Si el correo falla, queda en ConsolCorreoLog para reintento."],
      post="OC APROBADA; PDF generado; correo enviado al proveedor.",
      reglas=["El nivel requerido depende del rango de monto.",
              "Se puede consultar contadores y bandeja por aprobador."],
      tablas="LOGISTICA_FlujoAprobacionOC, LOGISTICA_UsuariosAprobadoresOC, "
             "LOGISTICA_AprobacionOC, LOGISTICA_LogAprobacionesOC, LOGISTICA_OCAdjunto, "
             "LOGISTICA_ConsolConfigEmpresa, LOGISTICA_ConsolCorreoLog"),
 dict(id="CU-054", nombre="Sincronizar OC con el ERP", actor="Analista de Logistica / Sistema",
      secundarios="ERP Spring", modulo="main/ordenes-compra",
      desc="La OC aprobada se sincroniza a WH_OrdenCompra del ERP Spring; los errores "
           "quedan en una cola de reintentos.",
      prec="OC APROBADA.",
      flujo=["1. El SP de sincronizacion arma el payload de la OC.",
             "2. Inserta/actualiza la OC en el ERP (WH_OrdenCompra).",
             "3. Marca la OC como sincronizada y registra el log.",
             "4. Si falla: se encola en LOGISTICA_ColaMigracionSpring con reintentos."],
      alt=["2a. Multiempresa: se sincroniza con el idEmpresa correspondiente.",
           "4a. Reintento automatico o manual desde la cola."],
      post="OC disponible en el ERP; trazabilidad en logs.",
      reglas=["No cambia el estado de la OC si la sincronizacion falla.",
              "Todo evento se registra en LOGISTICA_SincronizacionLog."],
      tablas="WH_OrdenCompra, LOGISTICA_ColaMigracionSpring, LOGISTICA_SincronizacionLog"),
 dict(id="CU-055", nombre="Incrementar cantidad de una OC", actor="Analista de Logistica",
      secundarios="Aprobador OC", modulo="main/ordenes-compra",
      desc="Permite aumentar cantidades en una OC segun la configuracion de "
           "incremento (porcentaje maximo, requiere reaprobacion).",
      prec="OC en estado que permita incremento; configuracion habilitada.",
      flujo=["1. Selecciona la OC y aplica el incremento en el item.",
             "2. El sistema valida contra LOGISTICA_ConfigIncrementoOC.",
             "3. Se registra el incremento en LOGISTICA_OCIncrementoLog.",
             "4. Si la configuracion lo exige, la OC vuelve a aprobacion."],
      alt=["2a. Excede porcentaje permitido: se rechaza el incremento."],
      post="OC con cantidades actualizadas; log generado.",
      reglas=["Todo incremento es auditable."],
      tablas="LOGISTICA_ConfigIncrementoOC, LOGISTICA_ConfigIncrementoOCAuditoria, "
             "LOGISTICA_OCIncrementoLog, LOGISTICA_OrdenCompraDet"),
 dict(id="CU-056", nombre="Dar seguimiento a la OC", actor="Analista de Logistica",
      secundarios="", modulo="main/ordenes-compra, seguimiento",
      desc="Registra y consulta las etapas de la OC: emitida, enviada al proveedor, "
           "confirmada, en transito, entregada.",
      prec="OC existente.",
      flujo=["1. Se actualizan las etapas en LOGISTICA_SeguimientoOrdenCompra.",
             "2. Se consulta el timeline de la OC."],
      alt=[],
      post="Seguimiento actualizado.",
      reglas=[],
      tablas="LOGISTICA_SeguimientoOrdenCompra, LOGISTICA_OrdenCompra"),
 dict(id="CU-057", nombre="Gestionar maestro de proveedores", actor="Administrador / Almacenero",
      secundarios="", modulo="main/maestro-proveedores",
      desc="Registra y mantiene proveedores: RUC, razon social, contacto, emails, "
           "categoria y evaluacion.",
      prec="Permisos de administrador o almacen.",
      flujo=["1. Registra/edita proveedor con sus datos.",
             "2. Asocia correos (ProveedoresEmail) y categoria.",
             "3. Evaluacion de proveedores disponible para analisis."],
      alt=[],
      post="Proveedor disponible para cotizaciones y ordenes.",
      reglas=[],
      tablas="LOGISTICA_Proveedor, LOGISTICA_CategoriaProveedor, ProveedoresEmail"),
 # --- Recepcion / Almacen ---
 dict(id="CU-060", nombre="Recepcionar mercaderia de una OC", actor="Almacenero",
      secundarios="", modulo="main/recepcion-mercaderia",
      desc="Registra la recepcion fisica contra la OC, con cantidades recibidas y "
           "aceptadas, numero de nota de ingreso y estado.",
      prec="OC aprobada/sincronizada pendiente de recepcion.",
      flujo=["1. Lista OCs pendientes de recepcion y selecciona una.",
             "2. Revisa el detalle y registra cantidad recibida/aceptada por item.",
             "3. Genera la recepcion con numeroNI (nota de ingreso).",
             "4. El sistema crea la transaccion de ingreso "
             "(LOGISTICA_TransaccionInventario) y actualiza Kardex y Stock.",
             "5. Se replica al ERP (SPRING_*) y se notifica disponibilidad."],
      alt=["2a. Recepcion parcial: la OC queda RECIBIDA_PARCIAL hasta completar.",
           "4a. Discrepancias: se puede derivar a devolucion a proveedor."],
      post="Stock incrementado; OC RECIBIDA (parcial/total); requerimiento origen "
           "puede marcarse COMPLETADO.",
      reglas=["La recepcion referencia la OC y su detalle.",
              "El ingreso actualiza kardex con costo unitario."],
      tablas="LOGISTICA_RecepcionOC, LOGISTICA_RecepcionOCDet, LOGISTICA_RecepcionOCLog, "
             "LOGISTICA_TransaccionInventario, LOGISTICA_DetalleTransaccionInventario, "
             "LOGISTICA_Kardex, LOGISTICA_Stock, SPRING_Stock, SPRING_Kardex"),
 dict(id="CU-061", nombre="Despachar requerimiento", actor="Almacenero",
      secundarios="", modulo="main/despachos",
      desc="Atiende requerimientos de consumo con el stock disponible; los saldos "
           "generan aprobacion de pendiente.",
      prec="Requerimiento APROBADO de consumo con items en stock.",
      flujo=["1. Lista requerimientos pendientes de despacho y valida stock.",
             "2. Registra el despacho (LOGISTICA_Despachos/DetalleDespacho).",
             "3. El sistema descuenta stock y registra salida en kardex.",
             "4. Si se despacho parcial: se crea LOGISTICA_SaldoPendienteAprobacion.",
             "5. El requerimiento pasa a ATENDIDO (parcial o total)."],
      alt=["2a. Sin stock suficiente: despacho parcial o no se despacha.",
           "4a. El saldo pendiente requiere resolucion por aprobador."],
      post="Salida registrada en kardex; estado del requerimiento actualizado.",
      reglas=["El despacho no puede exceder la cantidad aprobada ni el stock."],
      tablas="LOGISTICA_Despachos, LOGISTICA_DetalleDespacho, LOGISTICA_Kardex, "
             "LOGISTICA_Stock, LOGISTICA_SaldoPendienteAprobacion"),
 dict(id="CU-062", nombre="Consultar kardex y stock", actor="Almacenero / Logistica",
      secundarios="", modulo="main/kardex, main/listas-stock, main/reporte_logistico",
      desc="Consulta movimientos de kardex (entradas, salidas, saldos valorizados) "
           "por item, almacen, compania y fechas.",
      prec="Existen movimientos registrados.",
      flujo=["1. Filtra por item/almacen/empresa/fechas.",
             "2. Visualiza movimientos con saldos cantidad y valor.",
             "3. Puede exportar el reporte."],
      alt=[],
      post="Informacion de kardex/stock visualizada.",
      reglas=["El alcance por empresa depende del rol (permisos)."],
      tablas="LOGISTICA_Kardex, LOGISTICA_Stock, LOGISTICA_Item, LOGISTICA_Almacen, "
             "LOGISTICA_ListaStock, LOGISTICA_AlertaStock"),
 dict(id="CU-063", nombre="Gestionar inventario (ajustes y conteos)", actor="Almacenero",
      secundarios="", modulo="main/gestion-inventario",
      desc="Registra ajustes de inventario y conteos fisicos para corregir "
           "diferencias.",
      prec="Permisos de almacen.",
      flujo=["1. Registra ajuste (tipo, motivo, items y cantidades).",
             "2. O registra conteo fisico y compara contra stock de sistema.",
             "3. El ajuste genera transacciones de entrada/salida y actualiza kardex."],
      alt=["2a. Diferencias: se genera ajuste automatico con el resultado del conteo."],
      post="Stock corregido y trazado.",
      reglas=["Los ajustes requieren motivo y quedan auditados."],
      tablas="LOGISTICA_AjusteInventario, LOGISTICA_AjusteInventarioDet, "
             "LOGISTICA_ConteoInventario, LOGISTICA_ConteoInventarioDet"),
 dict(id="CU-064", nombre="Transferir stock entre almacenes", actor="Almacenero",
      secundarios="", modulo="main/transferencias-almacen",
      desc="Genera requerimientos de transferencia entre almacenes (incluido almacen "
           "virtual); al aprobarse se mueve el stock.",
      prec="Stock disponible en almacen origen.",
      flujo=["1. Crea LOGISTICA_RequerimientoTransferencia (origen, destino, items).",
             "2. La transferencia puede requerir aprobacion.",
             "3. Al aprobarse/ejecutarse se generan salida del origen e ingreso al "
             "destino en kardex."],
      alt=["2a. Rechazada: no se mueve stock."],
      post="Stock transferido entre almacenes.",
      reglas=["La transferencia genera dos movimientos de kardex (salida/ingreso)."],
      tablas="LOGISTICA_RequerimientoTransferencia, LOGISTICA_TransaccionInventario, "
             "LOGISTICA_Kardex, LOGISTICA_Stock"),
 dict(id="CU-065", nombre="Registrar devoluciones y reingresos", actor="Almacenero",
      secundarios="", modulo="main/devoluciones-proveedores, devoluciones-consumo, reingresos",
      desc="Devolucion a proveedor (mercaderia recibida), devolucion de consumo "
           "(interna) y reingreso de material.",
      prec="Documento origen existente (recepcion, despacho).",
      flujo=["1. Selecciona el tipo: devolucion a proveedor, devolucion de consumo o "
             "reingreso.",
             "2. Referencia el documento origen y registra cantidades y motivo.",
             "3. El sistema genera la transaccion inversa y actualiza kardex/stock."],
      alt=["1a. Devolucion a proveedor puede requerir aprobacion previa."],
      post="Kardex y stock actualizados segun el tipo.",
      reglas=["No se puede devolver mas de lo recibido/despachado."],
      tablas="LOGISTICA_DevolucionProveedor, LOGISTICA_DetalleDevolucionProveedor, "
             "logistica_devolucion_consumo, logistica_devolucion_consumo_det, "
             "LOGISTICA_Reingreso, LOGISTICA_DetalleReingreso"),
 dict(id="CU-066", nombre="Registrar conformidad de almacen", actor="Usuario Conformidad / Almacenero",
      secundarios="", modulo="main/conformidad-almacen",
      desc="Confirma que la recepcion/almacenaje es conforme; usa plantilla "
           "configurable.",
      prec="Recepcion registrada.",
      flujo=["1. Selecciona la recepcion pendiente de conformidad.",
             "2. Registra conformidad con observaciones/firma.",
             "3. Se actualiza el estado del documento."],
      alt=["2a. No conforme: se registra observacion y puede iniciar devolucion."],
      post="Conformidad registrada.",
      reglas=[],
      tablas="LOGISTICA_ConformidadAlmacen, LOGISTICA_PlantillaAlmacen"),
 # --- Servicios ---
 dict(id="CU-070", nombre="Registrar solicitud de servicio", actor="Solicitante / Logistica",
      secundarios="", modulo="main/solicitudes-servicio",
      desc="Registra una solicitud de servicio con detalle, alcance, fechas y "
           "adjuntos.",
      prec="Usuario autenticado.",
      flujo=["1. Completa la solicitud: tipo de servicio, descripcion, alcance, "
             "ubicacion, fechas, centro de costo/proyecto.",
             "2. Adjunta documentos (LOGISTICA_SolicitudServicioAdjuntos).",
             "3. Envia a aprobacion si aplica (LOGISTICA_AprobacionSS)."],
      alt=[],
      post="SS registrada PENDIENTE/enviada a aprobacion.",
      reglas=[],
      tablas="LOGISTICA_SolicitudServicio, LOGISTICA_SolicitudServicioDet, "
             "LOGISTICA_SolicitudServicioAdjuntos, LOGISTICA_AprobacionSS"),
 dict(id="CU-071", nombre="Generar orden de servicio", actor="Analista de Logistica",
      secundarios="", modulo="main/ordenes-servicio",
      desc="Genera la OS desde la SS aprobada o desde consolidacion "
           "(LOGISTICA_ConsolOS), con proveedor, fechas, plazo y monto.",
      prec="SS aprobada o consolidacion de servicios.",
      flujo=["1. Cotiza el servicio si aplica (LOGISTICA_CotizacionServicio).",
             "2. Genera la OS con numeroOrden correlativo, proveedor, rango de "
             "fechas, plazo, centro de costo y monto.",
             "3. Envia a aprobacion OS si requiereAprobacion."],
      alt=["2a. Desde consolidacion se crea LOGISTICA_ConsolOS(+Det)."],
      post="OS registrada PENDIENTE/PENDIENTE_APROBACION.",
      reglas=["Se consulta el historial de precios del servicio como referencia."],
      tablas="LOGISTICA_OrdenServicio, LOGISTICA_OrdenServicioDet, "
             "LOGISTICA_CotizacionServicio, LOGISTICA_ConsolOS, LOGISTICA_ConsolOSDet, "
             "LOGISTICA_HistorialPrecioServicio"),
 dict(id="CU-072", nombre="Aprobar orden de servicio", actor="Aprobador OS",
      secundarios="", modulo="main/aprobaciones-os",
      desc="Aprobacion multinivel de la OS; al aprobarse se envia al proveedor y se "
           "sincroniza con el ERP.",
      prec="OS en PENDIENTE_APROBACION.",
      flujo=["1. El flujo asigna niveles (LOGISTICA_FlujoAprobacionOS / "
             "UsuarioAprobadorOS).",
             "2. El aprobador revisa y decide; se registra LOGISTICA_AprobacionOS.",
             "3. Aprobada: se genera PDF, correo al proveedor y sincronizacion a "
             "WH_OrdenServicio."],
      alt=["2a. Rechazo: OS rechazada con motivo."],
      post="OS APROBADA y sincronizada.",
      reglas=[],
      tablas="LOGISTICA_FlujoAprobacionOS, LOGISTICA_UsuarioAprobadorOS, "
             "LOGISTICA_AprobacionOS, LOGISTICA_LogAprobacionesOS, WH_OrdenServicio"),
 dict(id="CU-073", nombre="Registrar conformidad de servicio", actor="Usuario Conformidad",
      secundarios="", modulo="main/conformidad-servicios",
      desc="Registra la conformidad del servicio ejecutado con firma digital del "
           "jefe de area y acta adjunta.",
      prec="OS ejecutada pendiente de conformidad.",
      flujo=["1. Lista OS pendientes de conformidad (filtradas por usuario/area).",
             "2. Registra CONFORME o NO_CONFORME con observaciones.",
             "3. Adjunta firma digital (firmaBase64) y/o acta (actaBase64).",
             "4. Se cierra el ciclo del servicio."],
      alt=["3a. NO_CONFORME: puede generar observaciones y acciones correctivas."],
      post="OS con conformidad registrada; servicio completado.",
      reglas=["La plantilla de conformidad es configurable "
              "(LOGISTICA_PlantillaConformidadServicio)."],
      tablas="LOGISTICA_ConformidadServicio, LOGISTICA_ConsolOSConformidad, "
             "LOGISTICA_ConfirmacionServicioRealizado"),
 # --- Notificaciones / reportes ---
 dict(id="CU-080", nombre="Recibir y gestionar notificaciones", actor="Todos los usuarios",
      secundarios="", modulo="main/notificaciones",
      desc="El sistema notifica aprobaciones pendientes, stock disponible, saldos y "
           "eventos del proceso; el contador se actualiza cada 30 segundos.",
      prec="Usuario autenticado.",
      flujo=["1. El sistema genera notificaciones (logistica_notificaciones, "
             "LOGISTICA_NotificacionesStock).",
             "2. El usuario las consulta y marca como leidas.",
             "3. Las alertas de stock avisan nueva disponibilidad o niveles criticos."],
      alt=[],
      post="Notificaciones gestionadas.",
      reglas=["Las notificaciones son por usuario y, en alertas de stock, por item."],
      tablas="logistica_notificaciones, LOGISTICA_NotificacionesStock, LOGISTICA_AlertaStock"),
 dict(id="CU-081", nombre="Consultar dashboards y reportes", actor="Todos los roles",
      secundarios="", modulo="dashboards por rol, reportes-*",
      desc="Dashboards con KPIs por rol (logistica, compras, aprobaciones, finanzas, "
           "almacen) y reportes de saldos, despachos, aprobaciones, requerimientos y "
           "compras.",
      prec="Usuario autenticado con rol asignado.",
      flujo=["1. Accede al dashboard de su rol: KPIs, contadores por estado, "
             "pendientes.",
             "2. Genera reportes filtrando por fechas/empresa/estado.",
             "3. Exporta a Excel/PDF segun el reporte."],
      alt=[],
      post="Informacion de gestion disponible.",
      reglas=["El alcance de datos depende del rol y empresa del usuario."],
      tablas="Consultas sobre todas las tablas del proceso"),
 # --- Administracion ---
 dict(id="CU-090", nombre="Gestionar usuarios y asignacion por area", actor="Administrador",
      secundarios="", modulo="administracion / admin-logistica",
      desc="Administra usuarios del sistema y del modulo logistico, y los asigna a "
           "areas (incluida la marca de jefe de area).",
      prec="Rol administrador.",
      flujo=["1. Registra/edita usuarios (LOGISTICA_AdminUsuarios / LOGISTICA_Usuarios).",
             "2. Asigna usuarios a areas y define jefes "
             "(LOGISTICA_UsuariosPorArea.esJefeArea).",
             "3. Migra usuarios desde el ERP si aplica."],
      alt=[],
      post="Usuarios configurados; los flujos de aprobacion pueden resolver jefes.",
      reglas=["Un usuario sin area puede no ver pendientes de aprobacion."],
      tablas="LOGISTICA_AdminUsuarios, LOGISTICA_Usuarios, LOGISTICA_UsuariosPorArea, "
             "LOGISTICA_Areas"),
 dict(id="CU-091", nombre="Configurar flujos de aprobacion", actor="Administrador",
      secundarios="", modulo="administracion / aprobadores",
      desc="Configura los niveles de aprobacion por area (requerimientos) y por "
           "monto (SC, OC, OS), asignando aprobadores.",
      prec="Rol administrador.",
      flujo=["1. Define niveles y tipo de aprobador por area "
             "(LOGISTICA_FlujoAprobacionPorArea).",
             "2. Define rangos de monto y roles para SC/OC/OS.",
             "3. Asigna usuarios aprobadores con su nivel y rango."],
      alt=[],
      post="Flujos activos; los documentos se enrutan automaticamente.",
      reglas=["Los cambios de configuracion aplican a documentos nuevos/en curso "
              "segun implementacion."],
      tablas="LOGISTICA_FlujoAprobacionPorArea, LOGISTICA_FlujoAprobacionSC, "
             "LOGISTICA_FlujoAprobacionOC, LOGISTICA_FlujoAprobacionOS, "
             "LOGISTICA_UsuariosAprobadoresOC, LOGISTICA_UsuariosAprobadoresSC, "
             "LOGISTICA_UsuarioAprobadorOS"),
 dict(id="CU-092", nombre="Configurar menu dinamico y permisos", actor="Administrador",
      secundarios="", modulo="administracion",
      desc="Configura el menu visible por rol/empresa y los permisos (incluidos "
           "montos maximos y autoaprobacion).",
      prec="Rol administrador.",
      flujo=["1. Define items de menu por rol (LOGISTICA_ConfigMenu).",
             "2. Define permisos y valores maximos por rol "
             "(LOGISTICA_ConfigPermisos / ConfiguracionPermiso).",
             "3. Los usuarios ven el menu configurado al iniciar sesion."],
      alt=[],
      post="Menu y permisos aplicados por rol.",
      reglas=["Si no hay menu dinamico configurado se usa el menu legacy."],
      tablas="LOGISTICA_ConfigMenu, LOGISTICA_ConfigPermisos, LOGISTICA_ConfiguracionPermiso"),
 dict(id="CU-093", nombre="Administrar configuracion general", actor="Administrador Logistica / Sistema",
      secundarios="", modulo="admin-logistica",
      desc="Gestiona parametros, SMTP, plantillas de correo y PDF, branding por "
           "empresa, unidades de medida, correlativos y consulta auditoria y "
           "despliegues.",
      prec="Rol administrador.",
      flujo=["1. Ajusta parametros generales (LOGISTICA_Parametros).",
             "2. Configura SMTP y plantillas (ConfigSmtp, PlantillasCorreo, "
             "PlantillasPdf).",
             "3. Configura branding/empresa para documentos.",
             "4. Consulta auditoria y versiones desplegadas."],
      alt=[],
      post="Configuracion vigente aplicada en los modulos.",
      reglas=["Toda configuracion critica queda auditada."],
      tablas="LOGISTICA_Parametros, ConfigSmtp, PlantillasCorreo, PlantillasPdf, "
             "LOGISTICA_ConfigBranding, LOGISTICA_ConfigEmpresa, LOGISTICA_UnidadMedida, "
             "LOGISTICA_Correlativo, LOGISTICA_Auditoria, AppDeployments"),
]

for cu in CU:
    doc.add_heading(f"{cu['id']} - {cu['nombre']}", level=2)
    add_table(doc, ["Campo", "Detalle"], [
        ("Identificador", cu["id"]),
        ("Nombre", cu["nombre"]),
        ("Actor principal", cu["actor"]),
        ("Actores secundarios", cu["secundarios"] or "-"),
        ("Modulo / pantalla", cu["modulo"]),
        ("Descripcion", cu["desc"]),
        ("Precondiciones", cu["prec"]),
        ("Flujo basico", "\n".join(cu["flujo"])),
        ("Flujos alternos", "\n".join(cu["alt"]) if cu["alt"] else "-"),
        ("Postcondiciones", cu["post"]),
        ("Reglas de negocio", "\n".join(cu["reglas"]) if cu["reglas"] else "-"),
        ("Tablas involucradas", cu["tablas"]),
    ], widths=[1.5, 5.8], font_size=8.5)
    doc.add_paragraph()

out2 = os.path.join(DOCS_DIR, "DOCUMENTACION_CASOS_DE_USO.docx")
doc.save(out2)
print("OK", out2)


# #####################################################################
# DOCUMENTO 3: DESCRIPCION DE FLUJOS
# #####################################################################
doc = new_doc()
portada(doc, "DESCRIPCION DE FLUJOS DEL SISTEMA",
        "Flujos de proceso, estados y tablas involucradas")

doc.add_heading("1. Vision general", level=1)
par(doc, "El sistema implementa el ciclo completo de abastecimiento: un area solicita "
         "(requerimiento), su jefatura aprueba, Logistica consolida y cotiza, se "
         "generan las ordenes (compra o servicio), se aprueban por niveles segun "
         "monto, se sincronizan con el ERP Spring, Almacen recepciona/despacha y el "
         "inventario queda actualizado en kardex y stock.")
figura(doc, FLUJO, "Figura 1. Flujo general del proceso logistico.")
doc.add_paragraph()
add_table(doc, ["#", "Etapa", "Actor", "Resultado"], [
    ("1", "Requerimiento", "Solicitante", "Documento PENDIENTE (logistica_req/dreq)"),
    ("2", "Aprobacion por area", "Jefe de Area", "APROBADO o ANULADO (por niveles)"),
    ("3", "Consolidacion", "Logistica", "Items agrupados (ConsolidacionCab/Det/Origen)"),
    ("4", "Cotizacion", "Logistica + Proveedor", "Proveedor seleccionado"),
    ("5", "SC / OC / OS", "Logistica", "Documento generado"),
    ("6", "Aprobacion de orden", "Finanzas/Gerencia", "APROBADA por niveles de monto"),
    ("7", "Sincronizacion", "Sistema/ERP Spring", "Orden en ERP (WH_*), cola de reintentos"),
    ("8", "Recepcion", "Almacen", "Nota de ingreso, OC recibida"),
    ("9", "Kardex / Stock", "Sistema", "Inventario actualizado"),
    ("10", "Despacho", "Almacen", "Requerimiento atendido"),
    ("11", "Saldo pendiente", "Aprobador", "Saldo aprobado / cerrado / re-consolidado"),
], widths=[0.4, 1.8, 1.6, 3.2])
doc.add_page_break()


def flujo_seccion(num, titulo, objetivo, actores, pasos, estados_img=None,
                  estados_cap="", tablas="", sps="", pantallas=""):
    doc.add_heading(f"{num}. {titulo}", level=1)
    par(doc, "Objetivo: ", bold=True).add_run(objetivo)
    par(doc, "Actores: ", bold=True).add_run(actores)
    if pantallas:
        par(doc, "Pantallas: ", bold=True).add_run(pantallas)
    doc.add_heading(f"{num}.1 Pasos del flujo", level=2)
    for p in pasos:
        bullet(doc, p)
    if estados_img:
        doc.add_heading(f"{num}.2 Diagrama de estados", level=2)
        figura(doc, os.path.join(OUT_DIR, estados_img), estados_cap, 6.6)
    nxt = 3 if estados_img else 2
    doc.add_heading(f"{num}.{nxt} Tablas y procedimientos", level=2)
    if tablas:
        par(doc, "Tablas: ", bold=True).add_run(tablas)
    if sps:
        par(doc, "Stored procedures / endpoints principales: ", bold=True).add_run(sps)
    doc.add_page_break()


flujo_seccion(
    2, "Registro de requerimiento",
    "Registrar la necesidad de items (compra/consumo), commodities, activos o "
    "servicios por parte de un area.",
    "Solicitante (Operativo); sistema para asignacion de aprobador.",
    ["El usuario elige tipo de requerimiento: COMPRA (adquisicion) o CONSUMO (uso de "
     "stock), y el tipo de item: items de catalogo, commodities, activo fijo o "
     "activo menor.",
     "Completa la cabecera: empresa (ruc), fundo, area, almacen, centro de costo, "
     "proyecto, prioridad y glosa; agrega el detalle en logistica_dreq.",
     "Adjunta documentos de sustento si corresponde "
     "(LOGISTICA_RequerimientoAdjuntos, en base64).",
     "Si no existe el item, puede solicitar su creacion "
     "(LOGISTICA_SolicitudCreacionItem).",
     "Al guardar, el estado es PENDIENTE; el sistema resuelve el aprobador segun el "
     "flujo del area (LOGISTICA_FlujoAprobacionPorArea + LOGISTICA_UsuariosPorArea).",
     "Si no hay conexion, el requerimiento queda local; el usuario debe sincronizar "
     "pendientes (no puede cerrar sesion con pendientes).",
     "Se genera notificacion al aprobador y/o solicitante."],
    estados_img="est_req.png",
    estados_cap="Figura 2. Estados del requerimiento.",
    tablas="logistica_req, logistica_dreq, LOGISTICA_RequerimientoAdjuntos, "
           "LOGISTICA_ReqCommodity, LOGISTICA_DReqCommodity, "
           "LOGISTICA_SolicitudCreacionItem, LOGISTICA_LogCambioEstadoRequerimiento",
    sps="SPs de registro/sincronizacion de requerimientos; "
        "LOGISTICA_guardarSolicitudesEnriquecidas; sincronizarRequerimientos",
    pantallas="main/requerimientos")

flujo_seccion(
    3, "Aprobacion de requerimientos por area",
    "Validar y autorizar los requerimientos segun la jerarquia del area.",
    "Jefe de Area (aprobador por niveles); solicitante (notificado).",
    ["El aprobador ve la bandeja de pendientes filtrada por su area "
     "(LOGISTICA_UsuariosPorArea con esJefeArea=1).",
     "Revisa el detalle completo: items, cantidades, solicitante, adjuntos y glosa.",
     "Aprueba o rechaza; cada accion se guarda en LOGISTICA_LogAprobacionesPorArea "
     "con nivel, usuario y fecha.",
     "Si el flujo tiene mas niveles (nivelOrden), el documento pasa al siguiente "
     "aprobador; al completarse queda APROBADO.",
     "Existe un flujo hibrido: aprobacion por area o por aprobadores corporativos "
     "segun configuracion (LOGISTICA_ConfigPermisos permite autoaprobacion por "
     "monto).",
     "El rechazo exige motivo y deja el requerimiento ANULADO.",
     "Los cambios de estado se trazan en LOGISTICA_LogCambioEstadoRequerimiento."],
    tablas="logistica_req, LOGISTICA_FlujoAprobacionPorArea, "
           "LOGISTICA_UsuariosPorArea, LOGISTICA_Areas, "
           "LOGISTICA_LogAprobacionesPorArea, logistica_notificaciones",
    sps="SP_listarRequerimientoAprobacion, LOGISTICA_aprobarRequerimientoArea, "
        "SP_obtenerRequerimientosConAprobacion, reporte aprobaciones por area",
    pantallas="main/aprobaciones, main/aprobaciones-area, "
              "main/dashboard-aprobaciones-area, main/reporte-aprobaciones-area")

flujo_seccion(
    4, "Consolidacion de requerimientos",
    "Agrupar items aprobados de multiples requerimientos para cotizar o comprar "
    "consolidado.",
    "Analista de Logistica.",
    ["Los detalles aprobados se marcan como pendientes en "
     "LOGISTICA_ItemPendienteConsolidacion (por tipo: COMPRA, CONSUMO, commodity, "
     "activo fijo/menor, servicio).",
     "Logistica lista pendientes por tab/empresa y filtra por almacen, item o tipo.",
     "Selecciona lineas compatibles y crea la consolidacion "
     "(LOGISTICA_ConsolidacionCab con codigo correlativo CONS-anio-N).",
     "Cada linea consolidada (ConsolidacionDet) mantiene trazabilidad a sus "
     "requerimientos origen (ConsolidacionOrigen).",
     "Los items consolidados pasan a estado CONSOLIDADO; la consolidacion queda "
     "PENDIENTE_COTIZACION.",
     "Puede anularse la consolidacion completa o lineas (con motivo y usuario).",
     "Variante: Consolidacion Corporativa agrupa requerimientos de varias empresas "
     "del grupo (por ruc)."],
    estados_img="est_cons.png",
    estados_cap="Figura 3. Estados de la consolidacion.",
    tablas="LOGISTICA_ItemPendienteConsolidacion, LOGISTICA_ConsolidacionCab, "
           "LOGISTICA_ConsolidacionDet, LOGISTICA_ConsolidacionOrigen, "
           "LOGISTICA_Correlativo, RequerimientoConsolidadoCompra, "
           "DetalleRequerimientoConsolidado, ReqConsolidado",
    sps="02_SP_CONSOLIDACION, 03_SP_CONSOLIDACION_PARTE2, "
        "04_SP_HELPER_POBLAR_PENDIENTES, SP listar commodities/activos pendientes",
    pantallas="main/consolidacion-requerimientos, main/consolidacion-compras, "
              "main/consolidacion-compra, main/consolidacion-servicios, "
              "main/consolidacion-corporativa")

flujo_seccion(
    5, "Cotizaciones",
    "Solicitar y evaluar propuestas de proveedores para la consolidacion.",
    "Analista de Logistica; Proveedor (por correo).",
    ["Desde la consolidacion se genera LOGISTICA_SolicitudCotizacion(+Det) con "
     "correlativo SOL-COT-anio-N.",
     "Se seleccionan proveedores y se envia la solicitud por correo (ConfigSmtp); "
     "cada envio queda en LOGISTICA_CorreoLog.",
     "Se registran las cotizaciones recibidas (LOGISTICA_Cotizacion+Det): precio, "
     "plazo, condiciones de pago.",
     "Se comparan y se marca la cotizacion ganadora (seleccionada).",
     "La solicitud pasa a COTIZADA/CERRADA y la consolidacion a COTIZADO.",
     "Para OC multiempresa existe LOGISTICA_CotizacionOC(+Detalle)."],
    tablas="LOGISTICA_SolicitudCotizacion, LOGISTICA_SolicitudCotizacionDet, "
           "LOGISTICA_Cotizacion, LOGISTICA_CotizacionDet, LOGISTICA_CotizacionOC, "
           "LOGISTICA_CotizacionOCDetalle, LOGISTICA_Proveedor, ProveedoresEmail, "
           "LOGISTICA_CorreoLog",
    sps="05_SP_GENERAR_SOLICITUD_COTIZACION, 09_SP_COTIZACIONES, "
        "SP_REGISTRAR_COTIZACION, SP_LISTAR_COTIZACIONES",
    pantallas="main/cotizaciones, main/evaluacion-proveedores")

flujo_seccion(
    6, "Solicitud de compra y su aprobacion",
    "Formalizar la necesidad de compra y autorizarla por niveles de monto.",
    "Analista de Logistica; Aprobador SC.",
    ["Se genera LOGISTICA_SolicitudCompra(+Det) con numeroSolicitud unico; tipos: "
     "CONSOLIDADA, DIRECTA, URGENTE; prioridad NORMAL/URGENTE/CRITICA.",
     "Se adjuntan documentos obligatorios (LOGISTICA_SolicitudCompraAdjuntos).",
     "Al enviar, queda ENVIADA y entra al flujo de aprobacion SC.",
     "LOGISTICA_FlujoAprobacionSC define niveles por rango de monto "
     "(montoMinimo/montoMaximo); LOGISTICA_UsuariosAprobadoresSC define quien "
     "aprueba cada nivel.",
     "Cada aprobador decide; se registra LOGISTICA_AprobacionSC y "
     "LOGISTICA_LogAprobacionesSC.",
     "Al completar niveles: SC APROBADA. Rechazo: RECHAZADA con motivoRechazo."],
    estados_img="est_sc.png",
    estados_cap="Figura 4. Estados de la solicitud de compra.",
    tablas="LOGISTICA_SolicitudCompra, LOGISTICA_SolicitudCompraDet, "
           "LOGISTICA_SolicitudCompraAdjuntos, LOGISTICA_FlujoAprobacionSC, "
           "LOGISTICA_AprobacionSC, LOGISTICA_UsuariosAprobadoresSC, "
           "LOGISTICA_LogAprobacionesSC",
    sps="241_SP_SOLICITUDES_COMPRA, 250-252_SP_APROBACION_SOLICITUD_COMPRA, "
        "LOGISTICA_listarSolicitudesCompra",
    pantallas="main/solicitudes-compra, main/aprobaciones")

flujo_seccion(
    7, "Orden de compra: generacion, aprobacion y sincronizacion",
    "Emitir la OC al proveedor, autorizarla por niveles y llevarla al ERP.",
    "Analista de Logistica; Aprobador OC (Finanzas/Gerencia); Proveedor; ERP Spring.",
    ["Se genera LOGISTICA_OrdenCompra(+Det) desde la SC aprobada, o "
     "LOGISTICA_ConsolOC(+Det) desde el modulo de consolidacion; ConsolOCDist "
     "distribuye cantidades por requerimiento origen.",
     "La OC lleva proveedor (maestro), moneda (PEN/USD con tipo de cambio), "
     "condiciones/forma de pago, lugar y fecha de entrega; el detalle calcula "
     "precio, descuento, subtotal, IGV y total.",
     "Estado inicial PENDIENTE; al enviarse a aprobacion: PENDIENTE_APROBACION.",
     "LOGISTICA_FlujoAprobacionOC define niveles por monto; cada aprobador actua y "
     "se registra LOGISTICA_AprobacionOC / LogAprobacionesOC / OrdenCompraLog.",
     "Aprobada: se genera el PDF (datos de LOGISTICA_ConfigEmpresa / "
     "ConsolConfigEmpresa / PlantillasPdf), se envia al proveedor por correo y se "
     "marca correoEnviado.",
     "Sincronizacion ERP: el SP inserta la OC en WH_OrdenCompra de Spring "
     "(multiempresa por idEmpresa/ruc). Si falla, queda en "
     "LOGISTICA_ColaMigracionSpring para reintentos y en SincronizacionLog.",
     "Seguimiento: LOGISTICA_SeguimientoOrdenCompra registra las etapas "
     "(emitida, enviada, confirmada, entregada).",
     "Incrementos: si esta habilitado (LOGISTICA_ConfigIncrementoOC), se puede "
     "incrementar cantidades quedando log en LOGISTICA_OCIncrementoLog; puede "
     "requerir reaprobacion."],
    estados_img="est_oc.png",
    estados_cap="Figura 5. Estados de la orden de compra.",
    tablas="LOGISTICA_OrdenCompra, LOGISTICA_OrdenCompraDet, LOGISTICA_OrdenCompraLog, "
           "LOGISTICA_OCAdjunto, LOGISTICA_ConsolOC, LOGISTICA_ConsolOCDet, "
           "LOGISTICA_ConsolOCDist, LOGISTICA_ConsolAdjunto, "
           "LOGISTICA_FlujoAprobacionOC, LOGISTICA_UsuariosAprobadoresOC, "
           "LOGISTICA_AprobacionOC, LOGISTICA_LogAprobacionesOC, "
           "LOGISTICA_SeguimientoOrdenCompra, LOGISTICA_ConfigIncrementoOC, "
           "LOGISTICA_OCIncrementoLog, WH_OrdenCompra, "
           "LOGISTICA_ColaMigracionSpring, LOGISTICA_SincronizacionLog",
    sps="261_SP_ORDEN_COMPRA, 291_SP_APROBACION_OC, LOGISTICA_aprobarRechazarOC, "
        "enviarOCAprobacion, sincronizarOC / sincronizarOCConsolidacion, "
        "601_SINCRONIZAR_OC_CONSOLIDACION_SPRING",
    pantallas="main/ordenes-compra, main/aprobaciones-oc, main/dashboard-finanzas")

flujo_seccion(
    8, "Recepcion de mercaderia e ingreso a inventario",
    "Registrar la recepcion fisica contra la OC y reflejar el ingreso en "
    "kardex/stock.",
    "Almacenero; sistema (actualizacion de inventario y ERP).",
    ["Almacen lista las OCs aprobadas/sincronizadas pendientes de recepcion.",
     "Selecciona la OC, revisa el detalle y registra cantidades recibidas y "
     "aceptadas por item (LOGISTICA_RecepcionOC+Det), con numero de nota de "
     "ingreso (numeroNI) y guia del proveedor.",
     "La recepcion puede ser parcial (OC queda RECIBIDA_PARCIAL) o total "
     "(RECIBIDA_TOTAL).",
     "Se genera LOGISTICA_TransaccionInventario(+Detalle) tipo ingreso; se "
     "actualiza LOGISTICA_Kardex (movimiento, saldos cantidad y valor) y "
     "LOGISTICA_Stock.",
     "El ingreso replica al ERP (SPRING_Stock / SPRING_Kardex / "
     "SPRING_Transacciones).",
     "Se notifica disponibilidad de stock (LOGISTICA_NotificacionesStock / "
     "AlertaStock) y el requerimiento origen puede marcarse COMPLETADO.",
     "Opcionalmente se registra LOGISTICA_ConformidadAlmacen."],
    tablas="LOGISTICA_RecepcionOC, LOGISTICA_RecepcionOCDet, LOGISTICA_RecepcionOCLog, "
           "LOGISTICA_TransaccionInventario, LOGISTICA_DetalleTransaccionInventario, "
           "LOGISTICA_Kardex, LOGISTICA_Stock, SPRING_Stock, SPRING_Kardex, "
           "LOGISTICA_ConformidadAlmacen, logistica_notificaciones",
    sps="281_SP_RECEPCION_OC, 330_RECEPCION_INGRESO_ALMACEN, "
        "LOGISTICA_ingresarRecepcionKardex, SP_INGRESAR_KARDEX_SPRING",
    pantallas="main/recepcion-mercaderia, main/conformidad-almacen")

flujo_seccion(
    9, "Despacho y saldos pendientes",
    "Atender requerimientos de consumo con stock y gestionar los saldos "
    "pendientes.",
    "Almacenero; Jefe de Area/Logistica (resolucion del saldo).",
    ["Almacen lista requerimientos aprobados de consumo pendientes de despacho.",
     "Valida stock por item y registra LOGISTICA_Despachos(+DetalleDespacho).",
     "El sistema descuenta stock y registra la salida en kardex.",
     "Si el despacho no cubre lo solicitado, se genera "
     "LOGISTICA_SaldoPendienteAprobacion(+Det) en estado PENDIENTE para el "
     "usuario creador/aprobador.",
     "El saldo puede: aprobarse (se atiende luego o pasa a consolidacion), "
     "rechazarse o cerrarse (SP_cerrarSaldo).",
     "Se generan notificaciones de stock/almacen."],
    estados_img="est_desp.png",
    estados_cap="Figura 6. Estados del despacho y del saldo pendiente.",
    tablas="LOGISTICA_Despachos, LOGISTICA_DetalleDespacho, LOGISTICA_Kardex, "
           "LOGISTICA_Stock, LOGISTICA_SaldoPendienteAprobacion, "
           "LOGISTICA_SaldoPendienteAprobacionDet, LOGISTICA_ItemPendienteConsolidacion",
    sps="SPs de despacho y saldo pendiente (05_SP_SALDO_REQUERIMIENTO, "
        "SP_cerrarSaldo, registrarSaldoPendienteAprobacion)",
    pantallas="main/despachos, main/despachos-original")

flujo_seccion(
    10, "Servicios: solicitud, orden y conformidad",
    "Gestionar el ciclo de servicios: SS -> aprobacion -> cotizacion -> OS -> "
    "aprobacion -> ejecucion -> conformidad.",
    "Solicitante; Jefe de Area; Analista de Logistica; Aprobador OS; Usuario "
    "Conformidad; Proveedor; ERP Spring.",
    ["Se registra LOGISTICA_SolicitudServicio(+Det+Adjuntos) con tipo, alcance, "
     "ubicacion y fechas.",
     "La SS pasa por aprobacion (LOGISTICA_AprobacionSS).",
     "Logistica cotiza el servicio (LOGISTICA_CotizacionServicio+Det) usando el "
     "historial de precios (LOGISTICA_HistorialPrecioServicio).",
     "Se genera LOGISTICA_OrdenServicio(+Det) o LOGISTICA_ConsolOS(+Det) desde "
     "consolidacion de servicios.",
     "Aprobacion multinivel de la OS (FlujoAprobacionOS / UsuarioAprobadorOS / "
     "AprobacionOS / LogAprobacionesOS).",
     "Aprobada: PDF + correo al proveedor (LOGISTICA_ConsolCorreoLog) y "
     "sincronizacion a WH_OrdenServicio(+Detalle) del ERP.",
     "El proveedor ejecuta el servicio; se registra "
     "LOGISTICA_ConfirmacionServicioRealizado y seguimiento "
     "(SeguimientoOrdenServicio).",
     "El Usuario Conformidad registra CONFORME/NO_CONFORME con firma digital y "
     "acta (LOGISTICA_ConformidadServicio / ConsolOSConformidad), usando la "
     "plantilla configurable (PlantillaConformidadServicio)."],
    estados_img="est_os.png",
    estados_cap="Figura 7. Estados de la orden de servicio.",
    tablas="LOGISTICA_SolicitudServicio(+Det+Adjuntos), LOGISTICA_AprobacionSS, "
           "LOGISTICA_CotizacionServicio(+Det), LOGISTICA_OrdenServicio(+Det), "
           "LOGISTICA_ConsolOS(+Det+Conformidad), LOGISTICA_FlujoAprobacionOS, "
           "LOGISTICA_AprobacionOS, LOGISTICA_ConformidadServicio, "
           "LOGISTICA_ConfirmacionServicioRealizado, "
           "LOGISTICA_SeguimientoOrdenServicio, WH_OrdenServicio(+Detalle)",
    sps="301_SP_SOLICITUD_SERVICIO, 302_SP_APROBACION_SS_COTIZACION, "
        "303_SP_ORDEN_SERVICIO, 304_SP_APROBACION_OS, 304_SP_SINCRONIZAR_OS",
    pantallas="main/solicitudes-servicio, main/ordenes-servicio, "
              "main/aprobaciones-os, main/conformidad-servicios")

flujo_seccion(
    11, "Inventario: kardex, ajustes, transferencias, devoluciones y reingresos",
    "Mantener el inventario correcto y trazado.",
    "Almacenero.",
    ["Kardex: consulta de movimientos por item/almacen/empresa con saldos "
     "cantidad y valor (LOGISTICA_Kardex); stock actual en LOGISTICA_Stock.",
     "Listas de stock: LOGISTICA_ListaStock(+Det) para consulta y control; "
     "LOGISTICA_AlertaStock para niveles criticos.",
     "Ajustes: LOGISTICA_AjusteInventario(+Det) corrige diferencias con motivo.",
     "Conteos fisicos: LOGISTICA_ConteoInventario(+Det) compara stock fisico vs "
     "sistema.",
     "Transferencias: LOGISTICA_RequerimientoTransferencia mueve stock entre "
     "almacenes (incluye almacen virtual); al aprobarse genera salida e ingreso.",
     "Devolucion a proveedor: LOGISTICA_DevolucionProveedor(+Det) retorna "
     "mercaderia de una recepcion/OC.",
     "Devolucion de consumo: logistica_devolucion_consumo(+det) reingresa items "
     "despachados no usados.",
     "Reingreso: LOGISTICA_Reingreso(+Detalle) reincorpora material al almacen.",
     "Todo movimiento genera transaccion de inventario y asiento en kardex."],
    tablas="LOGISTICA_Item, LOGISTICA_Almacen, LOGISTICA_Kardex, LOGISTICA_Stock, "
           "LOGISTICA_TransaccionInventario(+Det), LOGISTICA_ListaStock(+Det), "
           "LOGISTICA_AlertaStock, LOGISTICA_AjusteInventario(+Det), "
           "LOGISTICA_ConteoInventario(+Det), LOGISTICA_RequerimientoTransferencia, "
           "LOGISTICA_DevolucionProveedor(+Det), logistica_devolucion_consumo(+det), "
           "LOGISTICA_Reingreso(+Detalle)",
    sps="27-28_SP_KARDEX, 550_SP_KARDEX_INVENTARIO, 247_SP_GESTION_INVENTARIO, "
        "245_SP_LISTAS_STOCK, 29_SP_DEVOLUCION_PROVEEDOR, "
        "30_SP_DEVOLUCION_CONSUMO, 700_TRANSFERENCIA_ALMACEN_VIRTUAL",
    pantallas="main/kardex, main/gestion-inventario, main/listas-stock, "
              "main/transferencias-almacen, main/devoluciones-proveedores, "
              "main/devoluciones-consumo, main/reingresos")

flujo_seccion(
    12, "Notificaciones, reportes y dashboards",
    "Informar a los usuarios y dar visibilidad de gestion por rol.",
    "Todos los roles.",
    ["Notificaciones: se generan por aprobaciones pendientes, stock disponible, "
     "saldos y eventos (logistica_notificaciones, LOGISTICA_NotificacionesStock); "
     "el contador se refresca cada 30 segundos.",
     "Dashboards por rol: usuario, operativo logistico, jefe/jemma/admin/ti "
     "logistica, compras, despacho, finanzas, aprobaciones por area.",
     "Reportes: saldos, despachos, aprobaciones por area, requerimientos, compras, "
     "reporte logistico; varios exportan a Excel/PDF.",
     "El alcance de datos depende del rol, la empresa (ruc) y el area del usuario."],
    tablas="logistica_notificaciones, LOGISTICA_NotificacionesStock, "
           "LOGISTICA_AlertaStock, LOGISTICA_CorreoLog + consultas de reporte",
    sps="LOGISTICA_obtenerKPIsRequerimientos, obtenerRequerimientosPorEstado, "
        "obtenerEntregasHoy, obtenerGastosMensuales, obtenerSeguimientoOrdenes, "
        "242_SP_DASHBOARD_COMPRAS, 243_SP_REPORTES_COMPRAS, 506_SP_REPORTES_DASHBOARD",
    pantallas="dashboards-* y reporte-* del frontend")

flujo_seccion(
    13, "Administracion y configuracion",
    "Parametrizar seguridad, flujos, menus y configuracion general.",
    "Administrador Logistica; Administrador Sistema.",
    ["Usuarios: LOGISTICA_AdminUsuarios (login admin con hash), LOGISTICA_Usuarios "
     "y esquema legacy; migracion de usuarios desde el ERP.",
     "Asignacion: LOGISTICA_UsuariosPorArea (esJefeArea) y LOGISTICA_Areas.",
     "Flujos: por area (requerimientos) y por monto (SC/OC/OS).",
     "Menu dinamico y permisos: LOGISTICA_ConfigMenu, LOGISTICA_ConfigPermisos, "
     "LOGISTICA_ConfiguracionPermiso.",
     "Maestros: items (cache del ERP), proveedores, unidades de medida, "
     "correlativos.",
     "Plantillas y branding: ConfigSmtp, PlantillasCorreo, PlantillasPdf, "
     "ConfigBranding, ConfigEmpresa, ConsolConfigEmpresa.",
     "Parametros y auditoria: LOGISTICA_Parametros, LOGISTICA_Auditoria, "
     "AppDeployments (versiones desplegadas)."],
    tablas="LOGISTICA_AdminUsuarios, LOGISTICA_Usuarios*, LOGISTICA_UsuariosPorArea, "
           "LOGISTICA_Areas, flujos de aprobacion, LOGISTICA_ConfigMenu, "
           "LOGISTICA_ConfigPermisos, LOGISTICA_Parametros, ConfigSmtp, "
           "PlantillasCorreo, PlantillasPdf, LOGISTICA_ConfigBranding, "
           "LOGISTICA_ConfigEmpresa, LOGISTICA_UnidadMedida, LOGISTICA_Correlativo, "
           "LOGISTICA_Auditoria, AppDeployments",
    sps="08_SP_LOGIN_ADMIN, 252_SP_CONFIGURACION_APROBADORES_SC, "
        "286_CONFIGURACION_ROLES_APROBADORES, 654_CONFIG_PERMISOS_ROLES, "
        "800_TABLA_CONFIGMENU_DINAMICO",
    pantallas="admin-login, admin-logistica, administracion, main/maestros, "
              "main/aprobadores, main/parametros")

flujo_seccion(
    14, "Integraciones externas",
    "Sincronizar informacion con sistemas externos.",
    "ERP Spring; API Maestra; Servidor SMTP; Servicio de tipo de cambio.",
    ["Autenticacion: el login valida credenciales contra la API Maestra.",
     "OC/OS aprobadas: se insertan en WH_OrdenCompra / WH_OrdenServicio del ERP "
     "Spring; si falla quedan en LOGISTICA_ColaMigracionSpring con reintentos y "
     "trazado en LOGISTICA_SincronizacionLog.",
     "Maestro de items: job nocturno sincroniza el maestro del ERP a "
     "LOGISTICA_MaestroItemCache/MaestroItemSync.",
     "Inventario: recepciones y despachos replican a SPRING_Stock/Kardex/"
     "Transacciones.",
     "Correos: solicitudes de cotizacion y ordenes se envian via SMTP; resultado "
     "en LOGISTICA_CorreoLog / ConsolCorreoLog.",
     "Tipo de cambio: consulta para documentos en USD."],
    tablas="WH_OrdenCompra, WH_OrdenServicio(+Detalle), SPRING_Stock, SPRING_Kardex, "
           "SPRING_Transacciones, LOGISTICA_ColaMigracionSpring, "
           "LOGISTICA_SincronizacionLog, LOGISTICA_MaestroItemCache, "
           "LOGISTICA_MaestroItemSync",
    sps="sincronizarOC, sincronizarOCConsolidacion, sincronizarOrdenServicio, "
        "JOB_SincronizarMaestroItem_11pm",
    pantallas="-")

par(doc, "Fin del documento.", italic=True, size=9, center=True)

out3 = os.path.join(DOCS_DIR, "DOCUMENTACION_FLUJOS_PROCESO.docx")
doc.save(out3)
print("OK", out3)
print("=== FIN ===")
