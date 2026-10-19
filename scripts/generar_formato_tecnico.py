# -*- coding: utf-8 -*-
"""
Genera el Excel TECNICO de configuracion del app LOGISTICA.
Las hojas replican 1:1 las tablas de la BD para que el implementador
pueda preparar los INSERTs/llamadas a SPs directamente.

Salida: ../docs/FORMATO_TECNICO_CONFIGURACION_LOGISTICA.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = "1F4E78"
AZUL_CLARO = "DDEBF7"
GRIS = "404040"
AMARILLO = "FFF2CC"
VERDE_CLARO = "E2EFDA"

F_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_NORMAL = Font(name="Consolas", size=9)
F_EJEMPLO = Font(name="Consolas", size=9, italic=True, color="7F7F7F")
F_NOTA = Font(name="Calibri", size=9, italic=True, color="595959")
F_TABLA = Font(name="Consolas", size=9, bold=True, color="1F4E78")

FILL_HEADER = PatternFill("solid", fgColor=GRIS)
FILL_EJEMPLO = PatternFill("solid", fgColor=AMARILLO)
FILL_SECCION = PatternFill("solid", fgColor=AZUL_CLARO)

BORDE = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

MAX_FILAS = 500


def hoja_tabla(wb, nombre, tabla, headers, anchos, nota):
    """Hoja estilo tabla BD: fila1 = tabla destino, fila2 = nota, fila3 = columnas."""
    ws = wb.create_sheet(nombre)
    n = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    c = ws.cell(row=1, column=1, value=tabla)
    c.font = Font(name="Consolas", size=12, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=GRIS)
    c.alignment = CENTER
    ws.row_dimensions[1].height = 22

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    cn = ws.cell(row=2, column=1, value=nota)
    cn.font = F_NOTA
    cn.alignment = LEFT
    ws.row_dimensions[2].height = 30

    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=i, value=h)
        cell.font = F_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = CENTER
        cell.border = BORDE
    ws.row_dimensions[3].height = 30

    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    return ws


def escribir_filas(ws, filas, inicio=4, ejemplo=False):
    for r, fila in enumerate(filas, start=inicio):
        for c, val in enumerate(fila, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = F_EJEMPLO if ejemplo else F_NORMAL
            cell.border = BORDE
            cell.alignment = LEFT
            if ejemplo:
                cell.fill = FILL_EJEMPLO


def validacion(ws, formula, col, ini=4, fin=MAX_FILAS):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}{ini}:{col}{fin}")


wb = Workbook()
wb.remove(wb.active)

# ================================================================
# README_TECNICO
# ================================================================
ws = wb.create_sheet("README_TECNICO")
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 44
ws.column_dimensions["D"].width = 55

ws.merge_cells("B2:D2")
c = ws["B2"]
c.value = "FORMATO TECNICO - CONFIGURACION APP LOGISTICA (carga directa a BD)"
c.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=GRIS)
c.alignment = CENTER
ws.row_dimensions[2].height = 28

secciones = [
    ("USO", "", ""),
    ("", "Cada hoja replica una tabla de la BD LOGISTICA con los nombres de columna reales.",
     "Completar los datos y generar INSERT/MERGE o llamar a los SPs de administracion."),
    ("", "Las columnas IDENTITY (idarea, idFlujo, id, etc.) se muestran solo como referencia de FK.",
     "En AREAS se puede fijar idarea deseado para que las demas hojas lo referencien."),
    ("", "", ""),
    ("ORDEN DE CARGA (dependencias)", "", ""),
    ("1", "EMPRESAS", "RUCs validos; no hay tabla propia, es referencia para el resto."),
    ("2", "AREAS", "LOGISTICA_Areas + LOGISTICA_SubAreas"),
    ("3", "USUARIOS_POR_AREA", "LOGISTICA_UsuariosPorArea  (FK: ruc, idarea)"),
    ("4", "FLUJO_REQ_AREA", "LOGISTICA_FlujoAprobacionPorArea  (FK: ruc, idarea)"),
    ("5", "FLUJO_OC + APROBADORES_OC", "LOGISTICA_FlujoAprobacionOC + LOGISTICA_UsuariosAprobadoresOC (join por rol)"),
    ("6", "FLUJO_OS + APROBADORES_OS", "LOGISTICA_FlujoAprobacionOS + LOGISTICA_UsuarioAprobadorOS"),
    ("", "", ""),
    ("MAPEO HOJA -> TABLA -> SP / ENDPOINT", "", ""),
    ("AREAS", "LOGISTICA_Areas", "admin-logistica/areas/crear  |  LOGISTICA_guardarArea"),
    ("SUBAREAS", "LOGISTICA_SubAreas", "admin-logistica/subareas/crear"),
    ("USUARIOS_POR_AREA", "LOGISTICA_UsuariosPorArea", "admin-logistica/usuarios-por-area/crear | asignar-usuario-area"),
    ("FLUJO_REQ_AREA", "LOGISTICA_FlujoAprobacionPorArea", "guardar-flujo-aprobacion-area | LOGISTICA_guardarFlujoAprobacionArea"),
    ("FLUJO_OC", "LOGISTICA_FlujoAprobacionOC", "aprobacion-oc/flujos/guardar | LOGISTICA_guardarFlujoAprobacionOC"),
    ("APROBADORES_OC", "LOGISTICA_UsuariosAprobadoresOC", "aprobacion-oc/usuarios/guardar | LOGISTICA_guardarUsuarioAprobadorOC"),
    ("FLUJO_OS", "LOGISTICA_FlujoAprobacionOS", "aprobacion-os/crear-flujo | LOGISTICA_crearFlujoAprobacionOS"),
    ("APROBADORES_OS", "LOGISTICA_UsuarioAprobadorOS", "aprobacion-os/crear-usuario-aprobador | LOGISTICA_crearUsuarioAprobadorOS"),
    ("", "", ""),
    ("CONVENCIONES", "", ""),
    ("bit", "activo / esJefeArea / esAprobador / mostrarAdmision", "1 = verdadero, 0 = falso"),
    ("documentoidentidad", "Clave de usuario en todo el sistema (DNI)", "varchar(20) - puede tener ceros a la izquierda"),
    ("rol", "Codigos validos en hoja CATALOGOS_TEC", "ej: JEFE_AREA, APLOGIST, FINANZAS, ADLOGIST"),
    ("tipoRequerimiento", "CONSUMO / COMPRA / SERVICIO / TRANSFERENCIA / COMPRA_AREA", "tipoRequerimiento = COMMODITY se mapea a SERVICIO en el SP"),
    ("nivel / nivelMaximo / niveles", "Orden de aprobacion multinivel OC/OS", "1 aprueba primero; niveles = total de firmas del rango"),
    ("montos", "montoMinimo/montoMaximo, montoDesde/montoHasta", "decimal(18,2) en soles; usar 999999999 como tope abierto"),
    ("", "", ""),
    ("NOTAS", "", ""),
    ("", "Las tablas de flujo OC/OS NO tienen columna ruc: la configuracion es global por base de datos.",
     "Si el grupo usa una sola BD LOGISTICA, el rango aplica a todas las empresas."),
    ("", "Filas amarillas = configuracion actual del sistema (referencia).", ""),
    ("", "Al aprobarse el levantamiento, este archivo alimenta los scripts de carga inicial.", ""),
]
fila = 4
for a, b, cc in secciones:
    ba = ws.cell(row=fila, column=2, value=a)
    bb = ws.cell(row=fila, column=3, value=b)
    bc = ws.cell(row=fila, column=4, value=cc)
    if b == "" and cc == "" and a:
        ba.font = Font(name="Calibri", size=11, bold=True, color=AZUL)
        ba.fill = FILL_SECCION
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
    else:
        ba.font = Font(name="Calibri", size=10, bold=True)
        ba.alignment = Alignment(horizontal="right", vertical="top")
        bb.font = Font(name="Consolas", size=9)
        bb.alignment = LEFT
        bc.font = F_NOTA
        bc.alignment = LEFT
    fila += 1

# ================================================================
# EMPRESAS (referencia - no tabla propia)
# ================================================================
ws = hoja_tabla(
    wb, "EMPRESAS", "EMPRESAS (referencia - el ruc viaja en todas las tablas)",
    ["ruc", "razonSocial", "codigoInterno", "activo"],
    [16, 45, 15, 10],
    "Referencia. El campo ruc (varchar 11/20) es obligatorio en las demas hojas.",
)
escribir_filas(ws, [
    ["20481121966", "HASS PERU S.A.", "HP", 1],
    ["20610773274", "BERRY HARVEST S.A.C.", "BH", 1],
])
validacion(ws, '"1,0"', "D")

# ================================================================
# AREAS -> LOGISTICA_Areas
# ================================================================
ws = hoja_tabla(
    wb, "AREAS", "LOGISTICA_Areas  (idarea es IDENTITY: fijarlo aqui para referenciarlo en otras hojas)",
    ["ruc", "idarea", "codigo", "nombre", "descripcion", "mostrarAdmision", "activo"],
    [15, 9, 12, 32, 32, 15, 9],
    "SP: admin-logistica/areas/crear. idarea conocidos en BD: 1-RIEGO, 2-LABORES, 3-BPM, 4-SST, 5-ALMACEN, 15-SERVICIO DE TRANSPORTE (ruc 20481121966).",
)
escribir_filas(ws, [
    ["20481121966", 1, "RIE", "RIEGO", "RIEGO", 1, 1],
    ["20481121966", 2, "LAB", "LABORES", "LABORES", 1, 1],
    ["20481121966", 3, "BPM", "BPM", "BPM", 1, 1],
    ["20481121966", 4, "SST", "SST", "SST", 1, 1],
    ["20481121966", 5, "ALM", "ALMACÉN", "ALMACÉN", 1, 1],
    ["20481121966", 15, "TRA", "SERVICIO DE TRANSPORTE", "SERVICIO DE TRANSPORTE", 1, 1],
    ["20481121966", "", "ACO", "ACOPIOS", "ACOPIOS", 1, 1],
    ["20481121966", "", "LOG", "LOGISTICA", "LOGISTICA", 1, 1],
    ["20481121966", "", "LIC", "LICITACIONES", "LICITACIONES", 1, 1],
    ["20481121966", "", "COS", "COSTOS", "COSTOS", 1, 1],
    ["20481121966", "", "SIS", "SISTEMAS", "SISTEMAS", 1, 1],
    ["20481121966", "", "TES", "TESORERÍA", "TESORERÍA", 1, 1],
    ["20481121966", "", "CON", "CONTABILIDAD", "CONTABILIDAD", 1, 1],
    ["20481121966", "", "AF", "ADMINISTRACIÓN Y FINANZAS", "ADMINISTRACIÓN Y FINANZAS", 1, 1],
    ["20481121966", "", "GER", "GERENCIA", "GERENCIA", 1, 1],
    ["20481121966", "", "RRHH", "RECURSOS HUMANOS", "RECURSOS HUMANOS", 1, 1],
    ["20481121966", "", "CEX", "COMERCIO EXTERIOR", "COMERCIO EXTERIOR", 1, 1],
    ["20610773274", "", "ALM", "ALMACÉN", "ALMACÉN", 1, 1],
    ["20610773274", "", "ACO", "ACOPIOS", "ACOPIOS", 1, 1],
    ["20610773274", "", "BPA", "BPA", "BPA", 1, 1],
])
validacion(ws, "EMPRESAS!$A$4:$A$20", "A")
validacion(ws, '"1,0"', "F")
validacion(ws, '"1,0"', "G")

# ================================================================
# SUBAREAS -> LOGISTICA_SubAreas
# ================================================================
ws = hoja_tabla(
    wb, "SUBAREAS", "LOGISTICA_SubAreas",
    ["ruc", "idsubarea", "idarea", "nombre", "estado"],
    [15, 11, 9, 32, 9],
    "SP: admin-logistica/subareas/crear. Opcional: solo si el area se subdivide (ej: fundos dentro de RIEGO).",
)
validacion(ws, "EMPRESAS!$A$4:$A$20", "A")
validacion(ws, '"1,0"', "E")

# ================================================================
# USUARIOS_POR_AREA -> LOGISTICA_UsuariosPorArea
# ================================================================
ws = hoja_tabla(
    wb, "USUARIOS_POR_AREA",
    "LOGISTICA_UsuariosPorArea",
    ["ruc", "documentoidentidad", "nombreCompleto", "idarea", "idsubarea",
     "rol", "esJefeArea", "esAprobador", "email", "telefono", "activo"],
    [15, 20, 32, 9, 10, 14, 11, 12, 30, 12, 9],
    "SP: admin-logistica/usuarios-por-area/crear o asignar-usuario-area. "
     "UK: (ruc + documentoidentidad + idarea) -> repetir fila si el usuario esta en varias areas. "
     "rol: ver CATALOGOS_TEC.",
)
escribir_filas(ws, [
    ["20481121966", "48073280", "JULCA BARRANTES JULIO CESAR", 15, "", "APLOGIST", 1, 1, "", "", 1],
    ["20481121966", "76019678", "CARRANZA LLANOS ANTONY JUNIOR", 15, "", "APLOGIST", 1, 1, "", "", 1],
    ["20481121966", "081279", "REYES RUIZ PABLO", 5, "", "APLOGIST", 1, 1, "", "", 1],
    ["20481121966", "60796", "SALVADOR PALACIOS EMILI LEONOR", 5, "", "APLOGIST", 1, 1, "", "", 1],
    ["20481121966", "74895320", "MIRANDA FLORES DANIEL MARCOS", "", "", "APLOGIST", 1, 1, "", "", 1],
    ["20610773274", "47857722", "GUTIERREZ GUZMAN LUISA MARIA", "", "", "APLOGIST", 1, 1, "", "", 1],
    ["20610773274", "47227505", "GUEVARA CONTRERAS FRANKLIN IVAN", "", "", "APLOGIST", 1, 1, "", "", 1],
], ejemplo=True)
validacion(ws, "EMPRESAS!$A$4:$A$20", "A")
validacion(ws, "CATALOGOS_TEC!$A$4:$A$25", "F")
for col in ["G", "H", "K"]:
    validacion(ws, '"1,0"', col)

# ================================================================
# CUENTAS_ACCESO -> LOGISTICA_AdminUsuarios (login a la app)
# ================================================================
ws = hoja_tabla(
    wb, "CUENTAS_ACCESO",
    "LOGISTICA_AdminUsuarios  (cuentas de login del panel/app)",
    ["ruc", "usuario (DNI)", "clave (texto plano)", "nombre", "email",
     "idrol", "rol (nombre)", "nrodocumento", "idempresa", "activo"],
    [15, 14, 16, 32, 30, 12, 24, 14, 11, 9],
    "La clave en texto plano se hashea SHA2_256 al cargar. "
    "idrol: ver CATALOGOS_TEC. idempresa: codigo interno ERP (ej: 000008).",
)
escribir_filas(ws, [
    ["20481121966", "48073280", "CLAVE_INICIAL", "JULCA BARRANTES JULIO CESAR",
     "jjulca@hassperu.com", "APLOGIST", "Aprobador Consumo", "48073280", "000008", 1],
], ejemplo=True)
validacion(ws, "EMPRESAS!$A$4:$A$20", "A")
validacion(ws, "CATALOGOS_TEC!$A$4:$A$25", "F")
validacion(ws, '"1,0"', "J")

# ================================================================
# FLUJO_REQ_AREA -> LOGISTICA_FlujoAprobacionPorArea
# ================================================================
ws = hoja_tabla(
    wb, "FLUJO_REQ_AREA",
    "LOGISTICA_FlujoAprobacionPorArea",
    ["ruc", "idarea", "tipoRequerimiento", "rolAprobador", "activo"],
    [15, 9, 20, 16, 9],
    "SP: guardar-flujo-aprobacion-area. Una fila por (ruc+idarea+tipo). "
    "tipoRequerimiento: CONSUMO/COMPRA/SERVICIO/TRANSFERENCIA/COMPRA_AREA. "
    "COMMODITY en el requerimiento mapea a SERVICIO.",
)
flujo_actual = []
for ida in [1, 2, 3, 4, 5]:
    for t in ["CONSUMO", "COMPRA", "SERVICIO"]:
        flujo_actual.append(["20481121966", ida, t, "JEFE_AREA", 1])
escribir_filas(ws, flujo_actual, ejemplo=True)
validacion(ws, "EMPRESAS!$A$4:$A$20", "A")
validacion(ws, '"CONSUMO,COMPRA,SERVICIO,TRANSFERENCIA,COMPRA_AREA"', "C")
validacion(ws, "CATALOGOS_TEC!$A$4:$A$25", "D")
validacion(ws, '"1,0"', "E")

# ================================================================
# FLUJO_OC -> LOGISTICA_FlujoAprobacionOC
# ================================================================
ws = hoja_tabla(
    wb, "FLUJO_OC",
    "LOGISTICA_FlujoAprobacionOC  (config global - sin ruc)",
    ["montoMinimo", "montoMaximo", "nivel", "rolAprobador", "nombreRol", "descripcion", "activo"],
    [14, 14, 8, 16, 22, 45, 9],
    "SP: aprobacion-oc/flujos/guardar | LOGISTICA_guardarFlujoAprobacionOC. "
    "Al generar la OC se busca el rango del montoTotal y se asigna el usuario con ese rol "
    "(join con LOGISTICA_UsuariosAprobadoresOC.rol).",
)
escribir_filas(ws, [
    [0, 10000, 1, "FINANZAS", "Finanzas", "Aprobación OC hasta S/ 10,000", 1],
    [10001, 50000, 1, "FINANZAS", "Finanzas", "Aprobación OC S/ 10,001 - 50,000", 1],
    [50001, 999999999, 1, "FINANZAS", "Finanzas", "Aprobación OC mayor a S/ 50,000", 1],
], ejemplo=True)
validacion(ws, "CATALOGOS_TEC!$A$4:$A$25", "D")
validacion(ws, '"1,0"', "G")

# ================================================================
# APROBADORES_OC -> LOGISTICA_UsuariosAprobadoresOC
# ================================================================
ws = hoja_tabla(
    wb, "APROBADORES_OC",
    "LOGISTICA_UsuariosAprobadoresOC",
    ["documentoIdentidad", "nombreCompleto", "email", "rol",
     "montoMinimo", "montoMaximo", "nivel", "activo"],
    [18, 32, 30, 14, 14, 14, 8, 9],
    "SP: aprobacion-oc/usuarios/guardar | LOGISTICA_guardarUsuarioAprobadorOC. "
    "El rol debe existir en FLUJO_OC para que el motor lo asigne.",
)
escribir_filas(ws, [
    ["DNI_AQUI", "NOMBRE APROBADOR OC", "correo@empresa.com", "FINANZAS", 0, 999999999, 1, 1],
], ejemplo=True)
validacion(ws, "CATALOGOS_TEC!$A$4:$A$25", "D")
validacion(ws, '"1,0"', "H")

# ================================================================
# FLUJO_OS -> LOGISTICA_FlujoAprobacionOS
# ================================================================
ws = hoja_tabla(
    wb, "FLUJO_OS",
    "LOGISTICA_FlujoAprobacionOS  (config global - sin ruc)",
    ["nombre", "montoDesde", "montoHasta", "niveles", "activo"],
    [35, 14, 14, 9, 9],
    "SP: aprobacion-os/crear-flujo | LOGISTICA_crearFlujoAprobacionOS. "
    "niveles = cuantas firmas requiere la OS en ese rango de monto.",
)
escribir_filas(ws, [
    ["OS hasta S/ 10,000", 0, 10000, 1, 1],
    ["OS S/ 10,001 a S/ 50,000", 10001, 50000, 1, 1],
    ["OS mayor a S/ 50,000", 50001, 999999999, 1, 1],
], ejemplo=True)
validacion(ws, '"1,0"', "E")

# ================================================================
# APROBADORES_OS -> LOGISTICA_UsuarioAprobadorOS
# ================================================================
ws = hoja_tabla(
    wb, "APROBADORES_OS",
    "LOGISTICA_UsuarioAprobadorOS",
    ["documentoIdentidad", "nombre", "rol", "nivelMaximo", "activo"],
    [18, 32, 14, 12, 9],
    "SP: aprobacion-os/crear-usuario-aprobador | LOGISTICA_crearUsuarioAprobadorOS. "
    "nivelMaximo = hasta que nivel de firma puede aprobar esta persona.",
)
escribir_filas(ws, [
    ["DNI_AQUI", "NOMBRE APROBADOR OS", "FINANZAS", 3, 1],
], ejemplo=True)
validacion(ws, "CATALOGOS_TEC!$A$4:$A$25", "C")
validacion(ws, '"1,0"', "E")

# ================================================================
# CATALOGOS_TEC
# ================================================================
ws = wb.create_sheet("CATALOGOS_TEC")
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 4
ws.column_dimensions["E"].width = 20
ws.column_dimensions["F"].width = 45

ws.merge_cells("A1:F1")
c = ws["A1"]
c.value = "CATALOGOS TECNICOS - valores validos en BD"
c.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=GRIS)
c.alignment = CENTER

ws["A2"] = "ROLES (codigo en BD)"
ws["A2"].font = Font(bold=True, color=AZUL)
for i, h in enumerate(["codigo", "nombre", "descripcion"], start=1):
    cell = ws.cell(row=3, column=i, value=h)
    cell.font = F_HEADER; cell.fill = FILL_HEADER; cell.border = BORDE; cell.alignment = CENTER
roles = [
    ("TILOGIST", "Admin Sistema", "TI. Acceso total a configuracion."),
    ("ADLOGIST", "Admin Logistica", "Operativo logistica completo; puede aprobar si se configura."),
    ("JLOLOGIST", "Jefatura Logistica", "Visualizacion y supervision operativa."),
    ("JEMLOGIST", "Jefatura Licitaciones", "Jefatura licitaciones/empaque."),
    ("LOLOGIST", "Operador Logistica", "Solicitudes, cotizaciones, ordenes."),
    ("EMLOGIST", "Operador Licitaciones", "Operador licitaciones/empaque."),
    ("ALLOGIST", "Almacen", "Despachos, ingresos, kardex."),
    ("OPLOGIST", "Operativo Campo", "Requerimientos desde campo."),
    ("APLOGIST", "Aprobador Consumo", "Aprueba requerimientos de su area (uso actual en UsuariosPorArea)."),
    ("JEFE_AREA", "Jefe de Area", "Rol aprobador usado en LOGISTICA_FlujoAprobacionPorArea."),
    ("SOLICITANTE", "Solicitante", "Crea requerimientos."),
    ("FINANZAS", "Finanzas", "Unico aprobador OC/OS en la config actual."),
    ("GERENTE", "Gerencia", "Nivel superior opcional."),
]
fila = 4
for r in roles:
    for i, v in enumerate(r, start=1):
        cell = ws.cell(row=fila, column=i, value=v)
        cell.font = F_NORMAL; cell.border = BORDE; cell.alignment = LEFT
    fila += 1

ws["E2"] = "TIPO REQUERIMIENTO"
ws["E2"].font = Font(bold=True, color=AZUL)
for i, h in enumerate(["valor", "nota"], start=5):
    cell = ws.cell(row=3, column=i, value=h)
    cell.font = F_HEADER; cell.fill = FILL_HEADER; cell.border = BORDE; cell.alignment = CENTER
tipos = [
    ("CONSUMO", "Salida de almacen para uso interno."),
    ("COMPRA", "Deriva a solicitud de compra / consolidacion."),
    ("SERVICIO", "COMMODITY del requerimiento mapea aqui."),
    ("TRANSFERENCIA", "Entre almacenes / empresas."),
    ("COMPRA_AREA", "Tipo adicional usado en flujo por area (aprobacion financiera/gerencia)."),
]
fila = 4
for t, d in tipos:
    ws.cell(row=fila, column=5, value=t).font = F_NORMAL
    ws.cell(row=fila, column=5).border = BORDE
    ws.cell(row=fila, column=6, value=d).font = F_NORMAL
    ws.cell(row=fila, column=6).border = BORDE
    fila += 1

ws["E11"] = "BITS"
ws["E11"].font = Font(bold=True, color=AZUL)
ws["E12"] = "Columnas bit: usar 1 o 0"
ws["E12"].font = F_NOTA
ws.merge_cells("E12:F12")

# ================================================================
wb._sheets = [wb[n] for n in [
    "README_TECNICO", "EMPRESAS", "AREAS", "SUBAREAS", "CUENTAS_ACCESO",
    "USUARIOS_POR_AREA", "FLUJO_REQ_AREA", "FLUJO_OC", "APROBADORES_OC",
    "FLUJO_OS", "APROBADORES_OS", "CATALOGOS_TEC",
]]

salida = os.path.join(os.path.dirname(__file__), "..", "docs",
                      "FORMATO_TECNICO_CONFIGURACION_LOGISTICA.xlsx")
salida = os.path.abspath(salida)
wb.save(salida)
print(f"Excel tecnico generado: {salida}")
